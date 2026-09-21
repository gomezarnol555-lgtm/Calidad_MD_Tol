import streamlit as st
import pandas as pd
import altair as alt
import hashlib, os, base64, json, re
from supabase_db import conn, read_df, exec_sql, reset_autoincrement, upload_evidence, table_columns, DatabaseIntegrityError
from io import BytesIO
from openpyxl.styles import Font, PatternFill
from datetime import datetime, date, timedelta
from pathlib import Path
from uuid import uuid4

APP_NAME = "Calidad MD | PNC y ME"
LOGO_MUNDO_DULCE_BASE64 = "iVBORw0KGgoAAAANSUhEUgAAAMcAAABPCAYAAABf5tFzAAAAAXNSR0IArs4c6QAAAARnQU1BAACxjwv8YQUAAAAJcEhZcwAAFxEAABcRAcom8z8AADD/SURBVHhe7X0HmFRVtvW8iY6TnqOC5FSpxRwGRVFAgqAoSjQnTDOYMI0zijkHJKOAEROoiBIEURRJgkiSoAgSW2iq6lZ1TrXX+9Y+51RXXxqkMf6/tb/vUE3VjefudXbe91e/+pkTgP8B8DcRtBfBwFQKd4sgD7shEQzntmZIiPv7j5ulLP0/SQB+LSIHiuAmEbwNIAqgPM38qSUoKRqG0qJRKC0ZhVJ5CqWlo1C6fQTKit4CkKoCClAggqgAgwX4j4j08J8vS1n62ROlhIh0EsFkCwgAxagsewklhf9GMno8krHWSEQPh7ejMbz8JvA2N0F8TCPErm6I6KH14H0WRDLZGsloKxR4fVBeORopTEuDRQEjGEGJ4j9/lrL0syMHCkBBUQGUorL8ZRTl90cyeiy8vGY6EjtCZnhhJLbmwHs6jFjbIGIHBRBrEED8iTAS8YjdJoREfgDeZ40QvzkA7+hW8Fodj8rPP3UASRr1Sy4EsC+llf+6spSln5QAHCCCdxwoykvGIj92Mry8JvDymltA5JgRz0EimQPvwwhi7YKINQwYYHQMwpsVQSKRg0Q0B4miHCSWRxC/JYRYywBi9Vsg1rAZYq2aIrH0KBTkn43K1NQqSQLkisiZ/mvLUpZ+MhKRngJ8QxuhrGIskgqKpkjsCFYBwo1EDrxNOYjfFEIsFECsUQCxxgHETg0isSaCREEOEjEDIO/VMOInBxGrZ7dpFkB8YAiJzQRXGIn8FvA+b46Cpy9B8X33Arl5lCRrs6pWln5yEpG/CjBZaGQnP0XBmEsRn9MAiXgNoKAkoLSYFUHsxKBh9qZmeIPDSOTmIOFZqRLLQfzBkKpYsSYWGCcG4H0UMb+X5iCxNgfxe8KIHRlAtE5j7Pj7/iie/LBTtb4CEPFfb5ay9KOQBcYMMmPFzJmItQojWqeh2g7K6ARDphpFhr85hFjEMjyB0SIAb2gYiUK7DdWoJRGVIvo7t6tvpcXGHLNdYQ68t8OIdw0iVtfYKLGjA/CmEmAhFHhnA8iliuWJyN3+685Sln5QIjAATEdlJQquvhqxYBCxJs0MMx8bRGJLBjhoP2zLQfwGKwmaW1WKatSiiJEWTo16LYzYSdYGqR9A7PigkSo8BoHzZQ7i94YNcLhN4wDi14eQWBdRqUQj3ks0Rf6OjoDkWiki9/mvP0tZ+kHISgwFRv5llyFarx5izZurPRALBOC9EjYMT2AUWDWKdgOB0yJg7IfOQSS+iCCRb1UpLwfx+8MGEFaqxC8KwiN4SnLMdssiiHex0oLgUmkRMSDk7wTH+hzEB4URbd0IiXPaQ1Jb0gChJ81/L1nK0vdGAPYTkelIlSG/Xz9E69dXQCjThy0wqPpkAuMIy8yUGC0CiF8VQmK1ZWaOxRHEqCI5NYrS4J6QOQ6ZPj8H3ogwYodaadEooFIosTaSVrMSX+cgfnsIseOs8b5/ALFLGyE/0TETIFkvVpZ+OBKRsWS0ojcHIlq/rgEGx0EBeI+FzSrvgEE37eGW4QmeJgHEh4eRKLbSgoxPYJwQNOoWGb+NdeU6+2NNBPFzgkYqUapQWkyx0oLn2piD+JNhVb90f46WAcT/HTLxk2gz5Od3hGArbZAvs16sLP0gxHQNgRSVfz0TsYNzjCpFYHCl7xM0BjNtBwcMSgxneDtgODdtMgfxB0KIHWYNakoDMvSXZn/1as20hrkLDA7g71ZaWKkU72nVLJ6DHq3DAvAmWZBayeTNborkS6cAW7dTeqwTkYb+e8tSlvaaRCQsjOyVbkH84oMRa9zMAIMr+tFBoybRaM5UpQgMqlPtgoi/xJiE3SZBF2zIrPIExnFBtROUmblNYQ7iw6zRTcY/JlglLSh1qEINCBlpREBQXSN4Lg3CW23tGF4HwdUniFhOENE/HoTCJwc49Wqo//6ylKW9IhH5i4hMQXExCi6/AtH6jaqrU49bicDg3kcZqlTDAOL0SK21dgGZ9tMI4mda+4LAaBNEYmnEMD3B8WUE8VtDhukPCiDeP4TEiohRsYpz4L0QNnYFgUVwEHw5AXivRtQjpsehZ+wa6zImuALmON6XzVBZMYMxkAIRCfrvM0tZqjXZnCUU3n47onUz7AwycM8gEhusRKDHifaBtR/SwCDTcyWnfdE6aGwH/k6j27lgCa41EU0fiR0Q0MBe/N6QsT2oIlFaXBw05yXwnLToEoT3TsQEBUsY/4iYHC2Chr+3C8J72apZsRYo8Hq4IOGorPcqS9+JAPwvgGWyaRPiLVtW2RlctRnAm2lX9biNYxAwXNUdMJy3aU4GMOimfZDJhFVqlEc16tggYnWs0T03gkSZlRYv29gHQcfzUpWj0X1TCIntFjzfWGmRY6UJPVYE7iZ7DjoAkhF4XgtUiMYtqV6d57/fLGVpj0nKy49PSw3ntuUgAC4ImVU/mhHg8wOD9sV9lmnJ+CcH4H1gA3/8/asceA+GzX51AhpBT9BuoHq0IQfekHDa06XndWoUbRBKi9IcNcA1eZG/8TiMtYwOm0Ckk1qMkTB5sUtTJHp1Ajar92o5gN/77zlLWdojEmA0ioqQOOMMRBs3ri413osgUZEDb1rESAyCIxMY9EjdlQGaM6194WITq3MQ62DVqCMCCiIFjVWjuL16qng+qlH1AohfHjLHIDCctKBt4Yz7tkFNVlRwUVWj/fFf6xVzwcM+dVHx9QvEfL5IaY7/nrOUpW8llqQCWCqrVyPaKMMIV6kRNCoPDfCjrBrjgEE1ywHD2RcPhKtSSZgm8n6GfXFsAN68KqNc3cBnWDXKSYtDAsZ2yMtRdcujYU8bhOelVAnb4CLVKGv8e+9GEOtUZX9oKspTDFIGUJDfR1UrAP/233eWsvStJCKtVaUaOLAKHFzFgwFlboKA8Q2NRnexwKBRvCKC+LUhw7g0ih/OcOPSvhhk00SoRt0SgrfSSpOiHHhjrQvXqVGUSKzxYJpIuZEE3vNhIy34GwelGL1VTlowwZEer7AFFoFxerAqhrIjjET0aIjk0jBfLCJ/8N97lrK0WxJgjKpU3boh2qRJlaeoY1BXb3qTVFVxwKCqsyiC2HEBIxEY7X7T5lkRHCtyED/fumEPCyB+t1WjCIzNOfCGh4365AYlzj9tgRSPTRXpAgs6GuV09fYNwptvYxsEJiXS6TYdnsBoElAwqkQhOG2WsBdrjHJ5ltjPunWzVHsC8CFWrkK0YcNqKlV8SBiJeRGjxzeyxnFlBjAImJODVfEJjpURxNpnqFEL7G82WTDW3XqyXOziUEqDsImm0+h+06phrjCK0osqEoueaKPQKeBsC4KPtsrpQWP88xw8DoHBvymhVjZF8YIH1acrInTr/sZ//1nKUo0kIocKEK1Y8HGVIW7B4b0VQfw8o07F7w8ZdScTGKznIDCsqzY+LITY8QHEDmQ03ALDBg0VGN0y7AunRk23rtxvbOCPKpJl+vi5QWOjUOJQleK5L7GGP8HDakGqcuvteVzWrkc7JIx4vyBiRzVH/ORjgYoKeq3ey3qtsrTHJCLtVee4/urqxjhXdrpNqfYcHFDDmN4jjVpnSgyb/xR/zLppydS3huCtygj6fZWDOIHB33k8buPUKBe7oBpmjXqW1HqjrIuWv+dZicKIvJMo/7BuXBt7ccFJgo2gUnuG0uv4ZkguO8EFBO/333+WsrRLAtBWfZ3nn18dHBwuya+nWcEVGFRjnMTgap6Xg/gjYcO0rOS73cZEXP7Vc2HEWSrL3zkOD6ihrd4oqlFvhBHvbCUKQXNeUAOJCiqO1RHE/2WlBe0PunnPCCLxVUaNCM/zsc3q5TW7WpHTgvCWhJAsOA4iMYJjA1Nk/HOQpSzVSGlwXHjhzuDg4Ep/SlDTPJS5ncSgirU8gthZ1iimt4rAICio99NbNTrDI8XBFA+qUc4b9WLYNF5g0DASQHyklRYEBY3uGdbecd4qRtyZ37XBdi2h1KJUui+sLmCVOrxGxlIeCqUTGL38Zqgo13iHx0wA/xxkKUs10reCg4Mrts1x8p4Nq1HOlVrVLq7oBMYdNQCD+9mVXBMLKS2oJm3NMaqPc/OeH9LjORWN6hZryRUYBBX3p9H9bpV9o/EVSqxWGfUd3JZSZYmVasz6HRpG7JoGqMgfw9uMZcGRpT2mPQJHBkjiZwU1eq3/t/XdaWBQxWFMxAGDq/3Btn6D9gW9Rx9E9BhqtOcEEB9hpYU1uikt9HeCzjZmoBs3sc5uQ4ky39oVTqI4FepNm+7O7dZZgDUIIJpTHxWe1m5lwZGlPadagYPD1W7QYM9UpWhnbMyBx/pwK2W0IImpJzSaaV+wqQKP8feAxjG0/Y5z89LoZtDPlchapveesYFFlyr/aMgkLlLq2HOox4pp7E6q0J1MVZC/c7uBDVFRlpUcWaol1RocblCK3GG7hZC5v8pw1ZKxTwoaYDCot8X2pqIkYP05KwW3ViUsam3HVbYoypbJshbEm2aBQWkxL6Lql25jXbkqsZwrmeocVSkXZ3Hq3n9C8OK0OcZlbY4s1Y52663a1SCTtw5WrdYbMoBBhrw2ZH4jMGhfMPXkL7QtgppPlXa/kumnZhjdTo1iKe56u38sB95ISgsbW2EZbrcgvNdtRSGPw8EcK/bhdV0TqbJRqnlhJOOtAInSW7WRHVX8c5ClLNVI6TjHv/5VO3D8I2ii1oxY3x02Ll66YgkMu4ozpSTObFxt6GYaIaSlBWMbLKElMJxH67CgMfhdfGRDjtoraRWK0XQ2d6Bhz+PYzooKwEdtnIW2DK+NwUNNsw8hGUvHOe7133+WsrRLKisrO1KAwspPP60eIf+2QbWKcQWmg5Apjwzoyq2GN7NpxxnAUApolxFKC4KGwKD79YqQAZTLnTqLdeGW6Wm4z44g1qoqZqHSYrlj+KpjeQsjJh3e2hfx60ImYMnjaH5VCMn4cQ4c//Xff5aytEuyb2H6CCtWINqgwc4g2N1wSX+s6CMAYFd7Fhu1sTXnbBdKRrUteLzJbMpgG79xf9Z9MzWFapRt5eOxWwkNc01DCRqj3HZKrNZZkdIiaNPoj7Vp6i5abrfzEs1QtO0WB45sy9As1Y4EeBvRKLx27RBt2nRnEOxu1A9o6auqUazU6x5E/ELbvoeqkX0FgbpW7wyra1cBRWnR3daFExT0Zs1h7UbIBAXZBui/Ns3d5U0RIATPJxF13aaj7gwuugRHtx1tEGbu/rMxyia8xryqEqCirf/es5Sl3ZKIdFS7o3//Pbc73GDUmjGHXqYhm/e6VX24elsvkjcxonaAqj5UwY4KwLs3BI/SgkmHLIpitJxR+LpWhZpgO524dqPOTnnIxlh4rFAA8TttcJHbuU82eWMtSeNmiHc6GigtJTg+y2bkZqnWxOZnAqxL0e7ITFvf00EViarNeUETAaen6mvjiYqfEzJxD9oOzQOIX2kbtnGVp23CgN5lNijIKj9KIZddS1BowwQjLeKMXfBctFWobrG+wzWntkFGZuhSImnS4cENUTj3euI+JYIB/vvOUpb2iABMRDwOr02b2qtWHATAMaY+PN47ZHKdmDdFSUFpQIP7beuJIoBom7CpAjNtqWIxHkGDmxKC4HB5UUtstWFG/bj2pso0up09M96+ruBA05w6/mETVOINgiPBt1H57zlLWdojEpGTBKgsGTOmeveR2gwXGXf9puiJYkIh+1LR7ctCKfa8JShshq/2o2KxE1M+KAUICqpTuTlqiGurHkolSovWQXhjbODRGd3M41obMQVQdAmz9uQaupODSBbQS11KQ5zvLfyd/56zlKU9IurjAKYimdx76ZE5XA06g3U0or+KID44bHpaUeUJ2ixc24w6bVfQw/SMjY/Q2LaZtkxbV2nhGlg7afF62HRuZ307s3H5As7tLI9tivKSFwmMMgAn+u83S1mqFVF6wEkPvovDz/C1GQQHi6TuD5nVnz11ufrTiGYi4jKrQrnYB+2KhRGNU6Rzq2innGhdtFTHnHFOUDDl5I5QusZDXblsBqGltC30BZ58Qxtf7Om/zyxlaa9IgDcoPRIdO3536cHh8qWobjHtPLN0NhMU19uGcC4bl5LlmpA2aVNXL9UtCw56xJxtoQFCvvGJrmNVtSiNDkZF+TwCo1BETvDfY5aytEcEYB8R+aMdPUXkaVqwlcuWIcaIeYsWOzN8bQeZnb2uttgGDbQvqPr4QeGCimzY5tI/yPBUoVjtR+Ocnihu56QFazwyA4R5zVGUf5MN+slg//1mKUvVCMBv2bNJB9BFgGt1CO4BsB7AVjuqqLISxcOGmbjHdwUIjXNG0Z8IwRsZNv1y2YSNHi0/KFg7TuPdpakzTrIyovEM3d4lH/IdHmxs7dJJFBhUp06l55ZS4z0R+ZN/LrL0CyV21uAgCACcL8C/7FgCYDNHqrAihcoMEKRmobL8FaQqxqOkYCBKi1VwpKl46FBEm9Qi52pXw3myXBKhq093oHDtdyosMFhGy/pyBhitu1cNc9aBMBWeksLZIDr4yufDkar4hJddkDXCf8FkpcFRAlxpx2MCbHQjzd1bo0g98xoqJ0xA5fjxKBjcE4mH2iC5uj3yE+1VR/fyGsPLawIvrxGS0eORqiSOLKXKkOzbA9H6tcy7+rZBcBAUz1X3WCkgWMHH15sRUEwlIZjYbWS8LaW1jd90qMQIIhE9HJUV8xQYInK6f76y9AshAM8JZL1UVnKFNLRhIypfGY/Um2+i7O774f3jRHht28K7+VjELmuA+JRG8BY0hrctgEQ+mSlomGpHJGP1JaM1RWHigvRhGfvApk3wOnf8flQsN6ga9Qhq4I4A0Z65fGcgjXca2jZjVyXFMBsnYVBwahhxdmtnnOQ8lt/yHg5TYJj8KXT3z1eWfkHk2Lboin7wWhkQxI84AtE6ByFapz5ixzaC904LeIsDmrKt3h4ati7IlgmGnUYEibwQSou1rJRG7fsCrEZRIYqHDPl+AeJejsm4h00v0WKni0Mm9ZyuWddbl7XlTJF3qSgNA4i2aoTEpsORSqlnqlhELvXPVZZ+YZRKpT5FKgXk5qLs5VcRa9YCMed2JfOwNxQbMrss1Z0A8G2DqlZzlBY9S6YrEcFYAVYSLCU00nmu78PNy0EwXBTSVyxrzhWzcF2VIPOpHrLVfUxNd7XrzZsjWrcektf3RSq1yLlsu/nnKUu/QALwZxE5TIBnBahIrVuP+FFHI9bUNoGmPt8yAI/vt6hmtNZmUF05BpWVXxATZSJgN+bxKk1WrIDXqZNJUvyuUoRZvVeENO+KDRiYjasSgm1/nDfLls9yez3noYei6L77VOcTQTJrY2SpRhIRGgglZa+/iViQb161zMqWObeEqlIz9mpQpz8KZSXPkA8rRDBDO7IDual4HMXDhiogvzNImItl68/TuVkuP8tuwypFqnTJs89G5dy5zh6aKCLN/HOSpSylSYB7ySxem5Oq1B2Cg+nh1Nl3YvpajLwgvLwWKCtmg3KNHywV4AE11AFULFuKgquvTjPw96ZucVB9atjQgKJnT1TMMO/4A1Ra9GQA0z8XWcpSNRLgGgXH8W2qg4PZqcw1ch06/GOP7REa6XxDK9+StMAxaHXauBGlzzyjzgEys670lCi1BIvbj8eIH3008i+/HBXvsTm6SoqkAMNEJOCfg5+CGEuhd8wNEdnfv02WfmISQCt6qoGDKglfY9wlqI2adxosLWXDZpf6vSejoAW8GUEkup6BROeuKB7yKIBvyLRvCfCmcnBREfDllygZOgj5l14Kr2OHKpA0bWqAQ+Z3w0kbjsaNkezbF/mXXYbKmTOBvDwHvYQAw38uL58BsB+Atyg404uDodm/1LoREekE4AMAiwEZKyKH+7f5SUiA/nwy3gkZapUDiNPf/aNuQN2jmn7hB0FNg3UTMyPa+CDa0EqFSHNg0xKCY6Behwg7mdwsQFr/QUkxsGaNAgYrV6JgwA0Kmvx+/ZB/8cUoGTMI+OIL8zs/LQngCTDF5nuF/Pf8U5EFxoelFcDod1agbf/xaHfNeFxwzzQUlgov/W3/Pj8WUXKJSFc7fpRu8gCai8gs3ni8oAKzluS6R5jPtk/+7X9UYh2GAINl02bEWh6y52oMX07DNIw9AQeBwWYF7AhCYHH/Zs0QP/wIYNsOMvId/uvS6L1IGwCvA3hXzFhTxftpKhFgOpmK2wC4XURa/VykhJ9ERKX0mLeX4zdtnsBfOw7V8Zs2g/DxcmWMdSJSx7/fD0EicqCIdBZgkgALBPjKTSrn2n73qm7zAzSzA6QZ7DMd9sZSRM55Fvu0G4wTr34Vn6zaxhgZAdLRv9+PRkyuo4Gaevc9JHv0QiwQSuvsGrDbFVgIDrbTsa8S2wkQNUgMdam6/Wkkn9XTsfl6AeaKCBMZG/uv0ZG91u5Uw9LQALje0tpv7t/+50bMXAaw+vOv4zjwtBH4387D8JcOQ/H3U4fjLx2H4pL7iW0NnJ7l3/f7JhFp6WJPSHrAJ58A8+ci//J+KLi2P7DwE2DhQmg8zNCM7xMgfF6ArNmeKEP3f0/Sufhbp2E4oOsI/OmUIWjeawwSxcocr/v3/dEIwK+lsnJAmtXWfAksXYrCq/sj/7zzkOjUEbFmzXcGB1Uu2h7sSbsr28MBg/GGTGA0bYr4IYch9c02M/mLFwOff+6u4Bn2wvJfJ0lErhJglSQSKJ8zB+Xz56NigTHwRYQeqJ+ltHAkIu1oZ4ycuBT7tB2MFr3HYNiEJWh53nPKEG2vec3dyxn+fb9Pol0jwCe07/Kvugreia21tDlzUdS/69fXeFTB9ddnvvbtb/7j7Q2JyFO819tGfawSdP8uIxQYHPybQPngsy3chIbjT7vwiQjb+rWlN8caRkUoL1cXaMy9IdY/nO0xrQb1ygGjpQ8YjGc0a4aSZ59TRij/bLFGzFk5WP722/yKaRwH+66tjojcqjvkJ5HofiaiBx1kHmCdOiagZ5jqqMz9fm7EJhSJwnIcc9k4/KHtYJz5byMAe93+Dn530iD0ueMd/f8Pmd9l5/ITFBYi2aOHzmOsaTPj9nbAoMbQ3C6ITZroNrTvLEC+80pOG5Bv5n577nqVnA4YVC//3GEI9u8yXBeLyx6ktqzP9TD/MX4ysu8VLyx96ikzeX5QZA6+JZZvW8oEx66AQVdrvfpInt3LMQGK7rnHNGRo3BiJ004zkwGNpqv0EJEmAizl92UvvwzvpDbVemOxyyKPYelo/738nAjA0o+Xb8Mf2w3GnzsMxYAhs7BuSwJNzh6NP7R9ElPmsjTmh02XT4lM4kmKhw/XhUXnkM6Rli1ReNttkIULUXDddYiF+Y6UZul53lG3LspeU8lWJCLH+o9bG4JRg3HKtePx1w5DFRhULTvd8DomffwVDrvwBZ2PKx+xLniRQ/3H+MkIwPvIL4D3j+N2bXPsChxM8nuvZmBwFUp0PQ3Ynscb3s7sjaK77zb15i1aqFs2Q3qEROSfmj6fTKDg+msQa94CscYZUqx5c8SPPBJiXLYLfiwPy94Qy20ByX/h3VW6KtLG+Hz9Dkybt17BUu+MUVixPs77+KyGfZuJyPEZ40j/NntCAK4QoLz09deNBOecN2yExFlnIrV4kVtgvuE/5dOnIRZhzzAjQfh+ea/DKWACKUuh/cfeUxKRsEqNOetVUlBq8LPjtRNQUWHsm3/0ewm/P/lnCA4R6ctFuujWW/as920mOBww2LLTDwyK7foNkZr9MSd3m4gcI8D6FMtpg+ykTsa30oMdBkU2axR91SokzzRqlD/FRKXGQw+5Cezsv5faEFUyAA+wgwpBSZerf5vvQjYij3b9X1Ng/LXTMMxdvgWn3fimqlidrqdTTu/jBrcP7UGIMBi03XGuJWYXjK+NV0tEruR+ZW+8YSRCs2YqsQuuvVarNsmwAC62QOwvQAElRWabJc43TOrNMv/x95QAuYMH6HbzRFWpDug6HPudOhzzPjeFpKmU4NjLDDj6PWg8+j8LcAA4SICvZe48XVHSkWoXdW7QwIyGGZV9DhwVuwdGrGEjlI6ltqQ3ey3Px8Ac9VhGxVVCcSVr1AhlbxldvPSFF8z3Ndk8dAUfeSRSW9PVue8DeEdEnqhNeSvVNxF50h1kw/Zi9+dKEenntiMjUm1zozaMSaIHigftcN0EBQe9VTcO/VDVCUqSse+oQ4LZwcdxewUGoEbrR0u24py7puCcO6eg53/fxvNTjZNJRBbymfnP5SeNWzB3J5nUrAENpjZokAYGgcBtqu0DzEFJiS5WlBoKjnr1ILMZp1T3ayRz+z0h075JhvMAwd5jFRS0Mc69awoqKtNeMZUcdFgMHG1z4IwWwepUzj1LTWmcZY5eNcWyVCUXYd2zf3s3Ju5xZgKZVp/QHQOR6NwZiW5naDS7oOsZSD3/DDB7Nio/+ADFjw9CtJEFCMExO6KvIKsRGFztmzZBydgxfAgiwC3ufPbFNxXFDzxQ1cqncRMkz+mDgv7/qgKWHxhOajz4oE5ecWkFVm5MIDfOFlM6ma0yzkGXIZmfiHtLRLpk/PY7ERnCfeYv36qrWd1uo3DFwzNQUKyrKZ/YFYDGJtIxAEurSkpK9jhpkczHnRw4nJ5N4/OgbiOxagObKuJTbmtLlUfyi3+P/Bj7th+CP3UYooxEIP35lCE4dyD7zJHkEf+5MintmSooQLJXL5XOXOjyL7rAGdi0IU6tYb/Z2LwZ8aOOMs+AzzEYBD771N1/UUWF9PDvtzuyAdDExI/Xpe/9Lx2GYPB4tiE29Omqb1D/jFEKnE07dKGaKSI9XN5RbqwUX2wpwNqthfq5PVFuZsF4KzWqznsGQObYxt+8wlR6Hze+3FqQuRAu/tZFRkSY/LQTsUJOgGUUp4w8Vyziuzjsas7+tQ/YJs5+YDgDvIcxwEXkMd/5/q4ZukuWGFHvgECRn1kUxe8zjMNMqVFUWoE+AyfjV8c9hmsG0cmm59HV17pO+R4/nZBv4mX8jYGlY/g7eZW/rVwfRaDXWDUCybj8pJqTKDBgIy1fH8cVD7+nkex+D81AboxFg/i8JinF7+hhIW+DrYrMmM8dOl5fBQ4OqhadbtCWo7y2G7l/RUVFF/7/kZcW6bVkujcVVF2G4387DcMs4+rk/dUYG9IgH89Lz1SvXkZNpiQ+5BDIyuUOGNUkht2PjcGj5VOmpFVrag8VvXtg8VcxjUssWavTuiZjn3oMwGbc7xuWn37ttuEr4xgIf+zVxdin7ZN6Hw27P4VvYkXpeX7i1U81EFj39JH4+hstVP2agNoaK8HNw2ej8dlP46+dhmK/zsPUVmnSYzTuf0Hr/jl/l1vpoqkSDKpedN80HHrB87ot93GDc8lzXP7QDCSLuBDKYhFpkJ4EP1kVgyscuTk9HLORBLgMZaVInN41LW4VFBnp4Zmre6LjqUCeRsHJSA2rn1HPeRtvpHDgwBpffsNzxFu1QvKcvuk+WMbWMFLj0gem439OeBx9B05GYUklJ4j+Pz6EtoB481duQ9cb39QJ71il12uQjQGmvGQ5Dj7vOfzmpEH6edwVLysD7tt+MLrdMhELVubikvunp4NTHDzfvc+m4ytXZtwL6/LPB6DWLUNY67cV6Qq1LrcAqzfEcOLVr+jDdeDg8Ua/vYKbk1FV4gGYtCNZrhFjrqAcXGWpi/Mhcz9KEUo6S2lpnEkpEdWBigcNQrRuXQOMli1RMWfOboHBBE3GnxJdu6SfcXGThlhw8/1ofM7zOLbfS8iN6ar7hd2Hke4N/GJjXonec6lpwkFPSTrzmWAloz85YYk6IRw4cqM0d4BkURkOPvc5tTfOum0SSsvMQd7/dDPqn/mUzhXvn4xNacpFhs+1aY/R2Jynx/jMpaIwfuKeF+dPF6KOQ/HH9oN1GBfycD1X15vcPMqIapNRWxIR9bcmu/eokZnTTN2wIRKdTwV2RPkgqCTXuLolEon91CNVXg6vdetqx2RTBor1shkz4J1k876s1JDt27B47Q696Ybdn8bmPF19ltNjxdUqJfJZSqC5S3842UiEzlUrtFb9UefcGitVT9FNQz7E5u358PJLcdh5z+mEMmpL8W9871VBKq7ax13+Euyz095WVBlERI0lt8oRjHyQfAgcPJ47hpMG+3cdjlUbk9xtjj0O1cBtHy3dqte836nDEOwzFjMXbURRSTmueGiGHocjcg7jpUo3++eVqygdGmUTJxqHB+fu8MMVGHYOdorCW4mRoPJbMGBA2mWe36QR5hx6Apqf+iT26TAM80yai9hUGC4Iz/OLh15cqPdLaffMFLWLKKXTxrSIXE1VddBri9Pg4Oqd5xWrBkCJRGZudNbTmPmp6fEx/ZMN+nw0m6DjUBxx0Qs4/oqXdUGhDXbkxS/i8ofThvvH/Hh43EK1Wfa388xn2KLXWAx7fQneXfC1jhsGz1IJ4iTypI/X6TGqz0gtyQYKSwou6bdLcKjbr2MHYPt2lRi7AoYj+vUFiPLtUAQDVzl6pwqu7IfK1atRPGpU2mtCqVFspQZF5q9PfCIz7aK3lX4v8v+3jpitjO1WjVMHmMRfZoDyvCIygv9fuyVfvyfl7ihAsOdYnVDHxJxEPjT3fx6r13/fhrUhh9hjqRdm8tz1qNdtFPa1qxwZgKsUJVGwzzM+cAzXB//FZgUHFW9ee0v+58J7pqVXvnufq0rz37qjEAd1G6UP9JDzTSDVD45qwLAqKee0+LHHdGMR4USkV3Rre52hwCgvM8BwXqoGDZB/VnecfuXz+E27oXj8ZbU5CIzb7b5Mu8dTk5bpak5mI+M/N3U1v85lykzGdd3PLx8Zt1C3ob1x4b3TkF9UppKCc/S7k5/E7aMNgEmtr3xVnyHHmbe+hcLStOGueq0jC4x538RKEabEtQsbj0lgca6PueRF3DDkQx2XPThDgcnn8Ns2g3D/80Y1c9e61wRgIebPN8EjPzicy3aW6v/0ghzi378mssZ5DF99pR1QUlOnOhcjvBNOMCsf4xpHHw3kbcMna/LU40PmuWesqvOcoFMAjOXfn6zMVZHtmJyqzEn/fA1l5XrMl1jTwVVNRFjLslZEiLj3o8lytOgxJr0fGfvwC1/Awy98ooYiJ5li+cFxqjmVMn7BnCPGB5gsdyAli5UQfEC8xvPunooJH3yBpyctr2ZvcOx7ymC8M8esWkywhMir/OP8u6fqve3TfjCmzqfabWhbvCgNjnDftOS4NWMemXtWKd9sQexQk0jKRazgmmuAYlWFWKqsHeMtKDoJoKpIauFCJLp3S0sMAqSkdw9Mnb4Uf+k6ShmNQUxna4hIDjXHdVsTOPyiF1Si8p7I+C9NV7V/O1Upd202dw6PvrTIBEJPGYIbhszSe+Wc/rXjMJOdXGIy+Ret3oYG3Z9SEFEDSBSU8muKifNtvOdcDqv+09ifvnpjvpG4nYeplBnx5lK9Nh6f1081ioPSjXPI6wj1fQZL1mrszXPXutekOvWChTv3oiIwGjRC6VOm8wgNUv++uyMrQVYxM5WGGB9y+axZiIVM+S4fVuHNN+uBR01chn3aPanMS/EaS+pCkorll+vq1qL3WKvSGGY1npGhGPcuD6/EJYg6cToaDWB0okjS4HAM+Pm6Har/Nutpvs8AB3W5iIgM5H9uGfKRTjYZhBKDD5YRX0cdrp2Qthkc+Ki+caV7aNxC5Mb14Std+8QHuhJze11dC8uQLCxVg3W/zsOxb7vB+Nfj9F6raaN2j/XSLFTPVJ8+ph0rF5RWrUA3rta1iPwDQFMBHnWgQCKhUiXGDjS2ZJnew2T37oAXQ4+7jBT7e6dhWLBcY4Q7OFd8Rhu2FeCoS8ZVdzJ0HIpzBk5xt3JXxvzSYEffOyZX294lHTY+62lnnNOAKHnt/S/w+7ZPKogytIN03pmI1OUAhCUXExjH+XKTl1bBCCi9vYJSvDRjNS66912Nvh996Tgdh134PP756Eys2qDOBR47HWPaayJCsWSp6VqS4VHKBIYAt/n32xOyrzL4rTWq88pGjcKOAw4wD6xpU5TbUteL7pyCP9sJJoMxX2nA0I/UqOZkqiHbebiu3NRLudoSIGTaPndMVs/W3BX6oMkwmg4B4FkHjr91Nt6MpWtNwRTB58BBlenGYR+hsETdiFzmUlPmrcdBp4/U4/OayExX2Qjvxm35uPqxmTboZWwNpoyQIXh9vFa32vG66JpeszGuKhe35X4nX/UqTrjqFd2OD5/u37VblOEp9dinuJ4A89Qz1bdvlZepQQMU3nQTpFANVgKJ12uW5u3bUfz44yqNdXuqYHbkX3YxZMcOtatOHfCGnpfX8fL01diSV4hoQQor1sdwzKXj9Pt0lrG1kzSr1nj7yO2M1DK+sTheUIrGZ42upraqTdB7LN5fvMld46UAlrz36RY93p/aD8HF95s8KxsAZQCEK0P6XuhYGTJhicZMqILxmPRsfbRUPXppYjylslJ0ZMRW6Aa+hLEUPz/WmtitQ9t29O6dTlJj0LD0KS4misD/+PepLbFDiuYZrVkDr317fXjx41oBMYPyF99did+eNKja6sOV1nh2zOpF3Z8FRZUpwTkDJ6soVQnC+okTn8Cx/ca5673InvN5r0hwUNeReuzbRlH3FRa+F8XzS9HgzCo1jSA45doJuGXEbNw49CM1IjMNbl4HQdnz9nfSjOMYgdcwbf7XaH/NeBX1ztjnJ3XuaQuMGkXJ8LuTBxmj0a6uzuC9/em0YX01r10Azb4suPWmnfPhWrRA4vTT1SNYdOedKL7zTuRfdJEBD5+fdZNrsLdJE5S98opjGl11jYfNXD8XDIKWo87pI9P3xWuiF4tudV4jJV7vO95RG4lUWAps3F6Ibre8pb+rLWYXBB6jzT9fxVdbPHdahg128F1zHa57XefgqEtexIChH6qE3RItwbZ4KR4Zt0gXqYvve1evgZoEvVd81m6umvYYo/bE5u0FyPOqXMY+oufuJD8P7hUxkMYjJnv0NhHzRo1Q8rQFxl5KDD9RH66srGRVoM5Y5VsTgR3bIKlUOQM+G74pwAlXvqIJfHwQzoNDlYeTwtVu1deaq6SXtd0rQZcBb+BvnUwcgxNJ7wotS5cFC2AoNx43fQ1eff9Ls6PIKjpAKunDfnCGTr47F6WD6rBWf3VuxkyAkAEc8ztV7ZnJ6roFr4lBQdV9reHIzw/MCopkcQpn/WeSHofXzHulOvD0W8vdfTFQ+HteuzDoVVmJ+BGH1xw4ZXaty3DgyOyAT4O9Th0kundHapbRtFityC7z/OOCe6bqPToAO+8bB205XvM1T3yAdbn52JEosW7XwXrvlAjXDfoA4T7PKLB4D5y3A08bifPumqoxCKqiPA73u/rRmXh3wUZ4Nsa0bmsSra98ReMiXLB4Lu5b57SRut+vWj+ug5K4/bUTMHneBhQUV6D3f9/WeePiw32oOdC2oJOGgVUO/n3H6LmIJVWd/SjTebDXxPJFESnPv+AS9YKUPG2aR38fEsNP1JEFoPeJ/kOWvfLch9jG1njtgy9V3WBMIOfcZ3HX2PmYtiDd6pf7XFxZWXkVIIqUKfM2aNXZ6x+uddfMULO+JZZpBPp2KrMfPS3jysvLjzHfi1dSJho0otep5fnPVRu0TZh6zhWQD58Mnw46WZAQPM8aFyfPS/GgKJg872t9SFTbHnvFRKBFhKvNh/x71tJcjJi4PA1Y5r7RJ59Z90IbgrGJ+PGtqwdMdzdsfhU9hCWPPgrk53Nx4zxRh2f6ygEisik3WoS2/V+ziwozioekP8m4r76fLk/OS6VSFbOX5WqaC++dqz6Zl4sIGTXQe6w6KBasUsMen3/t6ZySebkNvY8ct47UEA1Ji4dfmfkFTu7/mj5nqs6Rc5/F0Re9gEdeXIRRb63AZhNNJ/HZqXrx3qJNGqvivBJcPP6vWj+GXx1vR6tHVRVespYmlNo5371WxZbWLpI5c1D8pElLoifCv933SQD2rcYMJgJNz5SR29WJBv3dmT2pSkpKmgOgL5NPhXoqffDLmCWaPok5D5mC59o383sbHK2uwO6CKHlaXf6yAoaAJZi63Pgmps7XOBnnaj4DoiUlJU0BcAINp1gSEcZr6lsPzI1c1ew2HM+LyMmZ16bXB5zK/KnSkSNNKnpN0sPW0qj6VK+eAQXdu1ZVZTUmjfVqxzWxjATNK+Z/jZy4XLNlh76+FONnpZ0NVEsoxQgmutP1gDMXbcawN5dh1KQVGPnWCjw7dRWMmaa0BVDvlR5k2boYhr2xDP0f/wD3PDsfXqGaEhXGxazHjqb3rE40vDgvdOWyBxufXWNr5+iC6BWl8NSkFQoiXosbQ99YhqVfmcPa4qvfZt77XlPa02HeQ7w4Go1+byWUtaGysrIjAC2EMkOkz+7Eo03tYLCLXqZaGWAm3UWZpep81YZeBw2ZtPLsozwRuc7/8k17TXRJ3lqTYWgDbdxmp1QVR7oNK/XKykw9Bh0YLAhjhR8H7ZBgUAHBCs+yMWMyQUEDhu7QGl8KygpLxnOsd0+9ANbQpsrJa65WZKYRc+O2pceD+2SOSYDc4JL9bPdNzgkXNF3CLS2rFOntjllYWHgQzbCd5xxMSmT59E6MbUCiz+TlGq7DjQ1M2/HP+XciZS7gSQEGiUgL/++/ZKJUAISlx//JGBf80PUmzCUTK0lZPszG3cUjRpgxZAjK2a7I1oTbHl67BYWfeP181vZ+jtjdIkSyjM990mNXJdB2xf8btE2UJhp+rw33/NeRMXa54GTp/zNiOMV6rqrC/lVEZDzBbi9M53bGfJaq6P8A55+BvfC5f/QAAAAASUVORK5CYII="
BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = None  # Evidencias almacenadas en Supabase Storage.

def preparar_sqlite():
    return None
# Seguridad: el administrador inicial solo se crea cuando se proporcionan variables de entorno.
# Nunca se restablece una contraseña en cada arranque.
FORCE_RESET_ADMIN = os.getenv("CALIDAD_FORCE_RESET_ADMIN", "0").strip() == "1"
ADMIN_USER = os.getenv("CALIDAD_ADMIN_USER", "admin").strip() or "admin"
ADMIN_PASS = os.getenv("CALIDAD_ADMIN_PASS", "").strip()
# SEED_PRODUCTS se administra en Supabase; ver migrate_seed_data.py.
# SEED_DEFECTS se administra en Supabase; ver migrate_seed_data.py.
# SEED_NAVES_LINEAS se administra en Supabase; ver migrate_seed_data.py.
# SEED_FORMATOS_ENTREGA se administra en Supabase; ver migrate_seed_data.py.
# SEED_CATALOGS se administra en Supabase; ver migrate_seed_data.py.

st.set_page_config(page_title=APP_NAME, page_icon="✅", layout="wide", initial_sidebar_state="collapsed")


def now_iso(): return datetime.now().isoformat(timespec="seconds")
def hash_password(p):
    salt=os.urandom(16); key=hashlib.pbkdf2_hmac("sha256",p.encode(),salt,120000); return base64.b64encode(salt+key).decode()
def check_password(p, stored):
    try:
        raw=base64.b64decode(str(stored).encode()); return hashlib.pbkdf2_hmac("sha256",p.encode(),raw[:16],120000)==raw[16:]
    except Exception: return False

# Seguridad de entradas: bloquea URLs, dominios y etiquetas HTML/script en campos capturados.
_PATRON_ENTRADA_PELIGROSA = re.compile(
    r"(?is)(?:https?\s*:\s*/\s*/|www\s*\.|<\s*/?\s*[a-z][^>]*>|(?:javascript|data|vbscript)\s*:|"
    r"\b(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+(?:com|net|org|mx|io|co|gov|edu|info|biz|app|dev|tech|online|site|xyz|me|tv|cloud)\b)"
)


def detectar_entrada_peligrosa(valor):
    if valor is None or not isinstance(valor,str):
        return None
    texto=valor.strip()
    if not texto:
        return None
    coincidencia=_PATRON_ENTRADA_PELIGROSA.search(texto)
    return coincidencia.group(0) if coincidencia else None


def entrada_segura(valor, campo='Campo'):
    """Retorna (ok, texto). No transforma silenciosamente contenidos peligrosos."""
    texto='' if valor is None else str(valor)
    return (False,texto) if detectar_entrada_peligrosa(texto) else (True,texto)


def _instalar_filtro_entradas_streamlit():
    """Aplica el filtro a entradas libres sin cambiar selectores, fechas ni contraseñas."""
    if getattr(st,'_filtro_seguridad_instalado',False):
        return
    original_text_input=st.text_input
    original_text_area=st.text_area
    original_data_editor=st.data_editor

    def text_input_seguro(label,*args,**kwargs):
        valor=original_text_input(label,*args,**kwargs)
        if kwargs.get('type')=='password':
            return valor
        if detectar_entrada_peligrosa(valor):
            st.error(f'{label}: no se permiten enlaces, dominios ni etiquetas HTML/script.')
            return ''
        return valor

    def text_area_segura(label,*args,**kwargs):
        valor=original_text_area(label,*args,**kwargs)
        if detectar_entrada_peligrosa(valor):
            st.error(f'{label}: no se permiten enlaces, dominios ni etiquetas HTML/script.')
            return ''
        return valor

    def data_editor_seguro(data,*args,**kwargs):
        resultado=original_data_editor(data,*args,**kwargs)
        if isinstance(resultado,pd.DataFrame):
            peligros=[]
            for columna in resultado.columns:
                for indice,valor in resultado[columna].items():
                    if isinstance(valor,str) and detectar_entrada_peligrosa(valor):
                        peligros.append(f'{columna}, fila {indice+1 if isinstance(indice,int) else indice}')
                        resultado.at[indice,columna]=''
            if peligros:
                st.error('Se bloquearon enlaces o etiquetas HTML/script en: '+', '.join(peligros[:8])+'.')
        return resultado

    st.text_input=text_input_seguro
    st.text_area=text_area_segura
    st.data_editor=data_editor_seguro
    st._filtro_seguridad_instalado=True


_instalar_filtro_entradas_streamlit()



def opt_blank(values):
    return [''] + [str(v) for v in values if str(v).strip()]

def idx_or_zero(options, value):
    value='' if value is None else str(value)
    return options.index(value) if value in options else 0

def normalizar_catalogo(v):
    import unicodedata
    t=unicodedata.normalize('NFKD',str(v or '').strip().upper())
    return ' '.join(''.join(c for c in t if not unicodedata.combining(c)).split())

def nave_por_catalogo(linea,grupo=''):
    d=read_df('SELECT nave,linea_norm,sector_norm FROM catalogo_naves_lineas WHERE activo=1')
    for v in (normalizar_catalogo(linea),normalizar_catalogo(grupo)):
        if v and not d.empty:
            x=d[(d.linea_norm==v)|(d.sector_norm==v)]
            if not x.empty:return str(x.iloc[0].nave)
    return ''

def clasificar_filas(filas):
    t={'Nave 1':0.0,'Nave 2':0.0,'Nave 3':0.0,'Sin clasificar':0.0};out=[]
    for g,l,p,h,c,o,n in filas:
        nv=nave_por_catalogo(l,g) or 'Sin clasificar';t[nv]+=float(h or 0);out.append((g,l,p,h,c,o,n,nv))
    return t,out

def guardar_matriz(eid,fecha,analista,carga,tot):
    analista_limpio=str(analista or '').strip()
    exec_sql('INSERT INTO matriz_entrega(fecha,analista,entrega_id,total_carga_datos,horas_nave1,horas_nave2,horas_nave3,actualizado_en) VALUES(?,?,?,?,?,?,?,?)',(fecha,analista_limpio,eid,carga,tot.get('Nave 1',0),tot.get('Nave 2',0),tot.get('Nave 3',0),now_iso()))

def excel_matriz(df):
    b=BytesIO()
    with pd.ExcelWriter(b,engine='openpyxl') as w:
        base=df.rename(columns={'fecha':'Fecha','analista':'Analista','total_carga_datos':'Total carga de datos','horas_nave1':'Horas Nave 1','horas_nave2':'Horas Nave 2','horas_nave3':'Horas Nave 3'})
        base.to_excel(w,sheet_name='Base indicadores',index=False)
        for c,n in [('total_carga_datos','Total carga'),('horas_nave1','Horas Nave 1'),('horas_nave2','Horas Nave 2'),('horas_nave3','Horas Nave 3')]:df.pivot_table(index='analista',columns='fecha',values=c,aggfunc='sum',fill_value=0).to_excel(w,sheet_name=n)
        for ws in w.book.worksheets:
            ws.freeze_panes='A2'
            for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='062C36')
    return b.getvalue()

def pdf_pnc(rid):
    """Genera el formato corporativo de Producto No Conforme del registro seleccionado."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase.pdfmetrics import stringWidth
        from reportlab.lib.utils import ImageReader
    except ModuleNotFoundError:
        return None
    df=read_df('SELECT * FROM pnc_registros WHERE id=?',(rid,))
    if df.empty: return None
    r=df.iloc[0].to_dict()
    def valor(campo):
        v=r.get(campo,'')
        return '' if v is None or pd.isna(v) else str(v)
    def texto(c,contenido,x,y,tam=8,negrita=False):
        c.setFont('Helvetica-Bold' if negrita else 'Helvetica',tam)
        c.drawString(x,y,str(contenido or ''))
    def ajustar(c,contenido,x,y,ancho,tam=8,negrita=False):
        contenido=str(contenido or '').replace('\n',' / ')
        fuente='Helvetica-Bold' if negrita else 'Helvetica'; actual=tam
        while actual>5.5 and stringWidth(contenido,fuente,actual)>ancho: actual-=.25
        if stringWidth(contenido,fuente,actual)>ancho:
            while contenido and stringWidth(contenido+'...',fuente,actual)>ancho: contenido=contenido[:-1]
            contenido+='...'
        c.setFont(fuente,actual); c.drawString(x,y,contenido)
    def parrafo(c,contenido,x,y,ancho,tam=8.5,interlineado=11,max_lineas=7):
        lineas=[]
        for bloque in str(contenido or '').splitlines() or ['']:
            actual=''
            for palabra in bloque.split():
                prueba=(actual+' '+palabra).strip()
                if stringWidth(prueba,'Helvetica',tam)<=ancho: actual=prueba
                else:
                    if actual: lineas.append(actual)
                    actual=palabra
            if actual: lineas.append(actual)
        c.setFont('Helvetica',tam)
        for i,linea in enumerate(lineas[:max_lineas]): c.drawString(x,y-i*interlineado,linea)
        if len(lineas)>max_lineas: c.drawRightString(x+ancho,y-(max_lineas-1)*interlineado,'...')
    b=BytesIO(); c=canvas.Canvas(b,pagesize=A4); W,H=A4
    c.setTitle('Informe de Producto No Conforme')
    c.setLineWidth(1.1); c.rect(10,10,W-20,H-20)
    izq=58; der=W-58; ancho=der-izq
    # Encabezado
    sup=742; alto=52; logo=132; revision=58
    c.setLineWidth(.7); c.rect(izq,sup-alto,ancho,alto)
    c.line(izq+logo,sup-alto,izq+logo,sup); c.line(der-revision,sup-alto,der-revision,sup)
    # Logotipo corporativo integrado directamente en el codigo mediante Base64.
    try:
        logo_bytes=BytesIO(base64.b64decode(LOGO_MUNDO_DULCE_BASE64))
        c.drawImage(ImageReader(logo_bytes),izq+8,sup-alto+7,width=logo-16,height=alto-14,preserveAspectRatio=True,anchor='c',mask='auto')
    except Exception:
        texto(c,'Mundo Dulce',izq+28,sup-31,11,True)
    centro_izq=izq+logo; centro_der=der-revision
    c.setFont('Helvetica',8.3); c.drawCentredString((centro_izq+centro_der)/2,sup-19,'Anexo 5. Informe de Producto No Conforme')
    c.drawCentredString((centro_izq+centro_der)/2,sup-39,'PG-CAL01-2301-01760-2007')
    c.drawCentredString(der-revision/2,sup-31,'Rev. 0')
    # Folio
    y=665; texto(c,'PNC No.',der-100,y+5,8.5); c.rect(der-58,y,58,18)
    fecha_pnc=valor('fecha_apertura')
    try: anio_pnc=str(pd.to_datetime(fecha_pnc).year)
    except Exception: anio_pnc=str(datetime.now().year)
    ajustar(c,f'{rid}/{anio_pnc}',der-55,y+5,52,8,True)
    # Datos registrados
    producto=' - '.join(x for x in [valor('item'),valor('descripcion_producto')] if x)
    cantidad=valor('cantidad_observada')
    try: cantidad=f"{float(cantidad):.2f} kg"
    except Exception: pass
    responsables=' / '.join(x for x in [valor('supervisor'),valor('analista')] if x)
    codigo_nc='_'.join(x.strip() for x in [valor('etapa'),valor('codigo_defecto')] if x.strip())
    filas=[('Fecha:',valor('fecha_apertura')),('Sector:',valor('linea_sector')),('Código de la No Conformidad:',codigo_nc),('Item y Producto o SE:',producto),('Lote:',valor('lote')),('Cantidad observada:',cantidad),('Responsables:',responsables)]
    arriba=646; fila_alto=19; etiqueta=150
    c.rect(izq,arriba-fila_alto*len(filas),ancho,fila_alto*len(filas)); c.line(izq+etiqueta,arriba-fila_alto*len(filas),izq+etiqueta,arriba)
    for i,(nombre,dato) in enumerate(filas):
        limite=arriba-i*fila_alto
        if i: c.line(izq,limite,der,limite)
        texto(c,nombre,izq+2,limite-fila_alto+6,7.8); ajustar(c,dato,izq+etiqueta+3,limite-fila_alto+6,ancho-etiqueta-6,7.8)
    # Acciones inmediatas
    acciones_sup=495; acciones_alto=112
    c.rect(izq,acciones_sup-acciones_alto,ancho,acciones_alto); c.line(izq,acciones_sup-18,der,acciones_sup-18)
    c.setFont('Helvetica',8.5); c.drawCentredString(izq+ancho/2,acciones_sup-13,'Acciones inmediatas')
    parrafo(c,valor('acciones_inmediatas'),izq+6,acciones_sup-32,ancho-12)
    # Responsable
    responsable_y=348; c.rect(izq,responsable_y,ancho,18)
    ajustar(c,'Responsable: '+(valor('analista') or valor('supervisor')),izq+3,responsable_y+5,ancho-6,8)
    # Observaciones
    obs_sup=331; obs_alto=115
    c.rect(izq,obs_sup-obs_alto,ancho,obs_alto); c.line(izq,obs_sup-18,der,obs_sup-18)
    c.setFont('Helvetica',8.5); c.drawCentredString(izq+ancho/2,obs_sup-13,'Observaciones')
    observaciones=valor('observaciones')
    if valor('descripcion_defecto'):
        observaciones=('Descripción de la no conformidad: '+valor('descripcion_defecto')+'\n'+observaciones).strip()
    parrafo(c,observaciones,izq+6,obs_sup-32,ancho-12)
    c.showPage(); c.save(); return b.getvalue()

def pdf_pnc_fisico(rid):
    """Genera el Registro Fisico PNC en A4, llenado unicamente hasta Disposicion."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase.pdfmetrics import stringWidth
        from reportlab.lib.utils import ImageReader
    except ModuleNotFoundError:
        return None
    df=read_df('SELECT * FROM pnc_registros WHERE id=?',(rid,))
    if df.empty: return None
    r=df.iloc[0].to_dict()
    def v(campo):
        x=r.get(campo,'')
        return '' if x is None or pd.isna(x) else str(x).strip()
    def fecha_es(x):
        try: return pd.to_datetime(x).strftime('%d/%m/%Y')
        except Exception: return str(x or '')
    def t(c,texto,x,y,tam=10,negrita=False):
        c.setFont('Helvetica-Bold' if negrita else 'Helvetica',tam)
        c.drawString(x,y,str(texto or ''))
    def fit(c,texto,x,y,ancho,tam=10,negrita=False,centrado=False):
        texto=str(texto or '').replace('\n',' / ')
        fuente='Helvetica-Bold' if negrita else 'Helvetica'; actual=tam
        while actual>5.5 and stringWidth(texto,fuente,actual)>ancho: actual-=.25
        if stringWidth(texto,fuente,actual)>ancho:
            while texto and stringWidth(texto+'...',fuente,actual)>ancho: texto=texto[:-1]
            texto+='...'
        c.setFont(fuente,actual)
        if centrado: c.drawCentredString(x+ancho/2,y,texto)
        else: c.drawString(x,y,texto)
    def logo(c,x,y,w,h):
        try:
            data=BytesIO(base64.b64decode(LOGO_MUNDO_DULCE_BASE64))
            c.drawImage(ImageReader(data),x,y,width=w,height=h,preserveAspectRatio=True,anchor='c',mask='auto')
        except Exception: fit(c,'Mundo Dulce',x,y+h/2,w,11,True,True)
    def encabezado(c,reverso=False):
        x0,x1,x2,x3=38,132,494,557; top,bot=805,756
        c.setLineWidth(.8); c.rect(x0,bot,x3-x0,top-bot)
        c.line(x1,bot,x1,top); c.line(x2,bot,x2,top)
        logo(c,x0+6,bot+5,x1-x0-12,top-bot-10)
        fit(c,'Anexo 1. Registro de Productos No Conformes',x1+3,top-17,278,10,False,True)
        fit(c,'Rev. 3',412,top-17,72,10,False,True)
        fit(c,'PG-CAL01-2301-01760-2007',x1+4,bot+7,x2-x1-8,9.5,False,True)
        fit(c,'Reverso' if reverso else 'NUMERO',x2+2,top-28,x3-x2-4,10,False,True)
    b=BytesIO(); c=canvas.Canvas(b,pagesize=A4); W,H=A4
    c.setTitle('Registro Fisico de Producto No Conforme')
    # Pagina 1: formato A4 fiel a la referencia.
    encabezado(c)
    try: anio=str(pd.to_datetime(v('fecha_apertura')).year)
    except Exception: anio=str(datetime.now().year)
    fit(c,f'{rid:03d}/{anio[-2:]}',496,765,59,10,True,True)
    t(c,'Sector:',39,733,10); c.line(79,731,404,731); fit(c,v('linea_sector'),82,734,319,10)
    # Identificacion
    x0,x3=38,557; top_id,bot_id=711,498
    c.setLineWidth(.8); c.rect(x0,bot_id,x3-x0,top_id-bot_id)
    c.rect(x0,top_id-18,x3-x0,18); c.setFont('Helvetica-Bold',10); c.drawCentredString((x0+x3)/2,top_id-13,'IDENTIFICACION')
    codigo='_'.join(x for x in [v('etapa'),v('codigo_defecto')] if x)
    cantidad=''
    if v('cantidad_observada'):
        try: cantidad=f"{float(v('cantidad_observada')):.2f} kg"
        except Exception: cantidad=v('cantidad_observada')
    # ITEM y producto comparten una fila como en el formato fisico.
    campos=[('Código:',codigo,681),('Fecha:',fecha_es(v('fecha_apertura')),648),('ITEM:',v('item'),615),('Lote:',v('lote'),582),('Cantidad observada:',cantidad,549),('Responsable / Persona que detecta:',' / '.join(x for x in [v('supervisor'),v('analista')] if x),516)]
    for etiqueta,valor,y in campos:
        tam_campo=9.6
        t(c,etiqueta,40,y,tam_campo)
        valor_x=40+stringWidth(etiqueta,'Helvetica',tam_campo)+9
        fit(c,valor,valor_x,y,x3-valor_x-6,tam_campo)
    etiqueta_producto='Producto o Semielaborado:'
    t(c,etiqueta_producto,300,615,9.6)
    producto_x=300+stringWidth(etiqueta_producto,'Helvetica',9.6)+9
    fit(c,v('descripcion_producto'),producto_x,615,x3-producto_x-6,9.2)
    # Disposicion
    top_d,bot_d=482,350; medio=300
    c.rect(x0,bot_d,x3-x0,top_d-bot_d); c.rect(x0,top_d-18,x3-x0,18)
    c.setFont('Helvetica-Bold',10); c.drawCentredString((x0+x3)/2,top_d-13,'DISPOSICION')
    c.line(medio,bot_d,medio,top_d-18)
    t(c,'Marque con una X el recuadro correspondiente:',40,451,9.3)
    t(c,'Fecha, nombre y firma de responsable de disposición:',304,451,9.1)
    opciones=['Reproceso','Retrabajo','Decomiso','Inspeccion','Aprobado en segunda instancia','Otro']
    actual=v('disposicion').lower()
    for i,op in enumerate(opciones):
        y=430-i*16
        c.rect(40,y-3,82,16); t(c,op+':',124,y+1,9.3)
        normal_op=op.lower().replace('inspeccion','inspección')
        if normal_op==actual or op.lower()==actual:
            c.setFont('Helvetica-Bold',12); c.drawCentredString(81,y, 'X')
    # Fecha arriba y nombre del analista abajo. Sin firma autogenerada.
    fit(c,fecha_es(v('fecha_apertura')),307,430,242,10,False,True)
    fit(c,v('analista'),307,365,242,10,False,True)
    # Tratamiento: estructura vacia.
    top_t,bot_t=334,216
    c.rect(x0,bot_t,x3-x0,top_t-bot_t); c.rect(x0,top_t-18,x3-x0,18)
    c.setFont('Helvetica-Bold',10); c.drawCentredString((x0+x3)/2,top_t-13,'TRATAMIENTO')
    ancho_t=x3-x0
    col1=x0+ancho_t/3
    col2=x0+(ancho_t*2/3)
    for xx in (col1,col2):
        c.line(xx,bot_t,xx,top_t-18)
    separador_y=top_t-67
    c.line(x0,separador_y,x3,separador_y)
    titulos=[
        ('Kg Aprobados en\nsegunda instancia',x0+3,top_t-31),
        ('Kg Decomiso',col1+3,top_t-31),
        ('Kg Reproceso',col2+3,top_t-31),
        ('Fecha comienzo del tratamiento',x0+3,separador_y-15),
        ('Fecha final de tratamiento',col1+3,separador_y-15),
        ('Nombre y firma del responsable del\ntratamiento:',col2+3,separador_y-15)
    ]
    for texto,x,y in titulos:
        for j,linea in enumerate(texto.split('\n')):
            t(c,linea,x,y-j*11,8.9)
    # Consideraciones: estructura vacia.
    top_c,bot_c=198,76
    c.rect(x0,bot_c,x3-x0,top_c-bot_c); c.rect(x0,top_c-18,x3-x0,18)
    c.setFont('Helvetica-Bold',10); c.drawCentredString((x0+x3)/2,top_c-13,'CONSIDERACIONES DE TRABAJO')
    filas=[('Dias de retrabajo:',top_c-34),('No. de personas:',top_c-52),('Materiales:',top_c-70),('Otro:',top_c-88)]
    for etiqueta,y in filas: t(c,etiqueta,x0+2,y,9); c.line(x0,y-5,x3,y-5)
    c.line(300,bot_c,300,top_c-93); t(c,'Fecha:',x0+2,bot_c+7,9); t(c,'Nombre y firma:',302,bot_c+7,9)
    c.showPage()
    # Pagina 2: todo el reverso usa el mismo tamaño de 10.5 pt y el ancho completo del encabezado.
    encabezado(c,True)
    TAM_REVERSO=10.5
    X_REVERSO=38
    ANCHO_REVERSO=519

    def lineas_reverso(texto, sangria=0, negrita=False, espacio_despues=0):
        nonlocal yy
        fuente='Helvetica-Bold' if negrita else 'Helvetica'
        palabras=str(texto or '').split()
        linea=''
        ancho=ANCHO_REVERSO-sangria
        for palabra in palabras:
            prueba=(linea+' '+palabra).strip()
            if stringWidth(prueba,fuente,TAM_REVERSO)<=ancho:
                linea=prueba
            else:
                c.setFont(fuente,TAM_REVERSO); c.drawString(X_REVERSO+sangria,yy,linea)
                yy-=15; linea=palabra
        if linea:
            c.setFont(fuente,TAM_REVERSO); c.drawString(X_REVERSO+sangria,yy,linea)
            yy-=15
        yy-=espacio_despues

    yy=727
    lineas_reverso('Instrucciones para el uso del Registro de Productos No Conformes:',espacio_despues=21)
    lineas_reverso('ENCABEZADO:',negrita=True,espacio_despues=3)
    lineas_reverso('En el cuadro superior derecho se coloca el número consecutivo del registro que se genera. La numeración se compone de la siguiente manera:',sangria=22)
    c.setFont('Helvetica-Bold',TAM_REVERSO); c.drawCentredString((38+557)/2,yy,'NNN / AA'); yy-=21
    lineas_reverso('dónde:',espacio_despues=0)
    lineas_reverso('- NNN es el número consecutivo',sangria=16)
    lineas_reverso('- AA son los dos últimos dígitos del año en curso.',sangria=16,espacio_despues=8)
    lineas_reverso('IDENTIFICACIÓN:',negrita=True,espacio_despues=3)
    for texto_info in [
        '- Código: se coloca el código de retención de acuerdo con el Anexo 4. Codificación de retención',
        '- Fecha en la que se produjo el Producto No Conforme',
        '- ITEM de Semielaborado o Producto Terminado',
        '- Lote(s) del producto',
        '- Categoría de retención inicial',
        '- Cantidad observada: Producto que potencialmente o no cumple con las especificaciones']:
        lineas_reverso(texto_info,sangria=16)
    yy-=8
    lineas_reverso('DISPOSICIÓN:',negrita=True,espacio_despues=3)
    lineas_reverso('En el casillero de disposición marcar con una X acorde a la definición y colocar fecha, nombre y firma de dicha disposición.',sangria=22,espacio_despues=8)
    lineas_reverso('TRATAMIENTO:',negrita=True,espacio_despues=3)
    lineas_reverso('En el casillero la fecha del comienzo y final del tratamiento, firma de la persona que llevó a cabo el tratamiento del producto.',sangria=22)
    lineas_reverso('Detallar Kg aprobados en segunda instancia, Kg decomiso y kg reproceso acorde con el tratamiento realizado.',sangria=22,espacio_despues=4)
    linea_categoria_y=yy
    lineas_reverso('Categoría final: Después del tratamiento a qué categoría pasó')
    c.line(X_REVERSO,linea_categoria_y-3,X_REVERSO+390,linea_categoria_y-3); yy-=10
    lineas_reverso('CONSIDERACIONES DE TRABAJO:',negrita=True,espacio_despues=3)
    lineas_reverso('Colocar la información que para determinar el costo de la NO Calidad (Personas, materiales y tiempo del retrabajo) más la fecha y la firma de la persona que tuvo a cargo de validar dicha tarea. En caso de que el retrabajo se realice en diferentes días o turnos detallar información.',sangria=22)
    lineas_reverso('Este casillero solo lo puede completar personal del área de calidad.')
    c.showPage(); c.save(); return b.getvalue()

def _pdf_hallazgo_base(tabla,rid,tipo_formato):
    """Genera los formatos ME y Detector/RX sin alterar los datos ni la logica operativa."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase.pdfmetrics import stringWidth
        from reportlab.lib.utils import ImageReader
    except ModuleNotFoundError:
        return None
    df=read_df(f'SELECT * FROM {tabla} WHERE id=?',(rid,))
    if df.empty: return None
    r=df.iloc[0].to_dict()
    def v(campo):
        x=r.get(campo,'')
        return '' if x is None or pd.isna(x) else str(x)
    def fecha_registro():
        try: return date(int(float(v('anio'))),int(float(v('mes'))),int(float(v('dia')))).strftime('%d/%m/%Y')
        except Exception: return ''
    def t(c,texto,x,y,tam=8,negrita=False):
        c.setFont('Helvetica-Bold' if negrita else 'Helvetica',tam); c.drawString(x,y,str(texto or ''))
    def fit(c,texto,x,y,ancho,tam=8,negrita=False):
        texto=str(texto or '').replace('\n',' / '); fuente='Helvetica-Bold' if negrita else 'Helvetica'; actual=tam
        while actual>5.5 and stringWidth(texto,fuente,actual)>ancho: actual-=.25
        if stringWidth(texto,fuente,actual)>ancho:
            while texto and stringWidth(texto+'...',fuente,actual)>ancho: texto=texto[:-1]
            texto+='...'
        c.setFont(fuente,actual); c.drawString(x,y,texto)
    def wrap(c,texto,x,y,ancho,tam=8,leading=11,max_lines=8):
        lineas=[]
        for bloque in str(texto or '').splitlines() or ['']:
            linea=''
            for palabra in bloque.split():
                prueba=(linea+' '+palabra).strip()
                if stringWidth(prueba,'Helvetica',tam)<=ancho: linea=prueba
                else:
                    if linea: lineas.append(linea)
                    linea=palabra
            if linea: lineas.append(linea)
        c.setFont('Helvetica',tam)
        for i,linea in enumerate(lineas[:max_lines]): c.drawString(x,y-i*leading,linea)
        if len(lineas)>max_lines: c.drawRightString(x+ancho,y-(max_lines-1)*leading,'...')
    def line_value(c,etiqueta,valor,y,label_x=45,value_x=108,right=557,bold=False):
        t(c,etiqueta,label_x,y,7.7,bold); c.line(value_x,y-2,right,y-2); fit(c,valor,value_x+3,y, right-value_x-6,7.7)
    def logo(c,x,y,w,h):
        try:
            data=BytesIO(base64.b64decode(LOGO_MUNDO_DULCE_BASE64))
            c.drawImage(ImageReader(data),x,y,width=w,height=h,preserveAspectRatio=True,anchor='c',mask='auto')
        except Exception: t(c,'Mundo Dulce',x+12,y+h/2,10,True)
    b=BytesIO(); c=canvas.Canvas(b,pagesize=A4); W,H=A4
    c.setLineWidth(.8)
    if tipo_formato=='ME':
        c.setTitle('Reporte de Hallazgos de Materia Extraña')
        # Encabezado de la primera imagen
        x0=45; x1=198; x2=508; x3=558; top=800; bot=758
        c.rect(x0,bot,x3-x0,top-bot); c.line(x1,bot,x1,top); c.line(x2,bot,x2,top)
        logo(c,x0+15,bot+6,x1-x0-30,top-bot-12)
        c.setFont('Helvetica',7.7); c.drawCentredString((x1+x2)/2,top-17,'Anexo 1. Reporte de Hallazgos de Materia Extraña')
        c.drawCentredString((x1+x2)/2,top-32,'IT-CAL03-2301-03'); c.drawCentredString((x2+x3)/2,top-25,'Rev. 0')
        t(c,'N° Hallazgo:',46,733,7.5); fit(c,f'{rid}/{v("anio") or datetime.now().year}',108,733,140,7.5,True)
        t(c,'Fecha:',46,718,7.5); fit(c,fecha_registro(),108,718,140,7.5)
        line_value(c,'Línea:',v('linea_sector'),688)
        line_value(c,'Familia:',v('familia'),673)
        line_value(c,'Equipo:',v('equipo_hallazgo'),658)
        line_value(c,'Producto:',f'{v("item")} - {v("producto")}'.strip(' -'),643)
        line_value(c,'Lote:',v('lote'),628)
        t(c,'Descripción de material hallado:',46,596,7.5)
        c.line(222,594,557,594); c.line(45,579,557,579); c.line(45,564,557,564)
        wrap(c,v('descripcion_hallazgo'),225,596,329,7.5,14,3)
        # Cuadro de muestra
        c.rect(45,339,512,210); t(c,'Muestra Hallazgo de Materia Extraña:',48,537,7.2)
        detalle='Tipo: '+v('tipo')+' | Partículas halladas: '+v('particulas_halladas')
        fit(c,detalle,48,522,500,7)
        t(c,'Acción contingente:',46,312,7.5); c.line(45,296,557,296); c.line(45,281,557,281)
        wrap(c,v('accion_contingente') or v('acciones_inmediatas'),48,299,506,7.5,14,2)
        t(c,'Investigación del origen:',46,254,7.5); c.line(45,238,557,238); c.line(45,223,557,223); c.line(45,208,557,208)
        wrap(c,v('investigacion_origen'),48,241,506,7.5,14,3)
        # Firmas
        for cx,titulo in [(165,'Calidad'),(306,'Producción'),(448,'Mantenimiento')]:
            c.line(cx-57,142,cx+57,142); c.setFont('Helvetica',7); c.drawCentredString(cx,129,titulo); c.drawCentredString(cx,117,'(Nombre y Firma)')
    else:
        c.setTitle('Registro de Materiales Segregados por Detector de Metales')
        # Encabezado de la segunda imagen
        # A4 mide 595.28 pt de ancho. El formato anterior terminaba en x=609,
        # por lo que el borde derecho y varios campos quedaban fuera de la pagina.
        # Se conserva la estructura y la logica, ajustando solo la geometria horizontal.
        x0=27; x1=187; x2=506; x3=568; top=805; mid=772; bot=720
        c.rect(x0,bot,x3-x0,top-bot); c.line(x1,bot,x1,top); c.line(x2,bot,x2,top); c.line(x1,mid,x2,mid)
        logo(c,x0+18,bot+16,x1-x0-36,top-bot-30)
        c.setFont('Helvetica-Bold',7.5); c.drawCentredString((x1+x2)/2,top-18,'Anexo 1. Registro de Materiales Segregados por el detector de')
        c.drawCentredString((x1+x2)/2,top-30,'metales')
        c.setFont('Helvetica-Bold',7.1); c.drawCentredString((x1+x2)/2,mid-23,'IT-CAL03-2301-09240-2007 Instructivo para el Tratamiento de')
        c.drawCentredString((x1+x2)/2,mid-36,'producto separado por Detector de Metales')
        c.drawCentredString((x2+x3)/2,top-44,'REV: 02')
        t(c,'RECHAZO No. :',28,643,8,True); fit(c,f'{rid}/{v("anio") or datetime.now().year}',119,643,180,8,True)
        t(c,'FECHA:',28,602,8,True); fit(c,fecha_registro(),119,602,180,8)
        t(c,'LINEA DE ELABORACIÓN / EQUIPO:',28,547,8,True); c.line(270,545,x3,545); fit(c,v('linea_sector')+' / '+v('equipo_hallazgo'),273,547,x3-276,7.8)
        t(c,'SECTOR DEL HALLAZGO:',28,520,8,True)
        # Las opciones se marcan por coincidencia con la descripción, cuando existe.
        ubic=(v('descripcion_hallazgo')+' '+v('equipo_hallazgo')).upper()
        sector_explicito=normalizar_catalogo(v('sector_hallazgo'))
        ubicacion_explicita=normalizar_catalogo(v('ubicacion_material'))
        opciones=[('CONFORMADO',219,286),('ENVOLTURA',361,425),('EMPAQUE',485,532)]
        for nombre,label_x,box_x in opciones:
            t(c,nombre,label_x,515,7.0); c.rect(box_x,507,34,15)
            if normalizar_catalogo(nombre)==sector_explicito or (not sector_explicito and nombre in ubic): c.setFont('Helvetica-Bold',11); c.drawCentredString(box_x+17,509,'X')
        line_value(c,'PRODUCTO:',f'{v("item")} - {v("producto")}'.strip(' -'),468,label_x=28,value_x=119,right=x3,bold=True)
        line_value(c,'LOTE:',v('lote'),400,label_x=28,value_x=119,right=x3,bold=True)
        t(c,'DESCRIPCIÓN DEL MATERIAL HALLADO:',28,361,8,True); c.line(320,359,x3,359); c.line(x0,334,x3,334)
        wrap(c,v('descripcion_hallazgo'),323,362,x3-326,7.6,12,2)
        t(c,'UBICACIÓN:',28,308,8,True)
        for nombre,label_x,box_x in [('MASA',190,226),('RELLENO',337,383),('RECUBIERTO',478,535)]:
            t(c,nombre,label_x,303,7.0); c.rect(box_x,295,30,15)
            if normalizar_catalogo(nombre)==ubicacion_explicita or (not ubicacion_explicita and nombre in ubic): c.setFont('Helvetica-Bold',11); c.drawCentredString(box_x+15,297,'X')
        t(c,'ADJUNTAR AQUÍ MUESTRA DEL MATERIAL HALLADO',28,267,8,True)
        # Área amplia para colocar muestra física
        c.rect(x0,143,x3-x0,105)
        t(c,'ACCIÓN INMEDIATA:',28,112,8,True); c.line(187,110,x3,110)
        fit(c,v('acciones_inmediatas') or v('accion_contingente'),190,112,x3-194,7.5)
        t(c,'INVESTIGACIÓN DEL ORIGEN:',28,85,7.5); c.line(187,83,x3,83)
        fit(c,v('investigacion_origen'),190,85,x3-194,7.5)
        for cx,l1,l2 in [(112,'Nombre y Firma del Analista de','Calidad'),(298,'Nombre y Firma del','Supervisor de Producción'),(484,'Nombre y Firma de','Mantenimiento')]:
            c.line(cx-78,44,cx+78,44); c.setFont('Helvetica',6.7); c.drawCentredString(cx,31,l1); c.drawCentredString(cx,20,l2)
    c.showPage(); c.save(); return b.getvalue()

def pdf_materia_extrana(rid):
    return _pdf_hallazgo_base('me_registros',rid,'ME')

def pdf_detector_metales_rx(rid):
    return _pdf_hallazgo_base('ddm_rx_registros',rid,'DDM')

def pdf_entrega(eid):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4,landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
    except ModuleNotFoundError:return None
    h=read_df('SELECT * FROM entregas_turno WHERE id=?',(eid,));d=read_df('SELECT * FROM entregas_turno_lineas WHERE entrega_id=? ORDER BY orden_fila',(eid,));m=read_df('SELECT * FROM matriz_entrega WHERE entrega_id=?',(eid,));sg=read_df('SELECT * FROM entregas_turno_seguimientos WHERE entrega_id=? ORDER BY bloque,orden_fila,id',(eid,))
    if h.empty:return None
    r=h.iloc[0];b=BytesIO();doc=SimpleDocTemplate(b,pagesize=landscape(A4),leftMargin=28,rightMargin=28,topMargin=28,bottomMargin=28);sty=getSampleStyleSheet();story=[Paragraph('REPORTE DE ENTREGA DE TURNO',sty['Title']),Spacer(1,8)]
    meta=[['Registro',eid,'Fecha',r.fecha,'Analista',r.analista,'Turno',r.turno],['Nave',r.nave,'Referencia',r.referencia,'Creado por',r.creado_por,'Creado en',r.creado_en]];t=Table(meta);t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,colors.grey),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DDF4EF'))]));story+=[t,Spacer(1,10)]
    if not m.empty:
        x=m.iloc[0];mt=Table([['Total carga','Horas Nave 1','Horas Nave 2','Horas Nave 3'],[f'{x.total_carga_datos:.2f}',f'{x.horas_nave1:.2f}',f'{x.horas_nave2:.2f}',f'{x.horas_nave3:.2f}']],colWidths=[165]*4);mt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#062C36')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('ALIGN',(0,0),(-1,-1),'CENTER'),('GRID',(0,0),(-1,-1),.4,colors.grey)]));story+=[mt,Spacer(1,10)]
    data=[['Nave','Grupo','Línea/Sector','Producto/Análisis','Horas','Carga','Observaciones']]
    for q in d.itertuples():data.append([str(getattr(q,'nave_catalogo','') or ''),q.grupo or '',q.linea or '',q.producto_descripcion or '',f'{float(q.horas_trabajadas or 0):.2f}',f'{float(q.carga_spac or 0):.2f}',q.observaciones or ''])
    dt=Table(data,colWidths=[55,85,145,145,45,45,190],repeatRows=1);dt.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#0A4652')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTSIZE',(0,0),(-1,-1),7),('GRID',(0,0),(-1,-1),.3,colors.grey),('VALIGN',(0,0),(-1,-1),'MIDDLE')]));story.append(dt)
    if not sg.empty:
        columnas=['registro_numero','hoja_fisica','carga_electronica','correo','descripcion_seguimiento']
        sg=sg[sg[columnas].fillna('').astype(str).apply(lambda c:c.str.strip()).ne('').any(axis=1)]
        if not sg.empty:
            story += [Spacer(1,12),Paragraph('SEGUIMIENTOS REGISTRADOS',sty['Heading2']),Spacer(1,6)]
            filas=[['Seguimiento','Registro #','Hoja física','Carga electrónica','Correo','Descripción']]
            for q in sg.itertuples(): filas.append([q.bloque or '',q.registro_numero or '',q.hoja_fisica or '',q.carga_electronica or '',q.correo or '',q.descripcion_seguimiento or ''])
            ts=Table(filas,colWidths=[130,66,65,78,52,300],repeatRows=1);ts.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#062C36')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTSIZE',(0,0),(-1,-1),7),('GRID',(0,0),(-1,-1),.3,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP')]));story.append(ts)
    doc.build(story);return b.getvalue()

def estilo_faltantes_matriz(df, primeras_columnas=1):
    def pintar(valor):
        return 'background-color:#FECACA;color:#991B1B;font-weight:800;border:1px solid #EF4444' if pd.isna(valor) or str(valor).strip()=='' else ''
    return df.style.map(pintar,subset=list(df.columns[primeras_columnas:])).format(na_rep='')

def clave_periodo(fecha,tipo):
    if tipo=='SEMANA':
        iso=fecha.isocalendar(); return f'{iso.year}-S{iso.week:02d}'
    return fecha.strftime('%Y-%m')

def editor_metas_cumplimiento(tipo_entidad, entidades, teoricos, periodo_tipo, periodo_clave, key):
    """Datos teóricos vienen de registros; datos cargados se capturan manualmente."""
    guardados=read_df('SELECT entidad,meta FROM metas_carga_spac WHERE tipo_entidad=? AND periodo_tipo=? AND periodo_clave=?',(tipo_entidad,periodo_tipo,periodo_clave))
    mapa={str(r.entidad):float(r.meta or 0) for r in guardados.itertuples()}
    filas=[]
    for entidad in entidades:
        teorico=float(teoricos.get(entidad,0) or 0)
        cargado=float(mapa.get(entidad,0) or 0)
        porcentaje=(cargado/teorico*100) if teorico>0 else None
        filas.append({'Entidad':entidad,'Datos teóricos':teorico,'Datos cargados':cargado,'Cumplimiento %':porcentaje})
    base=pd.DataFrame(filas)
    if is_dev():
        editado=st.data_editor(base,use_container_width=True,hide_index=True,key=key,disabled=['Entidad','Datos teóricos','Cumplimiento %'],column_config={'Datos teóricos':st.column_config.NumberColumn(format='%.2f'),'Datos cargados':st.column_config.NumberColumn(min_value=0.0,step=1.0,format='%.2f'),'Cumplimiento %':st.column_config.NumberColumn(format='%.1f%%')})
        if st.button('Guardar datos cargados y calcular cumplimiento',key=key+'_guardar',type='primary'):
            c=conn();cur=c.cursor()
            try:
                for _,r in editado.iterrows():
                    cur.execute('INSERT INTO metas_carga_spac(tipo_entidad,entidad,periodo_tipo,periodo_clave,meta,actualizado_por,actualizado_en) VALUES(?,?,?,?,?,?,?) ON CONFLICT(tipo_entidad,entidad,periodo_tipo,periodo_clave) DO UPDATE SET meta=excluded.meta,actualizado_por=excluded.actualizado_por,actualizado_en=excluded.actualizado_en',(tipo_entidad,str(r['Entidad']),periodo_tipo,periodo_clave,float(r['Datos cargados'] or 0),st.session_state.auth['usuario'],now_iso()))
                c.commit()
            except Exception:
                c.rollback();raise
            finally:c.close()
            audit(st.session_state.auth['usuario'],'ACTUALIZAR_DATOS_CARGADOS_SPAC',f'{tipo_entidad} | {periodo_tipo} | {periodo_clave}')
            st.success('Datos cargados guardados. Cumplimiento recalculado.'); st.rerun()
    else:
        vista=base.copy(); vista['Cumplimiento %']=vista['Cumplimiento %'].map(lambda x:'' if pd.isna(x) else f'{x:.1f}%')
        st.dataframe(vista,use_container_width=True,hide_index=True)

def aplicar_ajustes_diarios(tabla,tipo_entidad,columna_entidad):
    ajustes=read_df('SELECT entidad,fecha,valor FROM ajustes_diarios_spac WHERE tipo_entidad=?',(tipo_entidad,))
    for r in ajustes.itertuples():
        if str(r.entidad) in tabla.index and str(r.fecha) in tabla.columns:
            tabla.loc[str(r.entidad),str(r.fecha)]=float(r.valor) if pd.notna(r.valor) else None
    return tabla

def tabla_diaria_editable(tabla,tipo_entidad,columna_entidad,key):
    """Tabla unica: vacios rojos y edicion directa para desarrolladores."""
    vista=tabla.reset_index().rename(columns={tabla.index.name or 'index':columna_entidad})
    if not is_dev():
        st.dataframe(estilo_faltantes_matriz(vista),use_container_width=True,hide_index=True)
        return

    # Streamlit no conserva el color de fondo en celdas numericas editables.
    # Para no duplicar la tabla, los vacios se muestran como un marcador rojo editable.
    editable=vista.copy()
    columnas_valor=[c for c in editable.columns if c!=columna_entidad]
    for c in columnas_valor:
        editable[c]=editable[c].map(lambda v:'🟥' if pd.isna(v) or str(v).strip()=='' else str(float(v)))

    st.caption('Edicion directa habilitada. La casilla 🟥 indica que falta el registro. Sustituye la casilla por 0 para justificar un dia sin carga o por el valor correspondiente.')
    editado=st.data_editor(
        editable,use_container_width=True,hide_index=True,key=key,
        disabled=[columna_entidad],num_rows='fixed',
        column_config={columna_entidad:st.column_config.TextColumn(columna_entidad,disabled=True),**{c:st.column_config.TextColumn(c) for c in columnas_valor}}
    )
    if st.button('Guardar ajustes diarios',key=key+'_guardar',type='primary'):
        errores=[]; cambios=[]
        for _,r in editado.iterrows():
            entidad=str(r[columna_entidad])
            for fecha_columna in columnas_valor:
                texto=str(r[fecha_columna] or '').strip()
                if texto in ('','🟥'):
                    valor=None
                else:
                    try: valor=float(texto.replace(',','.'))
                    except ValueError:
                        errores.append(f'{entidad} / {fecha_columna}')
                        continue
                    if valor<0:
                        errores.append(f'{entidad} / {fecha_columna}')
                        continue
                original=tabla.loc[entidad,fecha_columna] if entidad in tabla.index and fecha_columna in tabla.columns else None
                igual=(valor is None and pd.isna(original)) or (valor is not None and pd.notna(original) and float(valor)==float(original))
                if not igual: cambios.append((entidad,fecha_columna,valor))
        if errores:
            st.error('Corrige los valores no numericos o negativos en: '+', '.join(errores)+'.')
        else:
            c=conn();cur=c.cursor()
            try:
                for entidad,fecha_columna,valor in cambios:
                    if valor is None:
                        cur.execute('DELETE FROM ajustes_diarios_spac WHERE tipo_entidad=? AND entidad=? AND fecha=?',(tipo_entidad,entidad,fecha_columna))
                    else:
                        cur.execute('INSERT INTO ajustes_diarios_spac(tipo_entidad,entidad,fecha,valor,justificacion,actualizado_por,actualizado_en) VALUES(?,?,?,?,?,?,?) ON CONFLICT(tipo_entidad,entidad,fecha) DO UPDATE SET valor=excluded.valor,justificacion=excluded.justificacion,actualizado_por=excluded.actualizado_por,actualizado_en=excluded.actualizado_en',(tipo_entidad,entidad,fecha_columna,float(valor),'Ajuste administrativo',st.session_state.auth['usuario'],now_iso()))
                c.commit()
            except Exception:
                c.rollback();raise
            finally:c.close()
            audit(st.session_state.auth['usuario'],'AJUSTAR_MATRIZ_DIARIA_SPAC',tipo_entidad)
            st.success('Ajustes diarios guardados correctamente.');st.rerun()


def grafica_lineas_con_valores(df,titulo,key):
    """Grafica lineal con puntos y etiquetas visibles sin seleccionar la grafica."""
    if df.empty or df['Cumplimiento %'].dropna().empty:
        st.info(f'No hay porcentajes disponibles para {titulo}.')
        return
    datos=df.dropna(subset=['Cumplimiento %']).copy()
    datos['Cumplimiento %']=pd.to_numeric(datos['Cumplimiento %'],errors='coerce')
    base=alt.Chart(datos).encode(
        x=alt.X('Periodo:N',title='Periodo',sort=None),
        y=alt.Y('Cumplimiento %:Q',title='Cumplimiento (%)'),
        color=alt.Color('Entidad:N',title='Serie'),
        detail='Entidad:N',
        tooltip=['Periodo:N','Entidad:N',alt.Tooltip('Cumplimiento %:Q',format='.1f')]
    )
    linea=base.mark_line(point=True,strokeWidth=3)
    etiquetas=base.mark_text(dy=-12,fontSize=11,fontWeight='bold').encode(text=alt.Text('Cumplimiento %:Q',format='.1f'))
    regla=alt.Chart(pd.DataFrame({'Cumplimiento %':[100]})).mark_rule(color='#16A34A',strokeDash=[6,4]).encode(y='Cumplimiento %:Q')
    st.altair_chart((linea+etiquetas+regla).properties(title=titulo,height=330),use_container_width=True,key=key)


def datos_grafica_cumplimiento(tipo_entidad,periodo_tipo):
    cargados=read_df('SELECT entidad,periodo_clave,meta FROM metas_carga_spac WHERE tipo_entidad=? AND periodo_tipo=? ORDER BY periodo_clave,entidad',(tipo_entidad,periodo_tipo))
    if cargados.empty:
        return pd.DataFrame(columns=['Periodo','Entidad','Cumplimiento %'])

    registros=read_df('SELECT fecha,analista,total_carga_datos,horas_nave1,horas_nave2,horas_nave3 FROM matriz_entrega')
    registros['fecha_dt']=pd.to_datetime(registros['fecha'],errors='coerce').dt.date
    ajustes=read_df('SELECT entidad,fecha,valor FROM ajustes_diarios_spac WHERE tipo_entidad=?',(tipo_entidad,))
    resultado=[]

    for fila in cargados.itertuples():
        periodo=str(fila.periodo_clave)
        datos_cargados=float(fila.meta or 0)
        if periodo_tipo=='SEMANA':
            anio,semana=periodo.split('-S')
            inicio=date.fromisocalendar(int(anio),int(semana),1)
            fin=date.fromisocalendar(int(anio),int(semana),7)
        else:
            inicio=datetime.strptime(periodo+'-01','%Y-%m-%d').date()
            fin=(pd.Timestamp(inicio)+pd.offsets.MonthBegin(1)).date()-timedelta(days=1)
        fechas=[x.date().isoformat() for x in pd.date_range(inicio,fin,freq='D')]
        entidad=str(fila.entidad)

        if tipo_entidad=='ANALISTA':
            serie=registros[(registros['fecha_dt']>=inicio)&(registros['fecha_dt']<=fin)&(registros['analista'].astype(str)==entidad)].groupby('fecha')['total_carga_datos'].sum().reindex(fechas)
        elif tipo_entidad=='CALIDAD_NAVE':
            calidad=read_df("SELECT m.fecha,m.total_carga_datos,COALESCE(e.nave,'') AS nave FROM matriz_entrega m LEFT JOIN entregas_turno e ON e.id=m.entrega_id")
            calidad['fecha_dt']=pd.to_datetime(calidad['fecha'],errors='coerce').dt.date
            serie=calidad[(calidad['fecha_dt']>=inicio)&(calidad['fecha_dt']<=fin)&(calidad['nave'].astype(str)==entidad)].groupby('fecha')['total_carga_datos'].sum().reindex(fechas)
        else:
            campo={'Nave 1':'horas_nave1','Nave 2':'horas_nave2','Nave 3':'horas_nave3'}.get(entidad)
            if not campo: continue
            serie=registros[(registros['fecha_dt']>=inicio)&(registros['fecha_dt']<=fin)].groupby('fecha')[campo].sum(min_count=1).reindex(fechas)

        for ajuste in ajustes[ajustes['entidad'].astype(str)==entidad].itertuples():
            if str(ajuste.fecha) in serie.index:
                serie.loc[str(ajuste.fecha)]=float(ajuste.valor)

        datos_teoricos=float(serie.fillna(0).sum())
        porcentaje=(datos_cargados/datos_teoricos*100) if datos_teoricos>0 else None
        resultado.append({'Periodo':periodo,'Entidad':entidad,'Cumplimiento %':porcentaje})

    return pd.DataFrame(resultado)

def matriz_entregas():
    st.markdown('## Entregas de turno y matriz histórica')
    st.caption('Consulta los registros de las tres naves. Selecciona un registro para descargar su reporte en PDF.')
    df=read_df("SELECT m.id,m.fecha,m.analista,m.entrega_id,m.total_carga_datos,m.horas_nave1,m.horas_nave2,m.horas_nave3,COALESCE(e.nave,'') AS nave,COALESCE(e.turno,'') AS turno FROM matriz_entrega m LEFT JOIN entregas_turno e ON e.id=m.entrega_id ORDER BY m.fecha DESC,m.analista")
    if df.empty:
        st.info('Todavía no hay entregas de turno registradas.')
        return
    df['fecha_dt']=pd.to_datetime(df['fecha'],errors='coerce').dt.date
    validas=df['fecha_dt'].dropna()
    if validas.empty:
        st.warning('Los registros existentes no contienen fechas válidas.')
        return
    mn,mx=min(validas),max(validas)
    a,b,c=st.columns(3)
    desde=a.date_input('Fecha inicial',mn,min_value=mn,max_value=mx,key='mx_desde')
    hasta=b.date_input('Fecha final',mx,min_value=mn,max_value=mx,key='mx_hasta')
    naves=c.multiselect('Nave',['Nave 1','Nave 2','Nave 3'],key='mx_naves')
    if desde>hasta:
        st.error('La fecha inicial no puede ser posterior a la fecha final.')
        return
    nombres=sorted(df['analista'].dropna().astype(str).unique())
    analistas=st.multiselect('Analista',nombres,key='mx_names')
    f=df[(df['fecha_dt']>=desde)&(df['fecha_dt']<=hasta)].copy()
    if naves:f=f[f['nave'].isin(naves)]
    if analistas:f=f[f['analista'].isin(analistas)]
    if f.empty:
        st.info('No existen registros para los filtros seleccionados.')
        return
    vista=f[['id','entrega_id','fecha','nave','turno','analista','total_carga_datos','horas_nave1','horas_nave2','horas_nave3']].copy()
    vista['PDF']='Disponible'
    vista=vista.rename(columns={'id':'ID matriz','entrega_id':'Registro','fecha':'Fecha','nave':'Nave','turno':'Turno','analista':'Analista','total_carga_datos':'Total carga','horas_nave1':'Horas Nave 1','horas_nave2':'Horas Nave 2','horas_nave3':'Horas Nave 3'})
    nonce=st.session_state.get('mx_table_nonce',0)
    evento=st.dataframe(vista,use_container_width=True,hide_index=True,on_select='rerun',selection_mode='single-row',key=f'mx_table_{nonce}',column_config={'PDF':st.column_config.TextColumn('PDF del registro',help='Selecciona la fila para descargar el PDF.')})
    excel_df=f[['fecha','nave','turno','analista','total_carga_datos','horas_nave1','horas_nave2','horas_nave3']].copy()
    st.download_button('Descargar matriz en Excel',excel_matriz(excel_df),f'matriz_entregas_{desde}_{hasta}.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key='mx_excel')
    rows=getattr(evento,'selection',{}).get('rows',[]) if evento is not None else []
    valida=bool(rows) and isinstance(rows[0],int) and 0<=rows[0]<len(vista)
    if valida:
        seleccion=vista.iloc[rows[0]]
        rid=int(seleccion['ID matriz'])
        coincidencia=f[f['id']==rid]
        if not coincidencia.empty:
            row=coincidencia.iloc[0]
            eid=int(row['entrega_id']) if pd.notna(row['entrega_id']) else None
            st.markdown(f"### Registro seleccionado: {eid or 'Sin número'} | {row['fecha']} | {row['nave']} | {row['analista']}")
            if eid:
                contenido_pdf=pdf_entrega(eid)
                if contenido_pdf:
                    st.download_button('Descargar PDF de este registro',contenido_pdf,f'entrega_turno_{eid}_{row["fecha"]}.pdf','application/pdf',key=f'mx_pdf_{rid}_{eid}')
                else:
                    st.warning('No fue posible generar el PDF. Verifica que reportlab esté incluido en requirements.txt.')
            else:
                st.warning('Este indicador no tiene una entrega de turno relacionada.')
            if is_dev() and st.button('Eliminar registro seleccionado',key=f'mx_del_{rid}'):
                st.session_state.mx_confirm=rid
            if is_dev() and st.session_state.get('mx_confirm')==rid:
                st.warning('Se eliminarán de forma permanente la matriz y la entrega de turno relacionadas. Esta acción no se puede deshacer.')
                x,y=st.columns(2)
                if x.button('Confirmar eliminación',key=f'mx_ok_{rid}'):
                    cdb=conn();cur=cdb.cursor()
                    try:
                        if eid:
                            cur.execute('DELETE FROM entregas_turno_seguimientos WHERE entrega_id=?',(eid,))
                            cur.execute('DELETE FROM entregas_turno_lineas WHERE entrega_id=?',(eid,))
                            cur.execute('DELETE FROM entregas_turno WHERE id=?',(eid,))
                        cur.execute('DELETE FROM matriz_entrega WHERE id=?',(rid,))
                        cdb.commit()
                    except Exception:
                        cdb.rollback();raise
                    finally:cdb.close()
                    audit(st.session_state.auth['usuario'],'ELIMINAR_DIA_MATRIZ',f'Matriz {rid} | Entrega {eid}')
                    st.session_state.pop('mx_confirm',None)
                    st.session_state.mx_table_nonce=nonce+1
                    st.rerun()
                if y.button('Cancelar',key=f'mx_cancel_{rid}'):
                    st.session_state.pop('mx_confirm',None)
                    st.session_state.mx_table_nonce=nonce+1
                    st.rerun()
    elif rows:
        st.session_state.mx_table_nonce=nonce+1
        st.rerun()
    st.markdown('### Matrices por fecha y analista')
    for campo,titulo in [('total_carga_datos','Total de carga de datos'),('horas_nave1','Horas trabajadas Nave 1'),('horas_nave2','Horas trabajadas Nave 2'),('horas_nave3','Horas trabajadas Nave 3')]:
        st.markdown(f'**{titulo}**')
        st.dataframe(f.pivot_table(index='analista',columns='fecha',values=campo,aggfunc='sum',fill_value=0),use_container_width=True)

    st.markdown('## Control diario de datos SPAC')
    st.caption('Analistas: TOTAL DE CARGA DE DATOS. Naves: suma diaria de Horas Nave 1, Horas Nave 2 y Horas Nave 3. Las celdas rojas vacías indican ausencia de registro.')
    columnas_fecha=[x.date().isoformat() for x in pd.date_range(desde,hasta,freq='D')]
    analistas_catalogo=catalog('analista')
    datos_rango=df[(df['fecha_dt']>=desde)&(df['fecha_dt']<=hasta)].copy()
    pa=datos_rango.pivot_table(index='analista',columns='fecha',values='total_carga_datos',aggfunc='sum').reindex(index=analistas_catalogo,columns=columnas_fecha)
    pa.index.name='Analista'; pa=aplicar_ajustes_diarios(pa,'ANALISTA','Analista')
    st.markdown('### TOTAL DE CARGA DE DATOS diario por analista')
    tabla_diaria_editable(pa,'ANALISTA','Analista','tabla_diaria_analistas')

    calidad_diaria=datos_rango.pivot_table(index='nave',columns='fecha',values='total_carga_datos',aggfunc='sum').reindex(index=['Nave 1','Nave 2','Nave 3'],columns=columnas_fecha)
    calidad_diaria.index.name='Nave'; calidad_diaria=aplicar_ajustes_diarios(calidad_diaria,'CALIDAD_NAVE','Nave')
    st.markdown('### TOTAL DE CARGA DE DATOS diario por plantilla')
    st.caption('Suma el TOTAL DE CARGA DE DATOS segun la plantilla utilizada: Nave 1, Nave 2 o Nave 3.')
    tabla_diaria_editable(calidad_diaria,'CALIDAD_NAVE','Nave','tabla_diaria_calidad_naves')

    filas_nave=[]
    for nv,campo in [('Nave 1','horas_nave1'),('Nave 2','horas_nave2'),('Nave 3','horas_nave3')]:
        serie=datos_rango.groupby('fecha')[campo].sum(min_count=1).reindex(columnas_fecha)
        fila={'Nave':nv}; fila.update({f:(float(serie.get(f)) if pd.notna(serie.get(f)) and float(serie.get(f))>0 else None) for f in columnas_fecha}); filas_nave.append(fila)
    pn=pd.DataFrame(filas_nave,columns=['Nave']+columnas_fecha).set_index('Nave'); pn=aplicar_ajustes_diarios(pn,'NAVE','Nave')
    st.markdown('### Horas diarias consolidadas por nave')
    tabla_diaria_editable(pn,'NAVE','Nave','tabla_diaria_naves')
    st.markdown('## Datos teóricos, datos cargados y cumplimiento SPAC')
    st.info('Fórmula: Cumplimiento (%) = Datos cargados ÷ Datos teóricos × 100. Los datos teóricos provienen de los registros y ajustes diarios. Los datos cargados se capturan manualmente.')
    pt=st.radio('Periodo de evaluación',['SEMANA','MES'],horizontal=True,format_func=lambda x:'Semana' if x=='SEMANA' else 'Mes',key='meta_periodo_tipo')
    opciones=[]
    for x in pd.date_range(desde,hasta,freq='D').date:
        k=clave_periodo(x,pt)
        if k not in opciones: opciones.append(k)
    pc=st.selectbox('Semana o mes',opciones,key='meta_periodo_clave')
    if pt=='SEMANA':
        anio,semana=pc.split('-S'); inicio=date.fromisocalendar(int(anio),int(semana),1); fin=date.fromisocalendar(int(anio),int(semana),7)
    else:
        inicio=datetime.strptime(pc+'-01','%Y-%m-%d').date(); fin=(pd.Timestamp(inicio)+pd.offsets.MonthBegin(1)).date()-timedelta(days=1)
    dp=df[(df['fecha_dt']>=inicio)&(df['fecha_dt']<=fin)].copy()
    fechas_meta=[x.date().isoformat() for x in pd.date_range(inicio,fin,freq='D')]
    pa_meta=dp.pivot_table(index='analista',columns='fecha',values='total_carga_datos',aggfunc='sum').reindex(index=analistas_catalogo,columns=fechas_meta)
    pa_meta.index.name='Analista'; pa_meta=aplicar_ajustes_diarios(pa_meta,'ANALISTA','Analista')
    teoricos_a=pa_meta.fillna(0).sum(axis=1).to_dict()
    filas_meta=[]
    for nv,campo in [('Nave 1','horas_nave1'),('Nave 2','horas_nave2'),('Nave 3','horas_nave3')]:
        serie=dp.groupby('fecha')[campo].sum(min_count=1).reindex(fechas_meta)
        filas_meta.append([float(v) if pd.notna(v) and float(v)>0 else None for v in serie])
    pn_meta=pd.DataFrame(filas_meta,index=['Nave 1','Nave 2','Nave 3'],columns=fechas_meta); pn_meta.index.name='Nave'; pn_meta=aplicar_ajustes_diarios(pn_meta,'NAVE','Nave')
    teoricos_n=pn_meta.fillna(0).sum(axis=1).to_dict()
    calidad_meta=dp.pivot_table(index='nave',columns='fecha',values='total_carga_datos',aggfunc='sum').reindex(index=['Nave 1','Nave 2','Nave 3'],columns=fechas_meta)
    calidad_meta.index.name='Nave'; calidad_meta=aplicar_ajustes_diarios(calidad_meta,'CALIDAD_NAVE','Nave')
    teoricos_calidad=calidad_meta.fillna(0).sum(axis=1).to_dict()
    ca,cn=st.columns(2)
    with ca:
        st.markdown('### Cumplimiento por analista'); editor_metas_cumplimiento('ANALISTA',analistas_catalogo,teoricos_a,pt,pc,f'meta_a_{pt}_{pc}')
    with cn:
        st.markdown('### Indicador SPAC Produccion'); editor_metas_cumplimiento('NAVE',['Nave 1','Nave 2','Nave 3'],teoricos_n,pt,pc,f'meta_n_{pt}_{pc}')
    st.markdown('### Indicador SPAC Calidad')
    editor_metas_cumplimiento('CALIDAD_NAVE',['Nave 1','Nave 2','Nave 3'],teoricos_calidad,pt,pc,f'meta_calidad_{pt}_{pc}')



def formato_entrega(nave,tipo='PROCESO'):
    df=read_df("SELECT linea,sector,tipo_analisis,orden_linea,orden_sector FROM catalogo_formatos_entrega WHERE formato_nave=? AND tipo=? AND activo=1 ORDER BY orden_linea,orden_sector,id",(nave,tipo))
    if tipo=='ANALISIS':
        return [(str(r.linea),str(r.sector),str(r.tipo_analisis or '')) for r in df.itertuples()]
    resultado={}
    for r in df.itertuples(): resultado.setdefault(str(r.linea),[]).append(str(r.sector))
    return resultado
def formato_seguimientos():
    df=read_df('SELECT nombre FROM catalogo_seguimientos_entrega WHERE activo=1 ORDER BY orden,id')
    return df['nombre'].dropna().astype(str).str.strip().loc[lambda x:x.ne('')].tolist() if not df.empty else []

def configuracion_seguimientos():
    bloques=read_df('SELECT id,nombre FROM catalogo_seguimientos_entrega WHERE activo=1 ORDER BY orden,id')
    resultado=[]
    for bloque in bloques.itertuples():
        campos=read_df('SELECT id,nombre,tipo_campo,opciones,obligatorio,orden FROM catalogo_seguimientos_campos WHERE seguimiento_id=? AND activo=1 ORDER BY orden,id',(int(bloque.id),))
        definiciones=[]
        for campo in campos.itertuples():
            opciones=[x.strip() for x in str(campo.opciones or '').split('|') if x.strip()]
            definiciones.append({'id':int(campo.id),'nombre':str(campo.nombre),'tipo':str(campo.tipo_campo or 'Texto'),'opciones':opciones,'obligatorio':bool(campo.obligatorio)})
        resultado.append({'id':int(bloque.id),'nombre':str(bloque.nombre),'campos':definiciones})
    return resultado

def editor_seguimiento_dinamico(bloque,key):
    campos=bloque.get('campos',[])
    if not campos:
        st.info('Este seguimiento no tiene campos activos. Un administrador puede configurarlos desde Catálogos.')
        return pd.DataFrame(),{}
    base=pd.DataFrame([{c['nombre']:'' for c in campos} for _ in range(3)])
    config={}
    for c in campos:
        nombre=c['nombre']; tipo=c['tipo']; opciones=c.get('opciones',[])
        if tipo=='Lista desplegable': config[nombre]=st.column_config.SelectboxColumn(nombre,options=['']+opciones,required=c['obligatorio'])
        elif tipo=='Número': config[nombre]=st.column_config.NumberColumn(nombre,min_value=0.0,step=1.0,required=c['obligatorio'])
        elif tipo=='Sí / No / N/A': config[nombre]=st.column_config.SelectboxColumn(nombre,options=['','Sí','No','N/A'],required=c['obligatorio'])
        elif tipo=='Fecha': config[nombre]=st.column_config.DateColumn(nombre,required=c['obligatorio'])
        else: config[nombre]=st.column_config.TextColumn(nombre,required=c['obligatorio'])
    return st.data_editor(base,num_rows='dynamic',use_container_width=True,hide_index=True,key=key,column_config=config),{c['nombre']:c for c in campos}




def init_db():
    """Valida la conexión remota. El esquema se despliega con supabase_schema.sql."""
    read_df('SELECT 1 AS ok')

def reset_admin():
    pw=hash_password(ADMIN_PASS); c=conn(); cur=c.cursor(); cur.execute("SELECT id FROM usuarios WHERE usuario=?",(ADMIN_USER,)); exists=cur.fetchone()
    if exists: cur.execute("UPDATE usuarios SET nombre=?, password_hash=?, rol=?, activo=1, intentos_fallidos=0, requiere_cambio_pass=0 WHERE usuario=?",("Administrador del sistema",pw,"desarrollador",ADMIN_USER))
    else: cur.execute("INSERT INTO usuarios(usuario,nombre,password_hash,rol,activo,intentos_fallidos,requiere_cambio_pass,creado_en) VALUES(?,?,?,?,1,0,0,?)",(ADMIN_USER,"Administrador del sistema",pw,"desarrollador",now_iso()))
    c.commit(); c.close()

def auth_user(u,p):
    """Autentica y controla 8 fallos: 5 base, 3 advertencias finales y bloqueo."""
    usuario=str(u or '').strip()
    c=conn(); cur=c.cursor()
    cur.execute("SELECT id,usuario,nombre,password_hash,rol,activo,COALESCE(intentos_fallidos,0),COALESCE(requiere_cambio_pass,0) FROM usuarios WHERE usuario=?",(usuario,))
    row=cur.fetchone()
    if not row:
        c.close()
        return {'estado':'credenciales_invalidas','intentos':None,'restantes':None}
    uid,user,nombre,pw,rol,activo,intentos,requiere=row
    intentos=int(intentos or 0)
    if int(activo or 0)!=1:
        c.close()
        return {'estado':'bloqueado','usuario':user}
    if check_password(p,pw):
        cur.execute('UPDATE usuarios SET intentos_fallidos=0 WHERE id=?',(uid,)); c.commit(); c.close()
        return {'estado':'ok','usuario':user,'nombre':nombre,'rol':rol,'requiere_cambio_pass':int(requiere or 0)}
    nuevos=intentos+1
    bloqueado=nuevos>=8
    cur.execute('UPDATE usuarios SET intentos_fallidos=?, activo=? WHERE id=?',(min(nuevos,8),0 if bloqueado else 1,uid))
    c.commit(); c.close()
    if bloqueado:
        return {'estado':'bloqueado_por_intentos','usuario':user,'intentos':8,'restantes':0}
    return {'estado':'advertencia' if nuevos>=6 else 'credenciales_invalidas','usuario':user,'intentos':nuevos,'restantes':8-nuevos}


def cambiar_password_obligatoria():
    auth=st.session_state.get('auth') or {}
    if not auth or not int(auth.get('requiere_cambio_pass',0) or 0):
        return
    styles(False)
    st.markdown('<div class="login-card"><div class="login-title">Cambio de contraseña obligatorio</div></div>',unsafe_allow_html=True)
    st.warning('Debes actualizar la contraseña inicial antes de acceder al sistema.')
    with st.form('cambio_password_obligatorio'):
        nueva=st.text_input('Nueva contraseña',type='password')
        confirmar=st.text_input('Confirmar nueva contraseña',type='password')
        guardar=st.form_submit_button('Actualizar contraseña',type='primary')
    if guardar:
        errores=[]
        if len(nueva)<8: errores.append('mínimo 8 caracteres')
        if not re.search(r'[A-ZÁÉÍÓÚÑ]',nueva): errores.append('una mayúscula')
        if not re.search(r'[a-záéíóúñ]',nueva): errores.append('una minúscula')
        if not re.search(r'\d',nueva): errores.append('un número')
        if nueva!=confirmar: errores.append('las contraseñas deben coincidir')
        if errores:
            st.error('La contraseña requiere: '+', '.join(errores)+'.')
        else:
            exec_sql('UPDATE usuarios SET password_hash=?,requiere_cambio_pass=0,intentos_fallidos=0 WHERE usuario=?',(hash_password(nueva),auth['usuario']))
            audit(auth['usuario'],'CAMBIO_PASSWORD_OBLIGATORIO','Contraseña inicial actualizada')
            st.session_state.auth['requiere_cambio_pass']=0
            st.success('Contraseña actualizada correctamente.')
            st.rerun()
    st.stop()

def audit(u,a,d): exec_sql("INSERT INTO auditoria(usuario,accion,detalle,fecha_hora) VALUES(?,?,?,?)",(u,a,d,now_iso()))
def normalizar_nave(valor):
    texto=normalizar_catalogo(valor)
    mapa={'1':'Nave 1','NAVE 1':'Nave 1','2':'Nave 2','NAVE 2':'Nave 2','3':'Nave 3','NAVE 3':'Nave 3'}
    return mapa.get(texto,str(valor or '').strip())

def catalog(cat):
    df=read_df("SELECT valor FROM catalogos WHERE categoria=? AND activo=1 ORDER BY valor",(cat,))
    valores=df['valor'].tolist() if not df.empty else []
    if cat=='nave':
        # Evita que 1 y Nave 1 se presenten como entidades distintas.
        return list(dict.fromkeys(normalizar_nave(v) for v in valores if normalizar_nave(v)))
    return valores
def is_dev(): return st.session_state.get('auth',{}).get('rol')=='desarrollador'
def new_folio():
    pref=f"PNC-{datetime.now().year}-"; df=read_df("SELECT folio FROM pnc_registros WHERE folio LIKE ? ORDER BY folio DESC LIMIT 1",(f"{pref}%",)); n=0
    if not df.empty:
        try: n=int(str(df.iloc[0]['folio']).split('-')[-1])
        except Exception: n=0
    return f"{pref}{n+1:05d}"

def save_files(files,rid,folio,user):
    for file in files or []:
        ruta = upload_evidence(file, folio)
        exec_sql('INSERT INTO adjuntos(registro_id,folio,nombre_original,ruta_archivo,tipo_archivo,subido_por,subido_en) VALUES(?,?,?,?,?,?,?)',(rid,folio,file.name,ruta,file.type,user,now_iso()))

def styles(compact=False):
    st.markdown("""
    <style>
    :root { --navy:#062C36; --navy2:#0A4652; --navy3:#083640; --teal:#00A884; --blue:#3F7BFF; --violet:#5850EC; --bg:#F2F5F8; --text:#203047; }
    .stApp { background:var(--bg); }
    header[data-testid="stHeader"] { display:none; }
    section[data-testid="stSidebar"] { display:none !important; }
    [data-testid="stSidebarCollapsedControl"], button[title="Open sidebar"], button[title="Close sidebar"] { display:none !important; }
    .main .block-container { padding:1rem 1.25rem 2rem 1.25rem; max-width:100% !important; }

    /* Layout principal: columna izquierda SIEMPRE visible */
    div[data-testid="stHorizontalBlock"]:has(.left-menu-marker) { align-items:stretch !important; gap:1.2rem !important; }
    div[data-testid="stHorizontalBlock"]:has(.left-menu-marker) > div[data-testid="column"]:first-child {
        background:linear-gradient(180deg,#062C36 0%,#0A4652 68%,#083640 100%) !important;
        border-radius:0 28px 28px 0 !important;
        box-shadow:0 16px 34px rgba(7,49,61,.30) !important;
        padding:1.25rem 1rem !important;
        min-height:calc(100vh - 1rem) !important;
        box-sizing:border-box !important;
        overflow:hidden !important;
        align-self:stretch !important;
    }
    div[data-testid="stHorizontalBlock"]:has(.left-menu-marker) > div[data-testid="column"]:first-child * { color:#FFFFFF !important; }
    div[data-testid="stHorizontalBlock"]:has(.left-menu-marker) > div[data-testid="column"]:last-child {
        min-width:0 !important;
        padding-left:.25rem !important;
        box-sizing:border-box !important;
    }
    /* Fallback por si el navegador no aplica :has correctamente */
    .main .block-container > div > div > div[data-testid="stHorizontalBlock"]:first-of-type > div[data-testid="column"]:first-child {
        background:linear-gradient(180deg,#062C36 0%,#0A4652 68%,#083640 100%) !important;
        border-radius:0 28px 28px 0 !important;
        box-shadow:0 16px 34px rgba(7,49,61,.30) !important;
        padding:1.25rem 1rem !important;
        min-height:calc(100vh - 1rem) !important;
        box-sizing:border-box !important;
        overflow:hidden !important;
    }

    .menu-brand { display:flex; align-items:center; gap:.6rem; margin:1rem 0 1.55rem 0; font-size:1.12rem; font-weight:950; white-space:nowrap; color:#FFFFFF !important; }
    .menu-section { color:#BDEFE5 !important; font-size:.74rem; font-weight:900; letter-spacing:.06rem; margin:.4rem 0 .7rem 0; }
    div[data-testid="stHorizontalBlock"]:has(.left-menu-marker) > div[data-testid="column"]:first-child .stButton button {
        width:auto !important;
        max-width:100% !important;
        min-height:44px !important;
        display:inline-flex !important;
        align-items:center !important;
        justify-content:flex-start !important;
        text-align:left !important;
        background:rgba(255,255,255,.18) !important;
        border:1px solid rgba(255,255,255,.34) !important;
        border-radius:12px !important;
        color:#FFFFFF !important;
        font-weight:850 !important;
        margin:.18rem 0 !important;
        padding:.58rem .75rem !important;
        box-shadow:0 7px 16px rgba(0,0,0,.10) !important;
        overflow:hidden !important;
        white-space:nowrap !important;
    }
    div[data-testid="stHorizontalBlock"]:has(.left-menu-marker) > div[data-testid="column"]:first-child .stButton button:hover {
        background:rgba(255,255,255,.28) !important;
        border-color:rgba(255,255,255,.55) !important;
        color:#FFFFFF !important;
    }
    .menu-active {
        display:inline-flex !important;
        align-items:center !important;
        justify-content:flex-start !important;
        width:auto !important;
        max-width:100% !important;
        min-height:46px !important;
        background:linear-gradient(135deg,#00A884 0%,#3F7BFF 58%,#5850EC 100%) !important;
        color:#FFFFFF !important;
        border-radius:12px !important;
        padding:.62rem .8rem !important;
        font-weight:950 !important;
        margin:.32rem 0 .45rem 0 !important;
        box-shadow:0 10px 22px rgba(0,168,132,.32) !important;
        box-sizing:border-box !important;
        white-space:nowrap !important;
        overflow:hidden !important;
        text-align:left !important;
    }

    .topbar { height:76px; background:#FFFFFF; display:flex; justify-content:space-between; align-items:center; padding:0 1.6rem; border:1px solid #E4EAF2; border-radius:20px; box-shadow:0 10px 26px rgba(15,23,42,.06); margin-bottom:1.15rem; }
    .topbar-title { color:#0B3440; font-weight:950; font-size:1.05rem; }
    .topbar-user { display:flex; gap:1rem; align-items:center; color:#526078; font-weight:850; }
    .avatar { width:42px; height:42px; border-radius:50%; background:linear-gradient(135deg,#D6FFF6,#DFE1FF); display:flex; align-items:center; justify-content:center; color:#0B3440; font-weight:950; }
    .home-hero { background:linear-gradient(135deg,#FFFFFF 0%,#F7FAFC 62%,#ECFDF8 100%); border:1px solid #E2E8F0; border-radius:24px; padding:1.65rem 1.8rem; margin:0 0 1.15rem 0; box-shadow:0 14px 34px rgba(15,23,42,.07); }
    .home-hero-title { color:#203047; font-size:2.05rem; line-height:1.1; font-weight:950; letter-spacing:-.03em; margin:0; }
    .kpi { background:#FFFFFF; border-radius:18px; padding:1.35rem; min-height:126px; border:1px solid #E3E8EF; box-shadow:0 12px 28px rgba(15,23,42,.07); position:relative; overflow:hidden; }
    .kpi:before { content:""; position:absolute; left:0; top:0; bottom:0; width:5px; background:var(--c); }
    .kpi-label { color:var(--c); font-size:.78rem; font-weight:950; text-transform:uppercase; } .kpi-value { color:#394356; font-size:1.7rem; font-weight:950; margin-top:.35rem; } .kpi-foot { color:#7C8798; font-size:.8rem; margin-top:.35rem; }
    .panel { background:#FFFFFF; border:1px solid #E0E6EE; border-radius:18px; box-shadow:0 12px 28px rgba(15,23,42,.07); margin-top:1.25rem; overflow:hidden; } .panel-header { padding:1rem 1.25rem; border-bottom:1px solid #E2E8F0; color:#0B3440; font-weight:950; } .panel-body { padding:1.25rem; }
    div[data-testid="stForm"] { background:#FFFFFF !important; border:1px solid #E4EAF2 !important; border-radius:22px !important; padding:1.15rem 1.25rem !important; box-shadow:0 14px 34px rgba(15,23,42,.07) !important; }
    label, .stTextInput label, .stTextArea label, .stSelectbox label, .stNumberInput label, .stDateInput label, .stFileUploader label { color:#344054 !important; font-weight:850 !important; font-size:.82rem !important; }
    div[data-baseweb="input"] > div, div[data-baseweb="select"] > div, textarea, input { background:#F9FAFB !important; border:1px solid #D0D7E2 !important; border-radius:12px !important; color:#1F2937 !important; box-shadow:inset 0 1px 0 rgba(255,255,255,.8) !important; }
    div[data-baseweb="input"] > div:focus-within, div[data-baseweb="select"] > div:focus-within, textarea:focus, input:focus { border-color:#00A884 !important; box-shadow:0 0 0 3px rgba(0,168,132,.16) !important; }
    .login-card { max-width:480px; margin:8vh auto 1rem; background:linear-gradient(135deg,#FFFFFF 0%,#F8FBFC 62%,#ECFDF8 100%); border:1px solid #E4EAF2; border-radius:24px; padding:32px; box-shadow:0 20px 60px rgba(15,23,42,.12); text-align:center; position:relative; overflow:hidden; } .login-title { color:#0B3440; font-size:2rem; font-weight:950; } .login-title:before { content:"◆ "; color:#00A884; }
    div[data-testid="stExpander"], div[data-testid="stDataFrame"] { border-radius:16px!important; overflow:hidden!important; }
    .topbar, .home-hero, .kpi, .panel { box-sizing:border-box !important; max-width:100% !important; }

    /* Panel de color del menú: fondo independiente y estable */
    .left-menu-bg {
        position:fixed;
        top:0;
        left:0;
        width:282px;
        height:100vh;
        background:linear-gradient(180deg,#062C36 0%,#0A4652 68%,#083640 100%);
        border-radius:0 28px 28px 0;
        box-shadow:0 16px 34px rgba(7,49,61,.30);
        z-index:0;
        pointer-events:none;
    }
    .left-menu-marker, .menu-brand, .menu-section, .menu-active {
        position:relative;
        z-index:3;
    }
    div[data-testid="stHorizontalBlock"]:first-of-type > div[data-testid="column"]:first-child {
        position:relative !important;
        z-index:2 !important;
        padding:1.25rem 1rem !important;
        min-height:calc(100vh - 1rem) !important;
        box-sizing:border-box !important;
        overflow:hidden !important;
    }
    div[data-testid="stHorizontalBlock"]:first-of-type > div[data-testid="column"]:first-child * {
        color:#FFFFFF !important;
    }
    div[data-testid="stHorizontalBlock"]:first-of-type > div[data-testid="column"]:first-child .stButton,
    div[data-testid="stHorizontalBlock"]:first-of-type > div[data-testid="column"]:first-child .element-container {
        position:relative !important;
        z-index:3 !important;
    }
    div[data-testid="stHorizontalBlock"]:first-of-type > div[data-testid="column"]:first-child .stButton button {
        width:auto !important;
        max-width:100% !important;
        min-height:44px !important;
        display:inline-flex !important;
        align-items:center !important;
        justify-content:flex-start !important;
        text-align:left !important;
        background:rgba(255,255,255,.18) !important;
        border:1px solid rgba(255,255,255,.34) !important;
        border-radius:12px !important;
        color:#FFFFFF !important;
        font-weight:850 !important;
        margin:.18rem 0 !important;
        padding:.58rem .75rem !important;
        box-shadow:0 7px 16px rgba(0,0,0,.10) !important;
        overflow:hidden !important;
        white-space:nowrap !important;
    }
    div[data-testid="stHorizontalBlock"]:first-of-type > div[data-testid="column"]:first-child .stButton button:hover {
        background:rgba(255,255,255,.28) !important;
        border-color:rgba(255,255,255,.55) !important;
    }
    div[data-testid="stHorizontalBlock"]:first-of-type > div[data-testid="column"]:last-child {
        min-width:0 !important;
        padding-left:.25rem !important;
        box-sizing:border-box !important;
    }

    /* Nuevo registro y tablas: diseño corporativo sin afectar menú lateral */
    .registro-landing-hero{background:radial-gradient(circle at 12% 18%,rgba(0,168,132,.14),transparent 30%),linear-gradient(135deg,#fff 0%,#F7FAFC 55%,#EAFBF7 100%);border:1px solid #DDE7F0;border-radius:28px;padding:1.85rem 2rem;margin:0 0 1.4rem 0;box-shadow:0 18px 46px rgba(15,23,42,.08)}
    .registro-landing-title{color:#062C36;font-size:2.18rem;line-height:1.05;font-weight:950;letter-spacing:-.035em;margin:0 0 .55rem 0}.registro-landing-subtitle{color:#667085;font-size:1rem;font-weight:850;margin:0}
    .registro-card-slot{display:none}div[data-testid="column"]:has(.registro-card-slot) .stButton button{width:100%!important;min-height:145px!important;background:linear-gradient(180deg,#fff 0%,#F8FAFC 100%)!important;border:1px solid #DDE6F0!important;border-radius:26px!important;color:#102A43!important;font-weight:900!important;text-align:left!important;justify-content:flex-start!important;align-items:flex-start!important;padding:1.15rem 1.2rem!important;white-space:pre-line!important;line-height:1.45!important;box-shadow:0 18px 42px rgba(15,23,42,.10)!important;overflow:hidden!important}
    div[data-testid="column"]:has(.registro-card-slot) .stButton button:hover{transform:translateY(-4px)!important;border-color:#00A884!important;background:linear-gradient(180deg,#fff 0%,#ECFDF8 100%)!important;box-shadow:0 24px 54px rgba(15,23,42,.14),0 0 0 4px rgba(0,168,132,.13)!important;color:#062C36!important}
    .registro-full-panel{width:100%;background:radial-gradient(circle at 8% 10%,rgba(0,168,132,.12),transparent 32%),linear-gradient(135deg,#fff 0%,#F8FAFC 58%,#ECFDF8 100%);border:1px solid #DDE7F0;border-radius:30px;padding:1.7rem 1.85rem 1.9rem 1.85rem;box-shadow:0 22px 56px rgba(15,23,42,.11);box-sizing:border-box;overflow:hidden;margin-bottom:1.1rem}.registro-full-title{color:#062C36;font-size:2.05rem;font-weight:950;letter-spacing:-.035em;margin:0 0 .35rem 0}.registro-full-subtitle{color:#667085;font-size:1rem;font-weight:850;margin:0 0 1.2rem 0}.registro-pill{display:inline-flex;align-items:center;border-radius:999px;padding:.38rem .85rem;background:#DCFCE7;color:#027A48;font-size:.8rem;font-weight:950;margin-bottom:1rem}.registro-form-shell{background:#fff;border:1px solid #DDE7F0;border-radius:26px;padding:1.1rem 1.2rem;box-shadow:0 18px 44px rgba(15,23,42,.08);margin-top:.95rem}


    /* Tarjetas grandes de selección en las cuatro secciones principales */
    div[class*="st-key-card_pnc"] button,
    div[class*="st-key-card_me"] button,
    div[class*="st-key-card_ddm"] button,
    div[class*="st-key-card_reclamos"] button,
    div[class*="st-key-card_devoluciones"] button,
    div[class*="st-key-consulta_tarjeta_"] button,
    div[class*="st-key-card_muestras_"] button,
    div[class*="st-key-card_muestras"] button,
    div[class*="st-key-entrega_Nave"] button {
        width:100% !important;
        min-height:150px !important;
        height:auto !important;
        padding:1.2rem 1.25rem !important;
        display:flex !important;
        align-items:flex-start !important;
        justify-content:flex-start !important;
        text-align:left !important;
        white-space:pre-line !important;
        overflow:visible !important;
        text-overflow:clip !important;
        line-height:1.38 !important;
        font-size:.94rem !important;
        font-weight:850 !important;
        color:#203047 !important;
        background:linear-gradient(180deg,#FFFFFF 0%,#F7FAFC 100%) !important;
        border:1px solid #D7E1EC !important;
        border-radius:26px !important;
        box-shadow:0 18px 42px rgba(15,23,42,.10) !important;
        box-sizing:border-box !important;
    }
    div[class*="st-key-card_pnc"] button p,
    div[class*="st-key-card_me"] button p,
    div[class*="st-key-card_ddm"] button p,
    div[class*="st-key-card_reclamos"] button p,
    div[class*="st-key-card_devoluciones"] button p,
    div[class*="st-key-consulta_tarjeta_"] button p,
    div[class*="st-key-card_muestras_"] button p,
    div[class*="st-key-card_muestras"] button p,
    div[class*="st-key-entrega_Nave"] button p {
        white-space:pre-line !important;
        overflow:visible !important;
        text-overflow:clip !important;
        word-break:normal !important;
        overflow-wrap:anywhere !important;
        line-height:1.38 !important;
        margin:0 !important;
    }
    div[class*="st-key-card_pnc"] button:hover,
    div[class*="st-key-card_me"] button:hover,
    div[class*="st-key-card_ddm"] button:hover,
    div[class*="st-key-card_reclamos"] button:hover,
    div[class*="st-key-card_devoluciones"] button:hover,
    div[class*="st-key-consulta_tarjeta_"] button:hover,
    div[class*="st-key-card_muestras_"] button:hover,
    div[class*="st-key-card_muestras"] button:hover,
    div[class*="st-key-entrega_Nave"] button:hover {
        transform:translateY(-4px) !important;
        border-color:#00A884 !important;
        background:linear-gradient(180deg,#FFFFFF 0%,#ECFDF8 100%) !important;
        color:#062C36 !important;
        box-shadow:0 24px 54px rgba(15,23,42,.14),0 0 0 4px rgba(0,168,132,.12) !important;
    }
    /* El resto del contenido se mantiene dentro del ancho disponible */
    .registro-landing-hero, .registro-full-panel, .registro-form-shell,
    .topbar, .home-hero, .panel, div[data-testid="stForm"],
    div[data-testid="stDataFrame"], div[data-testid="stExpander"] {
        width:100% !important;
        max-width:100% !important;
        min-width:0 !important;
        box-sizing:border-box !important;
    }
    div[data-testid="stHorizontalBlock"], div[data-testid="column"] {
        min-width:0 !important;
    }
    div[data-testid="stDataFrame"] { overflow-x:auto !important; }
    .registro-full-panel, .registro-form-shell, .panel-body { overflow-x:auto !important; }
    input, textarea, div[data-baseweb="select"] { max-width:100% !important; box-sizing:border-box !important; }
    @media (max-width:1200px) {
        div[class*="st-key-card_pnc"] button,
        div[class*="st-key-card_me"] button,
        div[class*="st-key-card_ddm"] button,
    div[class*="st-key-card_reclamos"] button,
    div[class*="st-key-card_devoluciones"] button,
        div[class*="st-key-consulta_tarjeta_"] button,
        div[class*="st-key-card_muestras_"] button,
        div[class*="st-key-card_muestras"] button,
        div[class*="st-key-entrega_Nave"] button {
            min-height:140px !important;
            padding:1.05rem 1.1rem !important;
            font-size:.9rem !important;
        }
        .registro-landing-title { font-size:1.85rem !important; }
        .registro-full-title { font-size:1.75rem !important; }
    }
    @media (max-width:850px) {
        .main .block-container { padding:.65rem !important; }
        .registro-landing-hero, .registro-full-panel { padding:1.25rem !important; }
        .topbar { padding:0 1rem !important; }
        .topbar-user { gap:.45rem !important; font-size:.78rem !important; }
        div[class*="st-key-card_pnc"] button,
        div[class*="st-key-card_me"] button,
        div[class*="st-key-card_ddm"] button,
    div[class*="st-key-card_reclamos"] button,
    div[class*="st-key-card_devoluciones"] button,
        div[class*="st-key-consulta_tarjeta_"] button,
        div[class*="st-key-card_muestras_"] button,
        div[class*="st-key-card_muestras"] button,
        div[class*="st-key-entrega_Nave"] button {
            min-height:128px !important;
        }
    }

    /* Paneles generales compactos y separados del menu principal */
    div[data-testid="stHorizontalBlock"]:has(.left-menu-marker) {
        gap:1.45rem !important;
    }
    div[data-testid="stHorizontalBlock"]:has(.left-menu-marker) > div[data-testid="column"]:last-child {
        padding-left:.75rem !important;
        padding-right:.4rem !important;
        min-width:0 !important;
    }
    .registro-landing-hero {
        width:calc(100% - .8rem) !important;
        margin:0 .4rem 1.05rem .4rem !important;
        padding:1.35rem 1.5rem !important;
        border-radius:22px !important;
    }
    .registro-full-panel {
        width:calc(100% - .8rem) !important;
        margin:0 .4rem .95rem .4rem !important;
        padding:1.3rem 1.45rem 1.45rem 1.45rem !important;
        border-radius:24px !important;
    }
    .registro-form-shell {
        padding:.95rem 1rem !important;
        border-radius:20px !important;
    }
    .topbar {
        width:calc(100% - .8rem) !important;
        margin-left:.4rem !important;
        margin-right:.4rem !important;
        padding:0 1.25rem !important;
    }
    .registro-landing-title { font-size:1.9rem !important; }
    .registro-landing-subtitle { font-size:.92rem !important; }
    .registro-full-title { font-size:1.8rem !important; }
    /* Permite que columnas, controles y textos se adapten sin superponerse */
    div[data-testid="column"] { min-width:0 !important; overflow:visible !important; }
    .stButton button, .stDownloadButton button { max-width:100% !important; }
    .stButton button p, .stDownloadButton button p {
        overflow-wrap:anywhere !important;
        word-break:normal !important;
    }
    @media (max-width:1100px) {
        .registro-landing-hero, .registro-full-panel, .topbar {
            width:calc(100% - .4rem) !important;
            margin-left:.2rem !important;
            margin-right:.2rem !important;
        }
        div[data-testid="stHorizontalBlock"]:has(.left-menu-marker) > div[data-testid="column"]:last-child {
            padding-left:.35rem !important;
        }
    }

    /* Jerarquia tipografica de las tarjetas: icono y titulo primero, descripcion justificada */
    div[class*="st-key-card_pnc"] button p,
    div[class*="st-key-card_me"] button p,
    div[class*="st-key-card_ddm"] button p,
    div[class*="st-key-card_reclamos"] button p,
    div[class*="st-key-card_devoluciones"] button p,
    div[class*="st-key-consulta_tarjeta_"] button p,
    div[class*="st-key-card_muestras_"] button p,
    div[class*="st-key-entrega_Nave"] button p {
        width:100% !important;
        text-align:justify !important;
        text-justify:inter-word !important;
        white-space:pre-line !important;
        color:#5F6B7C !important;
        font-size:.88rem !important;
        font-weight:600 !important;
        line-height:1.42 !important;
    }
    div[class*="st-key-card_pnc"] button p::first-line,
    div[class*="st-key-card_me"] button p::first-line,
    div[class*="st-key-card_ddm"] button p::first-line,
    div[class*="st-key-consulta_tarjeta_"] button p::first-line,
    div[class*="st-key-card_muestras_"] button p::first-line,
    div[class*="st-key-entrega_Nave"] button p::first-line {
        color:#102A43 !important;
        font-size:1.08rem !important;
        font-weight:950 !important;
        line-height:1.75 !important;
    }

    /* Niveles visuales reales dentro de todas las tarjetas */
    div[class*="st-key-card_pnc"] button p,
    div[class*="st-key-card_me"] button p,
    div[class*="st-key-card_ddm"] button p,
    div[class*="st-key-card_reclamos"] button p,
    div[class*="st-key-card_devoluciones"] button p,
    div[class*="st-key-consulta_tarjeta_"] button p,
    div[class*="st-key-card_muestras_"] button p,
    div[class*="st-key-entrega_Nave"] button p {
        width:100% !important;
        margin:0 !important;
        color:#667085 !important;
        font-size:.88rem !important;
        font-weight:600 !important;
        line-height:1.5 !important;
        text-align:justify !important;
        text-justify:inter-word !important;
        white-space:normal !important;
        overflow:visible !important;
        text-overflow:clip !important;
    }
    div[class*="st-key-card_pnc"] button strong,
    div[class*="st-key-card_me"] button strong,
    div[class*="st-key-card_ddm"] button strong,
    div[class*="st-key-card_reclamos"] button strong,
    div[class*="st-key-card_devoluciones"] button strong,
    div[class*="st-key-consulta_tarjeta_"] button strong,
    div[class*="st-key-card_muestras_"] button strong,
    div[class*="st-key-entrega_Nave"] button strong {
        color:#102A43 !important;
        font-size:1.12rem !important;
        font-weight:950 !important;
        letter-spacing:-.01em !important;
    }
    div[class*="st-key-card_pnc"] button em,
    div[class*="st-key-card_me"] button em,
    div[class*="st-key-card_ddm"] button em,
    div[class*="st-key-card_reclamos"] button em,
    div[class*="st-key-card_devoluciones"] button em,
    div[class*="st-key-consulta_tarjeta_"] button em,
    div[class*="st-key-card_muestras_"] button em,
    div[class*="st-key-entrega_Nave"] button em {
        display:inline-block !important;
        color:#008C73 !important;
        font-size:.82rem !important;
        font-style:normal !important;
        font-weight:900 !important;
        text-align:left !important;
    }

    /* Ajuste definitivo de posicion superior en versiones nuevas y anteriores de Streamlit */
    [data-testid="stAppViewContainer"], [data-testid="stAppViewContainer"] > .main, .main { padding-top:0 !important; margin-top:0 !important; }
    [data-testid="stMain"], [data-testid="stMainBlockContainer"], .stMainBlockContainer, .main .block-container {
        padding-top:.25rem !important; margin-top:0 !important;
    }
    .topbar { height:60px !important; margin-top:0 !important; margin-bottom:.45rem !important; }
    .home-hero { min-height:68px !important; padding:.8rem 1.2rem !important; margin:0 0 .45rem 0 !important; display:flex; align-items:center; }
    .home-hero-title { font-size:1.7rem !important; white-space:normal !important; }
    .indicator-title { min-height:42px; display:flex; align-items:center; color:#0B3440; font-size:1.02rem; font-weight:950; }
    .toolbar-marker,.chart-marker { display:none; }
    div[data-testid="column"]:has(.toolbar-marker) { display:flex !important; align-items:center !important; justify-content:flex-end !important; }
    div[data-testid="column"]:has(.toolbar-marker) button { border-radius:999px !important; min-height:36px !important; padding:.3rem .72rem !important; font-size:.76rem !important; font-weight:850 !important; border:1px solid #D7E1EC !important; background:#FFF !important; color:#526078 !important; box-shadow:0 5px 14px rgba(15,23,42,.06) !important; }
    div[data-testid="stPopoverBody"] { min-width:360px !important; }
    .kpi { min-height:150px !important; height:100% !important; margin-bottom:.85rem !important; }
    div[data-testid="stHorizontalBlock"]:has(.kpi) { align-items:stretch !important; gap:1rem !important; margin-bottom:.75rem !important; }
    div[data-testid="stHorizontalBlock"]:has(.kpi) > div[data-testid="column"] { display:flex !important; min-width:0 !important; }
    div[data-testid="stHorizontalBlock"]:has(.chart-marker) { margin-top:.75rem !important; padding-top:.25rem !important; }
    </style>
    """, unsafe_allow_html=True)

def init_state():
    if 'auth' not in st.session_state: st.session_state.auth=None
    if 'page' not in st.session_state: st.session_state.page='Inicio'
    if 'inicio_indicador' not in st.session_state: st.session_state.inicio_indicador='PNC'

def login():
    if st.session_state.auth: return st.session_state.auth
    styles(False)
    c1,c2,c3=st.columns([.36,.28,.36])
    with c2:
        st.markdown('<div class="login-card"><div class="login-title">Sistema de Calidad</div></div>',unsafe_allow_html=True)
        with st.form('login_form'):
            u=st.text_input('Usuario'); p=st.text_input('Contraseña',type='password'); ok=st.form_submit_button('Ingresar')
    if ok:
        au=auth_user(u.strip(),p)
        estado=au.get('estado') if au else 'credenciales_invalidas'
        if estado=='ok':
            st.session_state.auth={k:au[k] for k in ('usuario','nombre','rol','requiere_cambio_pass')}
            audit(au['usuario'],'LOGIN','Ingreso correcto')
            st.rerun()
        elif estado=='bloqueado_por_intentos':
            audit(au.get('usuario') or u.strip(),'BLOQUEO_LOGIN','Cuenta bloqueada al alcanzar 8 intentos fallidos')
            st.error('Cuenta bloqueada por alcanzar 8 intentos fallidos. Solicita la reactivación a un desarrollador.')
        elif estado=='bloqueado':
            st.error('La cuenta está bloqueada o inhabilitada. Solicita la reactivación a un desarrollador.')
        elif estado=='advertencia':
            st.warning(f"Advertencia final: intento {au['intentos']} de 8. Quedan {au['restantes']} intento(s) antes del bloqueo.")
        else:
            restantes=au.get('restantes')
            detalle=f" Quedan {restantes} intento(s)." if restantes is not None else ''
            st.error('Usuario o contraseña incorrectos.'+detalle)
    st.stop()

def topbar(user):
    initials=''.join([x[0] for x in user['nombre'].split()[:2]]).upper() or 'AD'
    st.markdown(f'<div class="topbar"><div class="topbar-title">Sistema de Calidad MD</div><div class="topbar-user"><span>🔔</span><span>{user["nombre"].upper()}</span><span class="avatar">{initials}</span></div></div>',unsafe_allow_html=True)

def menu_button(page, full, icon):
    label=full
    if st.session_state.page==page:
        st.markdown(f'<div class="menu-active"><span>{label}</span></div>',unsafe_allow_html=True)
    else:
        if st.button(label,key=f'menu_{page}'):
            st.session_state.page=page
            st.rerun()


def left_menu():
    st.markdown('<div class="left-menu-bg"></div><span class="left-menu-marker"></span>', unsafe_allow_html=True)
    st.markdown('<div class="menu-brand"><span>◆</span><span class="menu-brand-text">CALIDAD MD</span></div><div class="menu-section">MENÚ PRINCIPAL</div>', unsafe_allow_html=True)
    menu_button('Inicio','🏠 Inicio','🏠')
    menu_button('Nuevo registro','📝 Nuevo registro','📝')
    menu_button('Consulta y descarga','📊 Consulta y descarga','📊')
    menu_button('Muestras de retención','🧪 Muestras de retención','🧪')
    menu_button('Entrega de turno','🔄 Entrega de turno','🔄')
    # Todos los usuarios pueden abrir Catálogos. El contenido se limita por rol
    # dentro de page_catalogos, sin exponer administración sensible.
    menu_button('Catálogos','🧩 Catálogos','🧩')
    if is_dev():
        menu_button('Usuarios','👤 Usuarios','👤')
        menu_button('Auditoría','🧾 Auditoría','🧾')
    st.markdown('<div style="height:1rem"></div>',unsafe_allow_html=True)
    if st.button('Cerrar sesión', key='logout_left_menu'):
        audit(st.session_state.auth['usuario'],'LOGOUT','Cierre de sesión')
        st.session_state.auth=None
        st.rerun()


def _selector_indicador_actual(indicador,key):
    opciones=['PNC','Materia extraña','Producto segregado por detector de metales y RX','Reclamos','Devoluciones','SPAC']
    with st.popover('Indicador',use_container_width=True):
        nuevo=st.radio('Seleccionar indicador',opciones,index=opciones.index(indicador),key=key,label_visibility='collapsed')
        if nuevo!=indicador:
            st.session_state.inicio_indicador=nuevo
            st.rerun()

def _grafica_conteo(data,campo,titulo,etiqueta,key,color='#00A884'):
    if data.empty or campo not in data.columns:
        st.info('No hay información disponible para generar esta gráfica.'); return
    serie=data[campo].fillna('').astype(str).str.strip(); serie=serie[serie.ne('')]
    if serie.empty:
        st.info('No hay información disponible para generar esta gráfica.'); return
    g=serie.value_counts().head(15).rename_axis(etiqueta).reset_index(name='Registros')
    orden=g[etiqueta].tolist()
    base=alt.Chart(g).encode(
        x=alt.X('Registros:Q',title='Número de registros',axis=alt.Axis(tickMinStep=1,grid=True,gridColor='#E8EDF3',domain=False,labelColor='#526078',titleColor='#344054')),
        y=alt.Y(f'{etiqueta}:N',title=None,sort=orden,axis=alt.Axis(domain=False,ticks=False,labelColor='#344054',labelLimit=320)),
        tooltip=[alt.Tooltip(f'{etiqueta}:N',title=etiqueta),alt.Tooltip('Registros:Q',title='Registros',format=',d')])
    barras=base.mark_bar(cornerRadiusEnd=7,color=color,size=22)
    textos=base.mark_text(align='left',dx=7,fontWeight='bold',fontSize=11,color='#203047').encode(text=alt.Text('Registros:Q',format=',d'))
    grafica=(barras+textos).properties(title=alt.TitleParams(titulo,anchor='start',fontSize=17,fontWeight=700,color='#0B3440',offset=18),height=max(270,min(520,38*len(g)))).configure_view(stroke=None).configure_axis(labelFontSize=11,titleFontSize=12)
    st.altair_chart(grafica,use_container_width=True,key=key)

def _grafica_mes(data,campo,titulo,key):
    fechas=pd.to_datetime(data[campo],errors='coerce').dropna() if campo in data.columns else pd.Series(dtype='datetime64[ns]')
    if fechas.empty:
        st.info('No hay fechas válidas para generar esta gráfica.'); return
    g=fechas.dt.to_period('M').astype(str).value_counts().sort_index().rename_axis('Mes').reset_index(name='Registros')
    base=alt.Chart(g).encode(
        x=alt.X('Mes:N',sort=None,title=None,axis=alt.Axis(labelAngle=-35,labelColor='#526078',domain=False,ticks=False)),
        y=alt.Y('Registros:Q',title='Número de registros',axis=alt.Axis(tickMinStep=1,grid=True,gridColor='#E8EDF3',domain=False,labelColor='#526078',titleColor='#344054')),
        tooltip=[alt.Tooltip('Mes:N',title='Mes'),alt.Tooltip('Registros:Q',title='Registros',format=',d')])
    area=base.mark_area(color='#5850EC',opacity=.09)
    linea=base.mark_line(point=alt.OverlayMarkDef(filled=True,fill='#FFFFFF',stroke='#5850EC',strokeWidth=3,size=90),strokeWidth=3,color='#5850EC')
    textos=base.mark_text(dy=-14,fontWeight='bold',fontSize=11,color='#203047').encode(text=alt.Text('Registros:Q',format=',d'))
    grafica=(area+linea+textos).properties(title=alt.TitleParams(titulo,anchor='start',fontSize=17,fontWeight=700,color='#0B3440',offset=18),height=315).configure_view(stroke=None).configure_axis(labelFontSize=11,titleFontSize=12)
    st.altair_chart(grafica,use_container_width=True,key=key)

def panel_indicadores_spac_inicio(indicador='SPAC'):
    # Mantiene SPAC seleccionado cuando cambian sus filtros internos.
    st.session_state.inicio_indicador='SPAC'
    registros=read_df("SELECT m.fecha,m.analista,m.total_carga_datos,m.horas_nave1,m.horas_nave2,m.horas_nave3,COALESCE(e.nave,'') AS nave FROM matriz_entrega m LEFT JOIN entregas_turno e ON e.id=m.entrega_id")
    h,sel,fil=st.columns([7.1,1.4,1.5],vertical_alignment='center')
    with h:
        st.markdown('<div class="indicator-title">Indicadores SPAC</div>',unsafe_allow_html=True)
    with sel:
        st.markdown('<span class="toolbar-marker"></span>',unsafe_allow_html=True)
        _selector_indicador_actual(indicador,'selector_spac')
    if registros.empty:
        with fil:
            st.markdown('<span class="toolbar-marker"></span>',unsafe_allow_html=True)
            st.button('Filtros',disabled=True,key='filtro_spac_vacio')
        st.info('Aún no hay información SPAC disponible.')
        return
    registros['fecha_dt']=pd.to_datetime(registros['fecha'],errors='coerce').dt.date
    registros=registros.dropna(subset=['fecha_dt'])
    registros=registros[registros['fecha_dt']<=date.today()]
    if registros.empty:
        st.info('La matriz no contiene información válida hasta la fecha actual.')
        return

    with fil:
        st.markdown('<span class="toolbar-marker"></span>',unsafe_allow_html=True)
        with st.popover('Filtros',use_container_width=True):
            tipo=st.radio(
                'Visualización',
                ['SEMANA','MES'],
                horizontal=True,
                format_func=lambda x:'Semana' if x=='SEMANA' else 'Mes',
                key='inicio_spac_tipo'
            )
            grafica_base=datos_grafica_cumplimiento('ANALISTA',tipo)
            nombres=sorted(grafica_base['Entidad'].dropna().astype(str).unique().tolist()) if not grafica_base.empty else []
            seleccionar_todos=st.checkbox(
                'Seleccionar todos los analistas',
                value=True,
                key='inicio_spac_todos_analistas'
            )
            if seleccionar_todos:
                analistas=nombres
                st.caption(f'{len(nombres)} analistas seleccionados.')
            else:
                analistas=st.multiselect(
                    'Selección de analistas',
                    nombres,
                    default=[],
                    key='inicio_spac_analistas'
                )

    ga=datos_grafica_cumplimiento('ANALISTA',tipo)
    gc=datos_grafica_cumplimiento('CALIDAD_NAVE',tipo)
    gp=datos_grafica_cumplimiento('NAVE',tipo)

    # Presenta todo el histórico disponible hasta el periodo actual.
    periodo_actual=clave_periodo(date.today(),tipo)
    if not ga.empty:
        ga=ga[ga['Periodo'].astype(str)<=periodo_actual]
        ga=ga[ga['Entidad'].astype(str).isin(analistas)] if analistas else ga.iloc[0:0]
    if not gc.empty:
        gc=gc[gc['Periodo'].astype(str)<=periodo_actual]
    if not gp.empty:
        gp=gp[gp['Periodo'].astype(str)<=periodo_actual]

    st.markdown('### Cumplimiento SPAC por analista')
    grafica_lineas_con_valores(ga,'Cumplimiento SPAC por analista','inicio_chart_analistas')
    c1,c2=st.columns(2)
    with c1:
        grafica_lineas_con_valores(gc,'Indicador SPAC Calidad','inicio_chart_calidad')
    with c2:
        grafica_lineas_con_valores(gp,'Indicador SPAC Producción','inicio_chart_produccion')

def _panel_registros_inicio(tipo):
    cfg={
        'PNC':('pnc_registros','fecha_apertura','defecto','Producto No Conforme','analista','linea_sector','status','cantidad_total_pnc'),
        'Materia extraña':('me_registros','_fecha','_codigo_defecto_panel','Materia extraña','analista_detecta','linea_sector',None,None),
        'Producto segregado por detector de metales y RX':('ddm_rx_registros','_fecha','_codigo_defecto_panel','Producto segregado por detector de metales y RX','analista_detecta','linea_sector',None,None),
        'Reclamos':('reclamos_registros','fecha','descripcion_defecto','Reclamos','creado_por','sector','estado_reclamo','cantidad_afectada'),
        'Devoluciones':('devoluciones_registros','fecha','defecto','Devoluciones','creado_por','sector','status','cantidad_afectada')}
    if tipo not in cfg:
        st.session_state.inicio_indicador='PNC'; st.rerun()
    tabla,fecha,defecto,titulo,c_analista,c_linea,c_estado,c_cantidad=cfg[tipo]
    d=read_df(f'SELECT * FROM {tabla}')
    if not d.empty and 'nave' in d.columns:
        d['nave']=d['nave'].map(normalizar_nave)
    if fecha=='_fecha':
        d[fecha]=pd.to_datetime(dict(year=pd.to_numeric(d.get('anio'),errors='coerce'),month=pd.to_numeric(d.get('mes'),errors='coerce'),day=pd.to_numeric(d.get('dia'),errors='coerce')),errors='coerce') if not d.empty else pd.Series(dtype='datetime64[ns]')
    else: d[fecha]=pd.to_datetime(d[fecha],errors='coerce') if fecha in d.columns else pd.NaT
    if defecto=='_codigo_defecto_panel':
        cat=read_df('SELECT codigo,defecto FROM defectos')
        mapa={str(r.codigo).strip():str(r.defecto).strip() for r in cat.itertuples()}
        cod=d.get('codigo_defecto',pd.Series('',index=d.index)).fillna('').astype(str).str.strip()
        d[defecto]=cod.map(lambda x:f'{x} | {mapa.get(x,"Defecto no catalogado")}' if x else '')
    fechas=d[fecha].dropna()
    limpio=lambda campo:sorted(d.get(campo,pd.Series(dtype=str)).dropna().astype(str).loc[lambda z:z.str.strip().ne('')].unique())
    lineas,naves,analistas=limpio(c_linea),limpio('nave'),limpio(c_analista)
    h,sel,fil=st.columns([7.1,1.4,1.5],vertical_alignment='center')
    with h: st.markdown(f'<div class="indicator-title">Indicadores de {titulo}</div>',unsafe_allow_html=True)
    with sel: st.markdown('<span class="toolbar-marker"></span>',unsafe_allow_html=True); _selector_indicador_actual(tipo,f'selector_{tabla}')
    with fil:
        st.markdown('<span class="toolbar-marker"></span>',unsafe_allow_html=True)
        with st.popover('Filtros',use_container_width=True):
            if not fechas.empty:
                mn,mx=fechas.min().date(),fechas.max().date(); a,b=st.columns(2)
                desde=a.date_input('Fecha inicial',mn,min_value=mn,max_value=mx,key=f'desde_{tabla}'); hasta=b.date_input('Fecha final',mx,min_value=mn,max_value=mx,key=f'hasta_{tabla}')
            else: desde=hasta=None
            sl=st.multiselect('Línea/Sector',lineas,key=f'lineas_{tabla}'); sn=st.multiselect('Nave',naves,key=f'naves_{tabla}'); sa=st.multiselect('Analista / usuario',analistas,key=f'analistas_{tabla}')
    x=d.copy()
    if desde and hasta:
        if desde>hasta: st.error('La fecha inicial no puede ser posterior a la fecha final.'); return
        x=x[(x[fecha].dt.date>=desde)&(x[fecha].dt.date<=hasta)]
    if sl:x=x[x[c_linea].isin(sl)]
    if sn:x=x[x.nave.isin(sn)]
    if sa:x=x[x[c_analista].isin(sa)]
    total=len(x); ln=x.get(c_linea,pd.Series(dtype=str)).replace('',pd.NA).nunique(); dn=x.get(defecto,pd.Series(dtype=str)).replace('',pd.NA).nunique(); meses=x[fecha].dropna().dt.to_period('M').nunique()
    if tipo=='PNC':
        e=x.get(c_estado,pd.Series('',index=x.index)).fillna('').astype(str).str.strip().str.upper(); ce=int(e.eq('CERRADO').sum()); ab=int(e.eq('ABIERTO').sum()); cantidad=pd.to_numeric(x.get(c_cantidad,pd.Series(0,index=x.index)),errors='coerce').fillna(0).sum()
        cards=[('Total de PNC',f'{total:,}','Número de registros totales','#00A884'),('% de cierre',f'{(ce/total*100 if total else 0):.1f}%','PNC cerrados respecto al total','#5850EC'),('Kg totales de PNC',f'{cantidad:,.2f} kg','Cantidad total registrada','#3F7BFF'),('PNC abiertos',f'{ab:,}','Registros pendientes de cierre','#F59E0B')]
    elif tipo in {'Reclamos','Devoluciones'}:
        e=x.get(c_estado,pd.Series('',index=x.index)).fillna('').astype(str).str.strip().str.upper(); ce=int(e.eq('CERRADO').sum()); ab=int(e.eq('ABIERTO').sum()); cantidad=pd.to_numeric(x.get(c_cantidad,pd.Series(0,index=x.index)),errors='coerce').fillna(0).sum(); singular='Reclamos' if tipo=='Reclamos' else 'Devoluciones'
        cards=[(f'Total de {tipo.lower()}',f'{total:,}','Número de registros totales','#00A884'),('% de cierre',f'{(ce/total*100 if total else 0):.1f}%',f'{singular} cerrados respecto al total','#5850EC'),('Cantidad afectada',f'{cantidad:,.2f}','Suma de la cantidad registrada','#3F7BFF'),('Registros abiertos',f'{ab:,}','Pendientes de cierre','#F59E0B')]
    else: cards=[('Registros',total,'Total filtrado','#00A884'),('Códigos / defectos',dn,'Tipos identificados','#5850EC'),('Líneas / sectores',ln,'Con registros','#3F7BFF'),('Meses',meses,'Periodos con actividad','#F59E0B')]
    cols=st.columns(4,gap='large')
    for c,(la,va,pi,co) in zip(cols,cards):
        with c: st.markdown(f'<div class="kpi" style="--c:{co}"><div class="kpi-label">{la}</div><div class="kpi-value">{va}</div><div class="kpi-foot">{pi}</div></div>',unsafe_allow_html=True)
    a,b=st.columns(2)
    with a:
        st.markdown('<span class="chart-marker"></span>',unsafe_allow_html=True)
        especial=tipo in {'Materia extraña','Producto segregado por detector de metales y RX'}
        _grafica_conteo(x,defecto,'Distribución por código y defecto' if especial else 'Distribución de defectos','Código / Defecto' if especial else 'Defecto',f'def_{tabla}','#00A884')
    with b: _grafica_conteo(x,c_linea,'Registros por línea o sector','Línea/Sector',f'lin_{tabla}','#3F7BFF')
    _grafica_mes(x,fecha,'Evolución mensual de registros',f'mes_{tabla}')

def page_inicio():
    st.markdown('<div class="home-hero"><div class="home-hero-title">Panel Calidad Mundo Dulce</div></div>',unsafe_allow_html=True)
    indicador=st.session_state.get('inicio_indicador','PNC')
    panel_indicadores_spac_inicio(indicador) if indicador=='SPAC' else _panel_registros_inicio(indicador)

def page_registro():
    if 'registro_tipo' not in st.session_state:
        st.session_state.registro_tipo=None
    if 'form_nonce' not in st.session_state:
        st.session_state.form_nonce=0
    flash=st.session_state.pop('flash_registro_guardado',None)
    if flash:
        st.success(flash)
    def limpiar_form():
        st.session_state.form_nonce += 1
    def volver_selector():
        st.session_state.registro_tipo=None
        st.rerun()
    def selector():
        st.markdown("""<div class="registro-landing-hero"><div class="registro-landing-title">Nuevo registro</div><div class="registro-landing-subtitle">Selecciona el tipo de registro que deseas capturar.</div></div>""",unsafe_allow_html=True)
        tarjetas=[
            ('PNC','📝  **PNC´s**\n\nCaptura y seguimiento de producto no conforme.\n\n*Abrir registro*','card_pnc'),
            ('ME','🧲  **Materia Extraña**\n\nRegistro de hallazgos y acciones de contención.\n\n*Abrir registro*','card_me'),
            ('DDM_RX','📦  **Detector de metales y RX**\n\nControl de producto segregado por detección.\n\n*Abrir registro*','card_ddm'),
            ('RECLAMOS','📣  **Reclamos**\n\nRegistro, investigación y seguimiento de reclamos.\n\n*Abrir registro*','card_reclamos'),
            ('DEVOLUCIONES','↩️  **Devoluciones**\n\nRegistro y seguimiento de producto devuelto.\n\n*Abrir registro*','card_devoluciones')
        ]
        # Misma estructura visual de Muestras de retención: hasta tres tarjetas
        # por fila, con ancho y separación uniformes.
        for inicio in range(0,len(tarjetas),3):
            lote=tarjetas[inicio:inicio+3]
            if len(lote)==2:
                columnas=st.columns([.5,1,1,.5],gap='large')[1:3]
            else:
                columnas=st.columns(3,gap='large')
            for columna,(valor,texto,clave) in zip(columnas,lote):
                with columna:
                    st.markdown('<span class="registro-card-slot"></span>',unsafe_allow_html=True)
                    if st.button(texto,key=clave):
                        st.session_state.registro_tipo=valor
                        st.rerun()

    def form_hallazgo(tabla,titulo,audit_action):
        nonce=st.session_state.form_nonce
        st.markdown(f"""<div class="registro-full-panel"><div class="registro-pill">Nuevo registro</div><div class="registro-full-title">{titulo}</div><div class="registro-full-subtitle">Los campos marcados con * son obligatorios.</div>""",unsafe_allow_html=True)
        if st.button('← Cambiar tipo de registro',key=f'volver_{tabla}'):
            volver_selector()
        st.markdown('<div class="registro-form-shell">',unsafe_allow_html=True)
        prod=read_df('SELECT * FROM productos WHERE activo=1 ORDER BY descripcion')
        defs=read_df('SELECT * FROM defectos WHERE activo=1 ORDER BY CAST(codigo AS INTEGER)')
        opt_prod=['']+[f'{r.item} | {r.descripcion}' for r in prod.itertuples()]
        opt_defs=['']+[f'{r.codigo} | {r.defecto}' for r in defs.itertuples()]
        with st.container():
            a,b,c=st.columns(3)
            fecha=a.date_input('Fecha *',value=date.today(),key=f'fecha_{tabla}_{nonce}')
            semana=b.number_input('Semana *',min_value=1,max_value=53,value=int(date.today().isocalendar().week),step=1,key=f'semana_{tabla}_{nonce}')
            nave=c.selectbox('Nave *',opt_blank(catalog('nave')),key=f'nave_{tabla}_{nonce}')
            a,b,c=st.columns(3)
            opt=a.selectbox('ITEM *',opt_prod,key=f'item_{tabla}_{nonce}')
            item=opt.split('|')[0].strip() if opt else ''
            row=prod[prod['item']==item].iloc[0] if item and item in prod['item'].values else None
            producto=str(row['descripcion']) if row is not None else ''
            cliente=str(row['cliente']) if row is not None else ''
            familia=str(row['familia']) if row is not None else ''
            lote=b.text_input('Lote *',key=f'lote_{tabla}_{nonce}')
            c.empty()
            a,b,c=st.columns(3)
            a.text_input('Descripción',value=producto,disabled=True)
            b.text_input('Cliente',value=cliente,disabled=True)
            c.text_input('Familia',value=familia,disabled=True)
            linea_sector=st.selectbox('Línea/Sector *',opt_blank(catalog('linea_sector')),key=f'linea_{tabla}_{nonce}')
            a,b,c=st.columns(3)
            optd=a.selectbox('Código / Defecto *',opt_defs,key=f'codigo_{tabla}_{nonce}')
            codigo=optd.split('|')[0].strip() if optd else ''
            dr=defs[defs['codigo']==codigo].iloc[0] if codigo and codigo in defs['codigo'].values else None
            defecto=str(dr['defecto']) if dr is not None else ''
            tipo_defecto=str(dr['tipo_defecto']) if dr is not None else ''
            categoria=str(dr['clasificacion']) if dr is not None else ''
            b.text_input('Defecto',value=defecto,disabled=True)
            c.text_input('Tipo de defecto',value=tipo_defecto,disabled=True)
            x1,x2=st.columns(2)
            x1.text_input('Clasificación *',value=categoria,disabled=True)
            turno=x2.selectbox('Turno *',opt_blank(catalog('turno')),key=f'turno_{tabla}_{nonce}')
            a,b,c=st.columns(3)
            supervisor=a.selectbox('Supervisor (Responsable) *',opt_blank(catalog('supervisor')),key=f'sup_{tabla}_{nonce}')
            analista=b.selectbox('Analista (Persona que detecta) *',opt_blank(catalog('analista')),key=f'ana_{tabla}_{nonce}')
            c.empty()
            descripcion=st.text_area('Descripción del defecto *',key=f'desc_{tabla}_{nonce}')
            acciones=st.text_area('Acciones inmediatas *',key=f'accion_{tabla}_{nonce}')
            a,b,c=st.columns(3)
            a.empty(); b.empty(); c.empty()
            a,b,c=st.columns(3)
            equipo=a.text_input('Equipo en donde se tiene el hallazgo',key=f'equipo_{tabla}_{nonce}')
            tipo=b.selectbox('Tipo',opt_blank(['Metal','Plástico duro','Plástico blando','Vidrio','Madera','Papel/Cartón','Cabello','Insecto','Otro']),key=f'tipo_{tabla}_{nonce}')
            particulas=c.number_input('# de partículas halladas',min_value=0,step=1,key=f'part_{tabla}_{nonce}')
            sector_hallazgo=''; ubicacion_material=''
            if tabla=='ddm_rx_registros':
                d1,d2=st.columns(2)
                sector_hallazgo=d1.selectbox('Sector del hallazgo', ['','CONFORMADO','ENVOLTURA','EMPAQUE'], key=f'sector_hallazgo_{tabla}_{nonce}')
                ubicacion_material=d2.selectbox('Ubicación del material', ['','MASA','RELLENO','RECUBIERTO'], key=f'ubicacion_material_{tabla}_{nonce}')
            investigacion=st.text_area('Investigación del origen',key=f'inv_{tabla}_{nonce}')
            evitar=st.text_area('Acciones a realizar para evitar la incidencia',key=f'evitar_{tabla}_{nonce}')
            ok=st.button('Guardar registro',key=f'guardar_{tabla}_{nonce}',type='primary')
        if ok:
            obligatorios={'Línea/Sector':linea_sector,'Nave':nave,'ITEM':item,'Descripción':producto,'Cliente':cliente,'Familia':familia,'Lote':lote,'Código':codigo,'Defecto':defecto,'Tipo de defecto':tipo_defecto,'Semana':semana,'Turno':turno,'Fecha':fecha,'Supervisor':supervisor,'Analista':analista,'Descripción del hallazgo':descripcion,'Acciones inmediatas':acciones,'Categoría inicial':categoria}
            faltantes=[k for k,v in obligatorios.items() if v is None or (isinstance(v,str) and not v.strip())]
            if faltantes:
                st.error('Completa los siguientes campos obligatorios: '+', '.join(faltantes)+'.')
            else:
                rid=exec_sql(f'INSERT INTO {tabla}(dia,mes,anio,nave,linea_sector,familia,equipo_hallazgo,item,producto,lote,descripcion_hallazgo,tipo,particulas_halladas,accion_contingente,investigacion_origen,analista_detecta,supervisor_responsable,acciones_evitar_incidencia,creado_por,creado_en,etapa,codigo_defecto,semana,turno,responsable_detecta,acciones_inmediatas,disposicion,cantidad_observada,status,categoria_inicial,sector_hallazgo,ubicacion_material) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(fecha.day,fecha.month,fecha.year,nave,linea_sector,familia,equipo,item,producto,lote,descripcion,tipo,int(particulas),acciones,investigacion,analista,supervisor,evitar,st.session_state.auth['usuario'],now_iso(),'',codigo,int(semana),turno,'',acciones,'',0.0,'',categoria,sector_hallazgo,ubicacion_material))
                audit(st.session_state.auth['usuario'],audit_action,f'ID {rid}')
                limpiar_form(); st.session_state.flash_registro_guardado=f'Registro guardado correctamente: Número {rid}'
                st.rerun()
        st.markdown('</div></div>',unsafe_allow_html=True)
    def form_devoluciones():
        nonce=st.session_state.form_nonce
        st.markdown('<div class="registro-full-panel"><div class="registro-pill">Nuevo registro</div><div class="registro-full-title">↩️ Devoluciones</div><div class="registro-full-subtitle">Los campos marcados con * son obligatorios. Producto y cliente se completan desde el catálogo vigente.</div>',unsafe_allow_html=True)
        if st.button('← Cambiar tipo de registro',key='volver_devoluciones'): volver_selector()
        st.markdown('<div class="registro-form-shell">',unsafe_allow_html=True)
        prod=read_df('SELECT * FROM productos WHERE activo=1 ORDER BY descripcion')
        defs=read_df('SELECT * FROM defectos WHERE activo=1 ORDER BY CAST(codigo AS INTEGER)')
        opt_prod=['']+[f'{r.item} | {r.descripcion}' for r in prod.itertuples()]
        opt_defs=['']+[f'{r.codigo} | {r.defecto}' for r in defs.itertuples()]
        a,b,c=st.columns(3)
        fecha=a.date_input('Fecha *',date.today(),key=f'dev_fecha_{nonce}')
        od=b.selectbox('Código / Defecto *',opt_defs,key=f'dev_def_{nonce}')
        cedis=c.text_input('Cedis *',key=f'dev_cedis_{nonce}')
        codigo=od.split('|')[0].strip() if od else ''
        dr=defs[defs.codigo.astype(str)==codigo].iloc[0] if codigo and codigo in defs.codigo.astype(str).values else None
        defecto='' if dr is None else str(dr.defecto)
        a,b,c=st.columns(3)
        pais_estado=a.text_input('País / Estado *',key=f'dev_pais_{nonce}')
        op=b.selectbox('ITEM *',opt_prod,key=f'dev_item_{nonce}')
        item=op.split('|')[0].strip() if op else ''
        pr=prod[prod.item.astype(str)==item].iloc[0] if item and item in prod.item.astype(str).values else None
        producto='' if pr is None or pd.isna(pr.descripcion) else str(pr.descripcion).strip()
        cliente='' if pr is None or pd.isna(pr.cliente) else str(pr.cliente).strip()
        familia='' if pr is None or pd.isna(pr.familia) else str(pr.familia).strip()
        c.text_input('Producto',producto,disabled=True,key=f'dev_producto_{nonce}_{item or "sin_item"}')
        a,b,c=st.columns(3)
        a.text_input('Cliente',cliente,disabled=True,key=f'dev_cliente_{nonce}_{item or "sin_item"}')
        b.text_input('Familia',value=familia,disabled=True,key=f'dev_familia_{nonce}_{item or "sin_item"}')
        lote=c.text_input('Lote *',key=f'dev_lote_{nonce}')
        a,b,c=st.columns(3)
        caducidad=a.text_input('Caducidad *',key=f'dev_cad_{nonce}')
        comprobado=b.selectbox('Comprobado *',['','Sí','No'],key=f'dev_comp_{nonce}')
        cantidad=c.number_input('Cant. afectada *',min_value=0.0,step=1.0,format='%.2f',key=f'dev_cant_{nonce}')
        descripcion=st.text_area('Descripción de la devolución *',key=f'dev_desc_{nonce}')
        a,b,c=st.columns(3)
        unidad=a.selectbox('Unidad *',['','Bulto','Bolsa','Pieza','Tarima'],key=f'dev_unidad_{nonce}')
        sector=b.text_input('Sector *',key=f'dev_sector_{nonce}')
        disposicion=c.selectbox('Disposición *',opt_blank(catalog('disposicion')),key=f'dev_disp_{nonce}')
        a,b,c=st.columns(3)
        tsp=a.text_input('TSP N°',key=f'dev_tsp_{nonce}')
        status=b.selectbox('Status *',['','Cerrado','Abierto'],key=f'dev_status_{nonce}')
        nave=c.selectbox('Nave *',opt_blank(catalog('nave')),key=f'dev_nave_{nonce}')
        if st.button('Guardar registro',key=f'guardar_dev_{nonce}',type='primary'):
            req={'Fecha':fecha,'Código / Defecto':codigo,'Cedis':cedis,'País / Estado':pais_estado,'ITEM':item,'Producto':producto,'Cliente':cliente,'Familia':familia,'Lote':lote,'Caducidad':caducidad,'Descripción del defecto':descripcion,'Comprobado':comprobado,'Cant. afectada':cantidad,'Unidad':unidad,'Sector':sector,'Disposición':disposicion,'Status':status,'Nave':nave}
            faltan=[k for k,v in req.items() if v is None or (isinstance(v,str) and not v.strip()) or (k=='Cant. afectada' and float(v)<=0)]
            if faltan: st.error('Completa los siguientes campos obligatorios: '+', '.join(faltan)+'.')
            else:
                rid=exec_sql('INSERT INTO devoluciones_registros(fecha,codigo_defecto,defecto,cedis,pais_estado,item,producto,cliente,linea,familia,lote,caducidad,descripcion_defecto,comprobado,cantidad_afectada,unidad,sector,disposicion,tsp_numero,status,nave,creado_por,creado_en) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(fecha.isoformat(),codigo,defecto,cedis.strip(),pais_estado.strip(),item,producto,cliente,'',familia,lote.strip(),caducidad.strip(),descripcion.strip(),comprobado,float(cantidad),unidad,sector.strip(),disposicion,tsp.strip(),status,nave,st.session_state.auth['usuario'],now_iso()))
                audit(st.session_state.auth['usuario'],'CREAR_DEVOLUCION',f'ID {rid}'); limpiar_form(); st.session_state.flash_registro_guardado=f'Devolución guardada correctamente: Número {rid}'; st.rerun()
        st.markdown('</div></div>',unsafe_allow_html=True)

    def form_reclamos():
        nonce=st.session_state.form_nonce
        st.markdown('<div class="registro-full-panel"><div class="registro-pill">Nuevo registro</div><div class="registro-full-title">📣 Reclamos</div><div class="registro-full-subtitle">Los campos marcados con * son obligatorios. Producto, cliente, familia y datos del defecto se completan desde los catálogos vigentes.</div>',unsafe_allow_html=True)
        if st.button('← Cambiar tipo de registro',key='volver_reclamos'): volver_selector()
        st.markdown('<div class="registro-form-shell">',unsafe_allow_html=True)
        prod=read_df('SELECT * FROM productos WHERE activo=1 ORDER BY descripcion')
        defs=read_df('SELECT * FROM defectos WHERE activo=1 ORDER BY CAST(codigo AS INTEGER)')
        opt_prod=['']+[f'{r.item} | {r.descripcion}' for r in prod.itertuples()]
        opt_defs=['']+[f'{r.codigo} | {r.defecto}' for r in defs.itertuples()]
        a,b,c=st.columns(3)
        fecha=a.date_input('Fecha *',date.today(),key=f'recl_fecha_{nonce}')
        od=b.selectbox('Código / Defecto *',opt_defs,key=f'recl_def_{nonce}')
        codigo=od.split('|')[0].strip() if od else ''
        dr=defs[defs.codigo.astype(str)==codigo].iloc[0] if codigo and codigo in defs.codigo.astype(str).values else None
        defecto=str(dr.defecto) if dr is not None else ''; tipo_defecto=str(dr.tipo_defecto) if dr is not None else ''; clasificacion=str(dr.clasificacion) if dr is not None else ''
        fuente=c.text_input('Fuente *',key=f'recl_fuente_{nonce}')
        a,b,c=st.columns(3); a.text_input('Descripción del defecto',defecto,disabled=True); b.text_input('Tipo de defecto',tipo_defecto,disabled=True); c.text_input('Clasificación del defecto',clasificacion,disabled=True)
        a,b,c=st.columns(3); mercado=a.selectbox('Mercado *',['','Interno','Externo'],key=f'recl_mercado_{nonce}'); pais_estado=b.text_input('País / Estado *',key=f'recl_pais_{nonce}'); numero_caso=c.text_input('No. Caso Right Now / MDLZ *',key=f'recl_caso_{nonce}')
        a,b,c=st.columns(3)
        op=a.selectbox('ITEM *',opt_prod,key=f'recl_item_{nonce}'); item=op.split('|')[0].strip() if op else ''
        pr=prod[prod.item.astype(str)==item].iloc[0] if item and item in prod.item.astype(str).values else None
        producto=str(pr.descripcion) if pr is not None else ''; cliente=str(pr.cliente) if pr is not None else ''; familia='' if pr is None or pd.isna(pr.familia) else str(pr.familia).strip()
        b.text_input('Producto',producto,disabled=True); c.text_input('Cliente',cliente,disabled=True)
        st.text_input('Familia',value=familia,disabled=True,key=f'recl_familia_{nonce}_{item or "sin_item"}')
        descripcion_reclamo=st.text_area('Descripción del reclamo *',key=f'recl_desc_{nonce}')
        causa_raiz=st.text_area('Causa raíz del defecto *',key=f'recl_causa_{nonce}')
        acciones=st.text_area('Acciones correctivas y contingentes *',key=f'recl_acciones_{nonce}')
        a,b,c=st.columns(3); comprobado=a.selectbox('Comprobado *',['','Sí','No'],key=f'recl_comp_{nonce}'); cantidad=b.number_input('Cantidad afectada *',0.0,step=1.0,format='%.2f',key=f'recl_cant_{nonce}'); unidad=c.selectbox('Unidad *',['','Bulto','Bolsa','Pieza','Lámina'],key=f'recl_unidad_{nonce}')
        a,b,c=st.columns(3); sector=a.text_input('Sector *',key=f'recl_sector_{nonce}'); estado=b.selectbox('Estado del reclamo *',['','Abierto','Cerrado'],key=f'recl_estado_{nonce}'); fecha_cierre=c.date_input('Fecha de cierre',value=None,key=f'recl_cierre_{nonce}')
        a,b,c=st.columns(3); tsp=a.text_input('TSP N°',key=f'recl_tsp_{nonce}'); caducidad=b.text_input('Caducidad',key=f'recl_cad_{nonce}',placeholder='Llenado libre'); lote=c.text_input('Lote *',key=f'recl_lote_{nonce}')
        a,b,c=st.columns(3); nave=a.selectbox('Nave *',opt_blank(catalog('nave')),key=f'recl_nave_{nonce}'); pmd=b.selectbox('P / M / D *',['','1 - Manipulación','2 - Producción','3 - Diseño'],key=f'recl_pmd_{nonce}'); red_social=c.selectbox('Red social *',['','Mail','Web','Instagram','No indica','Facebook'],key=f'recl_red_{nonce}')
        if st.button('Guardar registro',key=f'guardar_reclamos_{nonce}',type='primary'):
            req={'Código / Defecto':codigo,'Descripción del defecto':defecto,'Fuente':fuente,'Mercado':mercado,'País / Estado':pais_estado,'No. Caso Right Now / MDLZ':numero_caso,'ITEM':item,'Producto':producto,'Cliente':cliente,'Familia':familia,'Descripción del reclamo':descripcion_reclamo,'Causa raíz':causa_raiz,'Acciones correctivas y contingentes':acciones,'Comprobado':comprobado,'Cantidad afectada':cantidad,'Unidad':unidad,'Sector':sector,'Estado':estado,'Lote':lote,'Nave':nave,'P / M / D':pmd,'Tipo de defecto':tipo_defecto,'Clasificación':clasificacion,'Red social':red_social}
            faltan=[k for k,v in req.items() if v is None or (isinstance(v,str) and not v.strip()) or (k=='Cantidad afectada' and float(v)<=0)]
            duplicado=not read_df('SELECT id FROM reclamos_registros WHERE UPPER(TRIM(numero_caso))=UPPER(TRIM(?))',(numero_caso,)).empty if numero_caso.strip() else False
            if duplicado: st.error('Registro repetido. Ya existe un reclamo con el mismo No. Caso Right Now / MDLZ. Corrige el número de caso o edita el registro existente.')
            elif faltan: st.error('Completa los siguientes campos obligatorios: '+', '.join(dict.fromkeys(faltan))+'.')
            elif estado=='Cerrado' and not fecha_cierre: st.error('Captura la fecha de cierre para el reclamo cerrado.')
            elif fecha_cierre and fecha_cierre<fecha: st.error('La fecha de cierre no puede ser anterior a la fecha del reclamo.')
            else:
                try:
                    rid=exec_sql('INSERT INTO reclamos_registros(fecha,codigo_defecto,descripcion_defecto,fuente,mercado,pais_estado,numero_caso,item,producto,cliente,familia,descripcion_reclamo,causa_raiz,acciones_correctivas_contingentes,comprobado,cantidad_afectada,unidad,sector,estado_reclamo,fecha_cierre,tsp_numero,caducidad,lote,nave,pmd,tipo_defecto,clasificacion_defecto,red_social,creado_por,creado_en) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(fecha.isoformat(),codigo,defecto,fuente.strip(),mercado.strip(),pais_estado.strip(),numero_caso.strip(),item,producto,cliente,familia,descripcion_reclamo.strip(),causa_raiz.strip(),acciones.strip(),comprobado,float(cantidad),unidad,sector.strip(),estado,fecha_cierre.isoformat() if fecha_cierre else None,tsp.strip(),caducidad.strip(),lote.strip(),nave,pmd,tipo_defecto,clasificacion,red_social,st.session_state.auth['usuario'],now_iso()))
                    audit(st.session_state.auth['usuario'],'CREAR_RECLAMO',f'ID {rid} | Caso {numero_caso.strip()}'); limpiar_form(); st.session_state.flash_registro_guardado=f'Reclamo guardado correctamente: Número {rid}'; st.rerun()
                except DatabaseIntegrityError:
                    st.error('No fue posible guardar el reclamo porque existe un registro incompatible o repetido. Actualiza la página y verifica el número de caso.')
        st.markdown('</div></div>',unsafe_allow_html=True)

    def form_pnc():
        nonce=st.session_state.form_nonce
        st.markdown("""<div class="registro-full-panel"><div class="registro-pill">Nuevo registro</div><div class="registro-full-title">📝 PNC´s</div><div class="registro-full-subtitle">Los campos marcados con * son obligatorios.</div>""",unsafe_allow_html=True)
        if st.button('← Cambiar tipo de registro',key='volver_pnc'): volver_selector()
        st.markdown('<div class="registro-form-shell">',unsafe_allow_html=True)
        prod=read_df('SELECT * FROM productos WHERE activo=1 ORDER BY descripcion'); defs=read_df('SELECT * FROM defectos WHERE activo=1 ORDER BY CAST(codigo AS INTEGER)')
        opt_prod=['']+[f'{r.item} | {r.descripcion}' for r in prod.itertuples()]; opt_defs=['']+[f'{r.codigo} | {r.defecto}' for r in defs.itertuples()]
        with st.container():
            a,b,c=st.columns(3)
            fecha=a.date_input('Fecha *',value=date.today(),key=f'pnc_fecha_{nonce}')
            semana=b.number_input('Semana *',min_value=1,max_value=53,value=int(date.today().isocalendar().week),step=1,key=f'pnc_semana_{nonce}')
            nave=c.selectbox('Nave *',opt_blank(catalog('nave')),key=f'pnc_nave_{nonce}')
            a,b,c=st.columns(3)
            opt=a.selectbox('ITEM / Producto *',opt_prod,key=f'pnc_item_{nonce}')
            item=opt.split('|')[0].strip() if opt else ''
            row=prod[prod['item']==item].iloc[0] if item and item in prod['item'].values else None
            descp=str(row['descripcion']) if row is not None else ''; cliente=str(row['cliente']) if row is not None else ''; familia=str(row['familia']) if row is not None else ''
            lote=b.text_area('Lote *',key=f'pnc_lote_{nonce}'); etapa=c.selectbox('Etapa *',opt_blank(catalog('etapa')),key=f'pnc_etapa_{nonce}')
            a,b,c=st.columns(3)
            a.text_input('Descripción',value=descp,disabled=True); b.text_input('Cliente',value=cliente,disabled=True); c.text_input('Familia',value=familia,disabled=True)
            linea=st.selectbox('Línea/Sector *',opt_blank(catalog('linea_sector')),key=f'pnc_linea_{nonce}')
            a,b,c=st.columns(3)
            turno=a.selectbox('Turno *',opt_blank(catalog('turno')),key=f'pnc_turno_{nonce}'); status=b.selectbox('Status *',opt_blank(catalog('status')),key=f'pnc_status_{nonce}'); optd=c.selectbox('Código / Defecto *',opt_defs,key=f'pnc_def_{nonce}')
            cod=optd.split('|')[0].strip() if optd else ''; dr=defs[defs['codigo']==cod].iloc[0] if cod and cod in defs['codigo'].values else None
            defecto=str(dr['defecto']) if dr is not None else ''; tipo=str(dr['tipo_defecto']) if dr is not None else ''; clas=str(dr['clasificacion']) if dr is not None else ''
            d1,d2,d3=st.columns(3)
            d1.text_input('Defecto',value=defecto,disabled=True)
            d2.text_input('Tipo de defecto',value=tipo,disabled=True)
            d3.text_input('Clasificación *',value=clas,disabled=True)
            cat1,cat2=st.columns(2)
            categoria_inicial_pnc=cat1.selectbox('Categoría inicial *',['','1','2','3'],key=f'pnc_cat_inicial_{nonce}')
            categoria_final_pnc=cat2.selectbox('Categoría final',['','1','2','3'],key=f'pnc_cat_final_{nonce}')
            descripcion=st.text_area('Descripción del defecto *',key=f'pnc_desc_{nonce}'); acciones=st.text_area('Acciones inmediatas *',key=f'pnc_accion_{nonce}')
            a,b,c=st.columns(3)
            sup=a.selectbox('Supervisor (Responsable) *',opt_blank(catalog('supervisor')),key=f'pnc_sup_{nonce}'); ana=b.selectbox('Analista (Persona que detecta) *',opt_blank(catalog('analista')),key=f'pnc_ana_{nonce}'); resp=c.selectbox('Responsable de detectar el PNC *',opt_blank(catalog('responsable_detecta')),key=f'pnc_resp_{nonce}')
            a,b,c=st.columns(3)
            disp=a.selectbox('Disposición *',opt_blank(catalog('disposicion')),key=f'pnc_disp_{nonce}'); obs=b.number_input('Cantidad observada (kg) *',min_value=0.0,step=1.0,format='%.2f',key=f'pnc_obs_{nonce}'); fecha_final=c.date_input('Fecha final',value=date.today(),key=f'pnc_final_{nonce}') if status=='CERRADO' else None
            q1,q2,q3=st.columns(3); rep=q1.number_input('Reproceso kg',min_value=0.0,step=1.0,format='%.2f',key=f'pnc_rep_{nonce}'); dec=q2.number_input('Decomiso kg',min_value=0.0,step=1.0,format='%.2f',key=f'pnc_dec_{nonce}'); apr=q3.number_input('Aprobado 2da kg',min_value=0.0,step=1.0,format='%.2f',key=f'pnc_apr_{nonce}'); total=float(obs)
            mat=st.text_area('Material hallado / ME',key=f'pnc_mat_{nonce}'); notas=st.text_area('Observaciones',key=f'pnc_notas_{nonce}'); files=st.file_uploader('Adjuntar evidencia',accept_multiple_files=True,type=['pdf','png','jpg','jpeg','xlsx','csv','txt','docx'],key=f'pnc_files_{nonce}')
            ok=st.button('Guardar registro',key=f'guardar_pnc_{nonce}',type='primary')
        if ok:
            obligatorios={'Línea/Sector':linea,'Nave':nave,'ITEM':item,'Descripción':descp,'Cliente':cliente,'Familia':familia,'Lote':lote,'Etapa':etapa,'Código':cod,'Defecto':defecto,'Tipo de defecto':tipo,'Semana':semana,'Turno':turno,'Fecha':fecha,'Supervisor':sup,'Analista':ana,'Responsable de detectar el PNC':resp,'Descripción del defecto':descripcion,'Acciones inmediatas':acciones,'Disposición':disp,'Cantidad observada':obs,'Status':status,'Clasificación':clas,'Categoría inicial':categoria_inicial_pnc}
            faltantes=[k for k,v in obligatorios.items() if v is None or (isinstance(v,str) and not v.strip()) or (k=='Cantidad observada' and float(v)<=0)]
            if faltantes:
                st.error('Completa los siguientes campos obligatorios: '+', '.join(faltantes)+'.')
            else:
                folio=new_folio(); rid=exec_sql('INSERT INTO pnc_registros(folio,fecha_apertura,linea_sector,nave,item,descripcion_producto,cliente,familia,lote,etapa,codigo_defecto,defecto,tipo_defecto,clasificacion,turno,supervisor,analista,responsable_detecta,descripcion_defecto,acciones_inmediatas,disposicion,cantidad_observada,cantidad_reproceso,cantidad_decomiso,cantidad_aprobado_segunda,cantidad_total_pnc,status,fecha_final_tratamiento,observaciones,material_hallado,creado_por,creado_en,semana,categoria_inicial_pnc,categoria_final_pnc) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(folio,fecha.isoformat(),linea,nave,item,descp,cliente,familia,lote,etapa,cod,defecto,tipo,clas,turno,sup,ana,resp,descripcion,acciones,disp,float(obs),rep,dec,apr,total,status,fecha_final.isoformat() if fecha_final else None,notas,mat,st.session_state.auth['usuario'],now_iso(),int(semana),categoria_inicial_pnc,categoria_final_pnc))
                save_files(files,rid,folio,st.session_state.auth['usuario']); audit(st.session_state.auth['usuario'],'CREAR_PNC',f'ID {rid}')
                limpiar_form(); st.session_state.flash_registro_guardado=f'Registro guardado correctamente: Número {rid}'
                st.rerun()
        st.markdown('</div></div>',unsafe_allow_html=True)
    if st.session_state.registro_tipo is None: selector()
    elif st.session_state.registro_tipo=='PNC': form_pnc()
    elif st.session_state.registro_tipo=='ME': form_hallazgo('me_registros','🧲 Materia Extraña','CREAR_ME')
    elif st.session_state.registro_tipo=='DDM_RX': form_hallazgo('ddm_rx_registros','📦 Producto segregado por detector de metales y RX','CREAR_DDM_RX')
    elif st.session_state.registro_tipo=='RECLAMOS': form_reclamos()
    elif st.session_state.registro_tipo=='DEVOLUCIONES': form_devoluciones()
def page_consulta():
    if 'consulta_tipo' not in st.session_state:
        st.session_state.consulta_tipo=None

    if st.session_state.consulta_tipo is None:
        st.markdown('''<div class="registro-landing-hero"><div class="registro-landing-title">Consulta y descarga</div><div class="registro-landing-subtitle">Selecciona la sección que deseas consultar para acceder a su información.</div></div>''',unsafe_allow_html=True)
        tarjetas=[
            ('NO_CONFORMIDADES','📋  **Consulta y seguimiento de No Conformes**\n\nPNC, Materia Extraña, Detector de metales/RX, Reclamos y Devoluciones.\n\n*Abrir sección*'),
            ('MUESTRAS','🧪  **Muestras de retención**\n\nConsulta, edición, eliminación y descarga de muestras.\n\n*Abrir sección*'),
            ('MATRIZ','📊  **Matriz de entrega de turno**\n\nRegistros de las tres naves, indicadores y reportes por registro.\n\n*Abrir sección*')
        ]
        for columna,(valor,texto) in zip(st.columns(3,gap='large'),tarjetas):
            with columna:
                st.markdown('<span class="registro-card-slot"></span>',unsafe_allow_html=True)
                if st.button(texto,key=f'consulta_tarjeta_{valor}'):
                    st.session_state.consulta_tipo=valor
                    st.rerun()
        return

    if st.button('← Volver a Consulta y descarga',key='consulta_volver_tarjetas'):
        st.session_state.consulta_tipo=None
        st.rerun()

    tipo_consulta=st.session_state.consulta_tipo
    titulos={
        'NO_CONFORMIDADES':'Consulta y seguimiento de No Conformes',
        'MUESTRAS':'Muestras de retención',
        'MATRIZ':'Matriz de entrega de turno'
    }
    st.markdown(f'''<div class="registro-full-panel"><div class="registro-pill">Consulta y descarga</div><div class="registro-full-title">{titulos.get(tipo_consulta,'Consulta')}</div></div>''',unsafe_allow_html=True)
    def prep(df):
        v=df.copy()
        if all(c in v.columns for c in ['dia','mes','anio']):
            v['Fecha']=pd.to_datetime(dict(year=v['anio'].fillna(1900).astype(int),month=v['mes'].fillna(1).astype(int),day=v['dia'].fillna(1).astype(int)),errors='coerce').dt.date
            v=v.drop(columns=['dia','mes','anio'])
        for c in ['folio','numero']:
            if c in v.columns: v=v.drop(columns=[c])
        v=v.rename(columns={'id':'Número','fecha_apertura':'Fecha','descripcion_producto':'Producto','linea_sector':'Línea/Sector','descripcion_hallazgo':'Descripción del hallazgo','particulas_halladas':'# partículas','equipo_hallazgo':'Equipo','analista_detecta':'Analista','supervisor_responsable':'Supervisor'})
        cols=list(v.columns)
        if 'Número' in cols and 'Fecha' in cols: v=v[['Número','Fecha']+[c for c in cols if c not in ['Número','Fecha']]]
        return v
    def filtered(df,key):
        q=st.text_input('Buscar registro',placeholder='Buscar por número, fecha, producto, lote, línea, familia, analista...',key=f'buscar_{key}')
        if q and not df.empty:
            mask=df.astype(str).apply(lambda c: c.str.contains(q,case=False,na=False)).any(axis=1)
            df=df[mask]
        return df
    def table(df,key):
        df=filtered(df,key)
        if df.empty:
            st.info('No se encontraron registros para mostrar.'); return None,df
        v=prep(df)
        try:
            nonce_key=f'table_nonce_{key}'
            nonce=st.session_state.get(nonce_key,0)
            ev=st.dataframe(v,use_container_width=True,hide_index=True,on_select='rerun',selection_mode='single-row',key=f'table_{key}_{nonce}')
            rows=getattr(ev,'selection',{}).get('rows',[]) if ev is not None else []
            posicion_valida=bool(rows) and isinstance(rows[0],int) and 0 <= rows[0] < len(v)
            if posicion_valida:
                numero=v.iloc[rows[0]].get('Número')
                if pd.notna(numero): return int(numero),df
            elif rows:
                st.session_state[nonce_key]=nonce+1
                st.rerun()
        except Exception:
            st.dataframe(v,use_container_width=True,hide_index=True)
        return None,df
    def delete_confirm(table_name,key,audit_name,selected):
        if st.button('Eliminar registro',key=f'del_{key}_{selected}'):
            st.session_state[f'confirm_del_{key}']=selected
        if st.session_state.get(f'confirm_del_{key}')==selected:
            st.warning(f'Confirma la eliminación del registro Número {selected}. Esta acción no se puede deshacer.')
            c1,c2=st.columns(2)
            with c1:
                if st.button('Confirmar eliminación',key=f'ok_del_{key}_{selected}'):
                    exec_sql(f'DELETE FROM {table_name} WHERE id=?',(selected,)); reset_autoincrement(table_name); audit(st.session_state.auth['usuario'],audit_name,f'ID {selected}'); st.session_state.pop(f'confirm_del_{key}',None); st.session_state[f'table_nonce_{key}']=st.session_state.get(f'table_nonce_{key}',0)+1; st.success(f'Registro eliminado correctamente: Número {selected}'); st.rerun()
            with c2:
                if st.button('Cancelar',key=f'cancel_del_{key}_{selected}'):
                    st.session_state.pop(f'confirm_del_{key}',None); st.rerun()
    def edit_record(table_name,key,selected):
        row_df=read_df(f'SELECT * FROM {table_name} WHERE id=?',(selected,))
        if row_df.empty:
            st.warning('El registro seleccionado ya no está disponible.')
            return
        row=row_df.iloc[0].to_dict()
        prod=read_df('SELECT * FROM productos WHERE activo=1 ORDER BY descripcion')
        defs=read_df('SELECT * FROM defectos WHERE activo=1 ORDER BY CAST(codigo AS INTEGER)')
        opt_prod=['']+[f'{r.item} | {r.descripcion}' for r in prod.itertuples()]
        opt_defs=['']+[f'{r.codigo} | {r.defecto}' for r in defs.itertuples()]
        original_item=str(row.get('item') or '')
        original_codigo=str(row.get('codigo_defecto') or '')
        producto_actual=next((x for x in opt_prod if x.split('|')[0].strip()==original_item),'')
        defecto_actual=next((x for x in opt_defs if x.split('|')[0].strip()==original_codigo),'')
        labels={'fecha_apertura':'Fecha','dia':'Día','mes':'Mes','anio':'Año','linea_sector':'Línea/Sector *','nave':'Nave *','lote':'Lote *','etapa':'Etapa *','semana':'Semana *','turno':'Turno *','supervisor':'Supervisor *','supervisor_responsable':'Supervisor *','analista':'Analista *','analista_detecta':'Analista *','responsable_detecta':'Responsable de detectar el PNC *','descripcion_defecto':'Descripción del defecto *','descripcion_hallazgo':'Descripción del defecto *','acciones_inmediatas':'Acciones inmediatas *','accion_contingente':'Acciones inmediatas','disposicion':'Disposición *','cantidad_observada':'Cantidad observada (kg) *','status':'Status *','equipo_hallazgo':'Equipo del hallazgo','tipo':'Tipo de hallazgo','particulas_halladas':'Partículas halladas','investigacion_origen':'Investigación del origen','acciones_evitar_incidencia':'Acciones para evitar incidencia','observaciones':'Observaciones','material_hallado':'Material hallado / ME','fecha_final_tratamiento':'Fecha final','cantidad_reproceso':'Reproceso kg','cantidad_decomiso':'Decomiso kg','cantidad_aprobado_segunda':'Aprobado segunda instancia kg','categoria_inicial_pnc':'Categoría inicial *','categoria_final_pnc':'Categoría final'}
        protected={'folio','creado_por','creado_en','actualizado_por','actualizado_en','cantidad_total_pnc','item','descripcion_producto','producto','cliente','familia','codigo_defecto','defecto','tipo_defecto','clasificacion','categoria_inicial'}
        numeric={'dia','mes','anio','semana','particulas_halladas','cantidad_observada','cantidad_reproceso','cantidad_decomiso','cantidad_aprobado_segunda'}
        select_catalog={'linea_sector':'linea_sector','nave':'nave','etapa':'etapa','turno':'turno','supervisor':'supervisor','supervisor_responsable':'supervisor','analista':'analista','analista_detecta':'analista','responsable_detecta':'responsable_detecta','disposicion':'disposicion','status':'status','categoria_inicial_pnc':'categoria_pnc','categoria_final_pnc':'categoria_pnc'}
        with st.expander('✏️ Editar o completar registro',expanded=True):
            titulo={'pnc':'PNC','me':'Materia Extraña','ddm':'Detector de metales y RX'}.get(key,key)
            st.markdown(f'#### Editar registro de {titulo}: Número {selected}')
            st.caption('El registro original se carga completo. Al cambiar ITEM o Código, los campos relacionados se actualizan automáticamente.')
            a1,a2=st.columns(2)
            opt_item=a1.selectbox('ITEM *',opt_prod,index=idx_or_zero(opt_prod,producto_actual),key=f'edit_live_item_{key}_{selected}')
            opt_codigo=a2.selectbox('Código / Defecto *',opt_defs,index=idx_or_zero(opt_defs,defecto_actual),key=f'edit_live_codigo_{key}_{selected}')
            item=opt_item.split('|')[0].strip() if opt_item else ''
            pr=prod[prod['item']==item].iloc[0] if item and item in prod['item'].values else None
            descripcion=str(pr['descripcion']) if pr is not None else ''
            cliente=str(pr['cliente']) if pr is not None else ''
            familia=str(pr['familia']) if pr is not None else ''
            codigo=opt_codigo.split('|')[0].strip() if opt_codigo else ''
            dr=defs[defs['codigo']==codigo].iloc[0] if codigo and codigo in defs['codigo'].values else None
            defecto=str(dr['defecto']) if dr is not None else ''
            tipo_defecto=str(dr['tipo_defecto']) if dr is not None else ''
            clasificacion=str(dr['clasificacion']) if dr is not None else ''
            p1,p2,p3=st.columns(3)
            p1.text_input('Descripción',value=descripcion,disabled=True,key=f'edit_auto_descripcion_{key}_{selected}')
            p2.text_input('Cliente',value=cliente,disabled=True,key=f'edit_auto_cliente_{key}_{selected}')
            p3.text_input('Familia',value=familia,disabled=True,key=f'edit_auto_familia_{key}_{selected}')
            d1,d2,d3=st.columns(3)
            d1.text_input('Defecto',value=defecto,disabled=True,key=f'edit_auto_defecto_{key}_{selected}')
            d2.text_input('Tipo de defecto',value=tipo_defecto,disabled=True,key=f'edit_auto_tipo_defecto_{key}_{selected}')
            d3.text_input('Clasificación',value=clasificacion,disabled=True,key=f'edit_auto_clasificacion_{key}_{selected}')
            values={}
            columns=[c for c in row_df.columns if c!='id' and c not in protected]
            if table_name in {'me_registros','ddm_rx_registros'}:
                campos_retirados={'etapa','responsable_detecta','disposicion','status','cantidad_observada'}
                columns=[c for c in columns if c not in campos_retirados]
            with st.form(f'editar_form_{key}_{selected}'):
                for pos in range(0,len(columns),3):
                    cols_ui=st.columns(3)
                    for ui,col in zip(cols_ui,columns[pos:pos+3]):
                        raw=row.get(col); value='' if raw is None or pd.isna(raw) else raw
                        label=labels.get(col,col.replace('_',' ').title())
                        if col in select_catalog:
                            options=['','1','2','3'] if select_catalog[col]=='categoria_pnc' else opt_blank(catalog(select_catalog[col])); current=str(value)
                            if current and current not in options: options.append(current)
                            values[col]=ui.selectbox(label,options,index=idx_or_zero(options,current),key=f'edit_{key}_{selected}_{col}')
                        elif col in numeric:
                            if col in {'dia','mes','anio','semana','particulas_halladas'}: values[col]=ui.number_input(label,min_value=0,step=1,value=int(float(value or 0)),key=f'edit_{key}_{selected}_{col}')
                            else: values[col]=ui.number_input(label,min_value=0.0,step=1.0,format='%.2f',value=float(value or 0),key=f'edit_{key}_{selected}_{col}')
                        elif col in {'descripcion_defecto','descripcion_hallazgo','acciones_inmediatas','accion_contingente','investigacion_origen','acciones_evitar_incidencia','observaciones','material_hallado','lote'}:
                            values[col]=ui.text_area(label,value=str(value),key=f'edit_{key}_{selected}_{col}')
                        else: values[col]=ui.text_input(label,value=str(value),key=f'edit_{key}_{selected}_{col}')
                guardar=st.form_submit_button('Guardar cambios',type='primary')
            if guardar:
                auto={'item':item,'cliente':cliente,'familia':familia,'codigo_defecto':codigo,'defecto':defecto,'tipo_defecto':tipo_defecto}
                if table_name=='pnc_registros':
                    auto.update({'descripcion_producto':descripcion,'clasificacion':clasificacion,'cantidad_total_pnc':float(values.get('cantidad_observada') or 0)})
                else: auto.update({'producto':descripcion,'categoria_inicial':clasificacion})
                if table_name=='pnc_registros':
                    required=['linea_sector','nave','lote','etapa','semana','turno','responsable_detecta','acciones_inmediatas','disposicion','cantidad_observada','status','supervisor','analista','descripcion_defecto','categoria_inicial_pnc']
                else:
                    required=['linea_sector','nave','lote','semana','turno','acciones_inmediatas','supervisor_responsable','analista_detecta','descripcion_hallazgo']
                missing=[]
                for name,val in {'ITEM':item,'Descripción':descripcion,'Cliente':cliente,'Familia':familia,'Código':codigo,'Defecto':defecto,'Tipo de defecto':tipo_defecto,'Clasificación':clasificacion}.items():
                    if not str(val).strip(): missing.append(name)
                for col in required:
                    val=values.get(col)
                    if val is None or (isinstance(val,str) and not val.strip()) or (col=='cantidad_observada' and float(val)<=0): missing.append(labels.get(col,col))
                if missing: st.error('Completa los siguientes campos obligatorios: '+', '.join(dict.fromkeys(missing))+'.')
                else:
                    updates={**values,**auto}
                    columnas_tabla=table_columns(table_name)
                    if {'actualizado_por','actualizado_en'}.issubset(columnas_tabla):
                        updates['actualizado_por']=st.session_state.auth['usuario']
                        updates['actualizado_en']=now_iso()
                    update_cols=list(updates.keys())
                    assignments=', '.join([f'"{c}"=?' for c in update_cols])
                    exec_sql(f'UPDATE {table_name} SET {assignments} WHERE id=?',tuple(updates[c] for c in update_cols)+(selected,))
                    audit(st.session_state.auth['usuario'],'EDITAR_REGISTRO',f'Tabla {table_name} | ID {selected}')
                    st.success(f'Registro actualizado correctamente: Número {selected}')
                    st.rerun()
    if tipo_consulta=='MUESTRAS':
        consulta_muestras_retencion()
        return
    if tipo_consulta=='MATRIZ':
        matriz_entregas()
        return

    t1,t2,t3,t4,t5=st.tabs(['PNC´s','Materia Extraña','Detector de metales y RX','Reclamos','Devoluciones'])
    with t1:
        df=read_df('SELECT * FROM pnc_registros ORDER BY id ASC')
        auditoria=['creado_por','creado_en','actualizado_por','actualizado_en']
        base=[c for c in df.columns if c not in auditoria]
        df_tabla=df[base+[c for c in auditoria if c in df.columns]]
        selected,shown=table(df_tabla,'pnc')
        if not shown.empty: st.download_button('Descargar PNC CSV',prep(shown).to_csv(index=False).encode('utf-8-sig'),f"pnc_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",'text/csv')
        if selected:
            st.markdown(f'### Registro seleccionado: Número {selected}')
            contenido_pdf=pdf_pnc(selected)
            if contenido_pdf:
                st.download_button('Descargar informe PNC en PDF',contenido_pdf,f'informe_pnc_{selected}.pdf','application/pdf',key=f'pnc_pdf_{selected}')
            else:
                st.warning('No fue posible generar el informe PDF. Verifica que reportlab esté incluido en requirements.txt.')
            contenido_fisico=pdf_pnc_fisico(selected)
            if contenido_fisico:
                st.download_button('Descargar Registro Físico en PDF',contenido_fisico,f'registro_fisico_pnc_{selected}.pdf','application/pdf',key=f'pnc_fisico_pdf_{selected}')
            else:
                st.warning('No fue posible generar el Registro Físico. Verifica que reportlab esté incluido en requirements.txt.')
            edit_record('pnc_registros','pnc',selected)
            if is_dev(): delete_confirm('pnc_registros','pnc','ELIMINAR_PNC',selected)
    with t2:
        df=read_df('SELECT * FROM me_registros ORDER BY id ASC')
        campos_actuales=['id','dia','mes','anio','semana','nave','item','producto','familia','lote','linea_sector','codigo_defecto','turno','supervisor_responsable','analista_detecta','descripcion_hallazgo','acciones_inmediatas','equipo_hallazgo','tipo','particulas_halladas','investigacion_origen','acciones_evitar_incidencia','creado_por','creado_en','actualizado_por','actualizado_en']
        disponibles=[c for c in campos_actuales if c in df.columns]
        df_tabla=df[disponibles].copy()
        if all(c in df_tabla.columns for c in ['dia','mes','anio']):
            df_tabla.insert(1,'Fecha',pd.to_datetime(dict(year=df_tabla.pop('anio'),month=df_tabla.pop('mes'),day=df_tabla.pop('dia')),errors='coerce').dt.strftime('%Y-%m-%d'))
        selected,shown=table(df_tabla,'me')
        if not shown.empty: st.download_button('Descargar Materia Extraña CSV',prep(shown).to_csv(index=False).encode('utf-8-sig'),f"materia_extrana_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",'text/csv')
        if selected:
            st.markdown(f'### Registro seleccionado: Número {selected}')
            contenido_pdf=pdf_materia_extrana(selected)
            if contenido_pdf:
                st.download_button('Descargar reporte de Materia Extraña en PDF',contenido_pdf,f'materia_extrana_{selected}.pdf','application/pdf',key=f'me_pdf_{selected}')
            else:
                st.warning('No fue posible generar el PDF. Verifica que reportlab esté incluido en requirements.txt.')
            edit_record('me_registros','me',selected)
            if is_dev(): delete_confirm('me_registros','me','ELIMINAR_ME',selected)
    with t3:
        df=read_df('SELECT * FROM ddm_rx_registros ORDER BY id ASC')
        campos_actuales=['id','dia','mes','anio','semana','nave','item','producto','familia','lote','linea_sector','codigo_defecto','turno','supervisor_responsable','analista_detecta','descripcion_hallazgo','acciones_inmediatas','equipo_hallazgo','tipo','particulas_halladas','investigacion_origen','acciones_evitar_incidencia','creado_por','creado_en','actualizado_por','actualizado_en']
        disponibles=[c for c in campos_actuales if c in df.columns]
        df_tabla=df[disponibles].copy()
        if all(c in df_tabla.columns for c in ['dia','mes','anio']):
            df_tabla.insert(1,'Fecha',pd.to_datetime(dict(year=df_tabla.pop('anio'),month=df_tabla.pop('mes'),day=df_tabla.pop('dia')),errors='coerce').dt.strftime('%Y-%m-%d'))
        selected,shown=table(df_tabla,'ddm')
        if not shown.empty: st.download_button('Descargar Detector de metales y RX CSV',prep(shown).to_csv(index=False).encode('utf-8-sig'),f"ddm_rx_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",'text/csv')
        if selected:
            st.markdown(f'### Registro seleccionado: Número {selected}')
            contenido_pdf=pdf_detector_metales_rx(selected)
            if contenido_pdf:
                st.download_button('Descargar registro de Detector de metales y RX en PDF',contenido_pdf,f'detector_metales_rx_{selected}.pdf','application/pdf',key=f'ddm_pdf_{selected}')
            else:
                st.warning('No fue posible generar el PDF. Verifica que reportlab esté incluido en requirements.txt.')
            edit_record('ddm_rx_registros','ddm',selected)
            if is_dev(): delete_confirm('ddm_rx_registros','ddm','ELIMINAR_DDM_RX',selected)

    with t4:
        df=read_df('SELECT * FROM reclamos_registros ORDER BY id ASC')
        selected,shown=table(df,'reclamos')
        if not shown.empty:
            st.download_button('Descargar Reclamos CSV',prep(shown).to_csv(index=False).encode('utf-8-sig'),f"reclamos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",'text/csv')
        if selected:
            st.markdown(f'### Reclamo seleccionado: Número {selected}')
            row_df=read_df('SELECT * FROM reclamos_registros WHERE id=?',(selected,))
            if row_df.empty:
                st.warning('El registro seleccionado ya no está disponible.')
            else:
                row=row_df.iloc[0].to_dict()
                prod=read_df('SELECT * FROM productos WHERE activo=1 ORDER BY descripcion')
                defs=read_df('SELECT * FROM defectos WHERE activo=1 ORDER BY CAST(codigo AS INTEGER)')
                opt_prod=['']+[f'{r.item} | {r.descripcion}' for r in prod.itertuples()]
                opt_defs=['']+[f'{r.codigo} | {r.defecto}' for r in defs.itertuples()]
                item_actual=str(row.get('item') or '')
                codigo_actual=str(row.get('codigo_defecto') or '')
                producto_actual=next((x for x in opt_prod if x.split('|')[0].strip()==item_actual),'')
                defecto_actual=next((x for x in opt_defs if x.split('|')[0].strip()==codigo_actual),'')
                fecha_actual=pd.to_datetime(row.get('fecha'),errors='coerce')
                cierre_actual=pd.to_datetime(row.get('fecha_cierre'),errors='coerce')
                with st.expander('✏️ Editar registro de Reclamo',expanded=True):
                    with st.form(f'editar_reclamo_{selected}'):
                        a,b,c=st.columns(3)
                        fecha=a.date_input('Fecha *',value=fecha_actual.date() if pd.notna(fecha_actual) else date.today())
                        od=b.selectbox('Código / Defecto *',opt_defs,index=idx_or_zero(opt_defs,defecto_actual))
                        fuente=c.text_input('Fuente *',value=str(row.get('fuente') or ''))
                        codigo=od.split('|')[0].strip() if od else ''
                        dr=defs[defs['codigo'].astype(str)==codigo].iloc[0] if codigo and codigo in defs['codigo'].astype(str).values else None
                        defecto=str(dr['defecto']) if dr is not None else ''
                        tipo_defecto=str(dr['tipo_defecto']) if dr is not None else ''
                        clasificacion=str(dr['clasificacion']) if dr is not None else ''
                        a,b,c=st.columns(3)
                        a.text_input('Descripción del defecto',value=defecto,disabled=True)
                        b.text_input('Tipo de defecto',value=tipo_defecto,disabled=True)
                        c.text_input('Clasificación del defecto',value=clasificacion,disabled=True)
                        a,b,c=st.columns(3)
                        mercado_ops=['','Interno','Externo']; mercado_actual=str(row.get('mercado') or '')
                        if mercado_actual and mercado_actual not in mercado_ops: mercado_ops.append(mercado_actual)
                        mercado=a.selectbox('Mercado *',mercado_ops,index=idx_or_zero(mercado_ops,mercado_actual))
                        pais_estado=b.text_input('País / Estado *',value=str(row.get('pais_estado') or ''))
                        numero_caso=c.text_input('No. Caso Right Now / MDLZ *',value=str(row.get('numero_caso') or ''))
                        a,b,c=st.columns(3)
                        op=a.selectbox('ITEM *',opt_prod,index=idx_or_zero(opt_prod,producto_actual))
                        item=op.split('|')[0].strip() if op else ''
                        pr=prod[prod['item'].astype(str)==item].iloc[0] if item and item in prod['item'].astype(str).values else None
                        producto='' if pr is None or pd.isna(pr['descripcion']) else str(pr['descripcion']).strip()
                        cliente='' if pr is None or pd.isna(pr['cliente']) else str(pr['cliente']).strip()
                        familia='' if pr is None or pd.isna(pr['familia']) else str(pr['familia']).strip()
                        b.text_input('Producto',value=producto,disabled=True)
                        c.text_input('Cliente',value=cliente,disabled=True)
                        st.text_input('Familia',value=familia,disabled=True)
                        descripcion_reclamo=st.text_area('Descripción del reclamo *',value=str(row.get('descripcion_reclamo') or ''))
                        causa_raiz=st.text_area('Causa raíz del defecto *',value=str(row.get('causa_raiz') or ''))
                        acciones=st.text_area('Acciones correctivas y contingentes *',value=str(row.get('acciones_correctivas_contingentes') or ''))
                        a,b,c=st.columns(3)
                        comp_ops=['','Sí','No']; comprobado=a.selectbox('Comprobado *',comp_ops,index=idx_or_zero(comp_ops,str(row.get('comprobado') or '')))
                        cantidad=b.number_input('Cantidad afectada *',min_value=0.0,step=1.0,format='%.2f',value=float(row.get('cantidad_afectada') or 0))
                        unidad_ops=['','Bulto','Bolsa','Pieza','Lámina']; unidad_actual=str(row.get('unidad') or '')
                        if unidad_actual and unidad_actual not in unidad_ops: unidad_ops.append(unidad_actual)
                        unidad=c.selectbox('Unidad *',unidad_ops,index=idx_or_zero(unidad_ops,unidad_actual))
                        a,b,c=st.columns(3)
                        sector=a.text_input('Sector *',value=str(row.get('sector') or ''))
                        estado_ops=['','Abierto','Cerrado']; estado_actual=str(row.get('estado_reclamo') or '')
                        estado=b.selectbox('Estado del reclamo *',estado_ops,index=idx_or_zero(estado_ops,estado_actual))
                        fecha_cierre=c.date_input('Fecha de cierre',value=cierre_actual.date() if pd.notna(cierre_actual) else None)
                        a,b,c=st.columns(3)
                        tsp=a.text_input('TSP N°',value=str(row.get('tsp_numero') or ''))
                        caducidad=b.text_input('Caducidad',value=str(row.get('caducidad') or ''))
                        lote=c.text_input('Lote *',value=str(row.get('lote') or ''))
                        a,b,c=st.columns(3)
                        nave_ops=opt_blank(catalog('nave')); nave_actual=str(row.get('nave') or '')
                        if nave_actual and nave_actual not in nave_ops: nave_ops.append(nave_actual)
                        nave=a.selectbox('Nave *',nave_ops,index=idx_or_zero(nave_ops,nave_actual))
                        pmd_ops=['','1 - Manipulación','2 - Producción','3 - Diseño']; pmd_actual=str(row.get('pmd') or '')
                        if pmd_actual and pmd_actual not in pmd_ops: pmd_ops.append(pmd_actual)
                        pmd=b.selectbox('P / M / D *',pmd_ops,index=idx_or_zero(pmd_ops,pmd_actual))
                        red_ops=['','Mail','Web','Instagram','No indica','Facebook']; red_actual=str(row.get('red_social') or '')
                        if red_actual and red_actual not in red_ops: red_ops.append(red_actual)
                        red_social=c.selectbox('Red social *',red_ops,index=idx_or_zero(red_ops,red_actual))
                        guardar=st.form_submit_button('Guardar cambios',type='primary')
                    if guardar:
                        req={'Código / Defecto':codigo,'Descripción del defecto':defecto,'Fuente':fuente,'Mercado':mercado,'País / Estado':pais_estado,'No. Caso':numero_caso,'ITEM':item,'Producto':producto,'Cliente':cliente,'Familia':familia,'Descripción del reclamo':descripcion_reclamo,'Causa raíz':causa_raiz,'Acciones':acciones,'Comprobado':comprobado,'Cantidad afectada':cantidad,'Unidad':unidad,'Sector':sector,'Estado':estado,'Lote':lote,'Nave':nave,'P / M / D':pmd,'Tipo de defecto':tipo_defecto,'Clasificación':clasificacion,'Red social':red_social}
                        faltan=[k for k,v in req.items() if v is None or (isinstance(v,str) and not v.strip()) or (k=='Cantidad afectada' and float(v)<=0)]
                        duplicado=not read_df('SELECT id FROM reclamos_registros WHERE UPPER(TRIM(numero_caso))=UPPER(TRIM(?)) AND id<>?',(numero_caso,selected)).empty if numero_caso.strip() else False
                        if duplicado: st.error('Ya existe otro reclamo con el mismo No. Caso Right Now / MDLZ.')
                        elif faltan: st.error('Completa los siguientes campos obligatorios: '+', '.join(faltan)+'.')
                        elif estado=='Cerrado' and not fecha_cierre: st.error('Captura la fecha de cierre para el reclamo cerrado.')
                        elif fecha_cierre and fecha_cierre<fecha: st.error('La fecha de cierre no puede ser anterior a la fecha del reclamo.')
                        else:
                            exec_sql('UPDATE reclamos_registros SET fecha=?,codigo_defecto=?,descripcion_defecto=?,fuente=?,mercado=?,pais_estado=?,numero_caso=?,item=?,producto=?,cliente=?,familia=?,descripcion_reclamo=?,causa_raiz=?,acciones_correctivas_contingentes=?,comprobado=?,cantidad_afectada=?,unidad=?,sector=?,estado_reclamo=?,fecha_cierre=?,tsp_numero=?,caducidad=?,lote=?,nave=?,pmd=?,tipo_defecto=?,clasificacion_defecto=?,red_social=?,actualizado_por=?,actualizado_en=? WHERE id=?',(fecha.isoformat(),codigo,defecto,fuente.strip(),mercado,pais_estado.strip(),numero_caso.strip(),item,producto,cliente,familia,descripcion_reclamo.strip(),causa_raiz.strip(),acciones.strip(),comprobado,float(cantidad),unidad,sector.strip(),estado,fecha_cierre.isoformat() if fecha_cierre else None,tsp.strip(),caducidad.strip(),lote.strip(),nave,pmd,tipo_defecto,clasificacion,red_social,st.session_state.auth['usuario'],now_iso(),selected))
                            audit(st.session_state.auth['usuario'],'EDITAR_RECLAMO',f'ID {selected} | Caso {numero_caso.strip()}')
                            st.success(f'Reclamo actualizado correctamente: Número {selected}')
                            st.rerun()
            if is_dev():
                delete_confirm('reclamos_registros','reclamos','ELIMINAR_RECLAMO',selected)

    with t5:
        df=read_df('SELECT * FROM devoluciones_registros ORDER BY id ASC')
        df_tabla=df.drop(columns=['linea'],errors='ignore')
        selected,shown=table(df_tabla,'devoluciones')
        if not shown.empty: st.download_button('Descargar Devoluciones CSV',prep(shown).to_csv(index=False).encode('utf-8-sig'),f"devoluciones_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",'text/csv')
        if selected:
            st.markdown(f'### Devolución seleccionada: Número {selected}')
            row_df=read_df('SELECT * FROM devoluciones_registros WHERE id=?',(selected,))
            if row_df.empty: st.warning('El registro seleccionado ya no está disponible.')
            else:
                r=row_df.iloc[0].to_dict(); prod=read_df('SELECT * FROM productos WHERE activo=1 ORDER BY descripcion'); defs=read_df('SELECT * FROM defectos WHERE activo=1 ORDER BY CAST(codigo AS INTEGER)')
                opt_prod=['']+[f'{x.item} | {x.descripcion}' for x in prod.itertuples()]; opt_defs=['']+[f'{x.codigo} | {x.defecto}' for x in defs.itertuples()]
                item0=str(r.get('item') or ''); cod0=str(r.get('codigo_defecto') or ''); op0=next((x for x in opt_prod if x.split('|')[0].strip()==item0),''); od0=next((x for x in opt_defs if x.split('|')[0].strip()==cod0),'')
                fv=pd.to_datetime(r.get('fecha'),errors='coerce')
                with st.expander('✏️ Editar registro de Devolución',expanded=True):
                    with st.form(f'editar_devolucion_{selected}'):
                        a,b,c=st.columns(3); fecha=a.date_input('Fecha *',fv.date() if pd.notna(fv) else date.today()); od=b.selectbox('Código / Defecto *',opt_defs,index=idx_or_zero(opt_defs,od0)); cedis=c.text_input('Cedis *',str(r.get('cedis') or ''))
                        codigo=od.split('|')[0].strip() if od else ''; dr=defs[defs.codigo.astype(str)==codigo].iloc[0] if codigo and codigo in defs.codigo.astype(str).values else None; defecto='' if dr is None else str(dr.defecto)
                        a,b,c=st.columns(3); pais=a.text_input('País / Estado *',str(r.get('pais_estado') or '')); op=b.selectbox('ITEM *',opt_prod,index=idx_or_zero(opt_prod,op0)); item=op.split('|')[0].strip() if op else ''
                        pr=prod[prod.item.astype(str)==item].iloc[0] if item and item in prod.item.astype(str).values else None; producto='' if pr is None or pd.isna(pr.descripcion) else str(pr.descripcion).strip(); cliente='' if pr is None or pd.isna(pr.cliente) else str(pr.cliente).strip(); familia='' if pr is None or pd.isna(pr.familia) else str(pr.familia).strip()
                        c.text_input('Producto',producto,disabled=True); a,b,c=st.columns(3); a.text_input('Cliente',cliente,disabled=True); b.text_input('Familia',familia,disabled=True); lote=c.text_input('Lote *',str(r.get('lote') or ''))
                        a,b,c=st.columns(3); cad=a.text_input('Caducidad *',str(r.get('caducidad') or '')); comp=b.selectbox('Comprobado *',['','Sí','No'],index=idx_or_zero(['','Sí','No'],str(r.get('comprobado') or ''))); cant=c.number_input('Cant. afectada *',0.0,step=1.0,value=float(r.get('cantidad_afectada') or 0))
                        desc=st.text_area('Descripción de la devolución *',str(r.get('descripcion_defecto') or ''))
                        a,b,c=st.columns(3); uops=['','Bulto','Bolsa','Pieza','Tarima']; unidad=a.selectbox('Unidad *',uops,index=idx_or_zero(uops,str(r.get('unidad') or ''))); sector=b.text_input('Sector *',str(r.get('sector') or '')); dops=opt_blank(catalog('disposicion')); da=str(r.get('disposicion') or ''); dops=dops+([da] if da and da not in dops else []); disp=c.selectbox('Disposición *',dops,index=idx_or_zero(dops,da))
                        a,b,c=st.columns(3); tsp=a.text_input('TSP N°',str(r.get('tsp_numero') or '')); status=b.selectbox('Status *',['','Cerrado','Abierto'],index=idx_or_zero(['','Cerrado','Abierto'],str(r.get('status') or ''))); nops=opt_blank(catalog('nave')); na=str(r.get('nave') or ''); nops=nops+([na] if na and na not in nops else []); nave=c.selectbox('Nave *',nops,index=idx_or_zero(nops,na)); guardar=st.form_submit_button('Guardar cambios',type='primary')
                    if guardar:
                        req={'Código':codigo,'Cedis':cedis,'País/Estado':pais,'ITEM':item,'Producto':producto,'Cliente':cliente,'Familia':familia,'Lote':lote,'Caducidad':cad,'Descripción':desc,'Comprobado':comp,'Cantidad':cant,'Unidad':unidad,'Sector':sector,'Disposición':disp,'Status':status,'Nave':nave}; faltan=[k for k,v in req.items() if v is None or (isinstance(v,str) and not v.strip()) or (k=='Cantidad' and float(v)<=0)]
                        if faltan: st.error('Completa los siguientes campos obligatorios: '+', '.join(faltan)+'.')
                        else:
                            exec_sql('UPDATE devoluciones_registros SET fecha=?,codigo_defecto=?,defecto=?,cedis=?,pais_estado=?,item=?,producto=?,cliente=?,familia=?,lote=?,caducidad=?,descripcion_defecto=?,comprobado=?,cantidad_afectada=?,unidad=?,sector=?,disposicion=?,tsp_numero=?,status=?,nave=?,actualizado_por=?,actualizado_en=? WHERE id=?',(fecha.isoformat(),codigo,defecto,cedis.strip(),pais.strip(),item,producto,cliente,familia,lote.strip(),cad.strip(),desc.strip(),comp,float(cant),unidad,sector.strip(),disp,tsp.strip(),status,nave,st.session_state.auth['usuario'],now_iso(),selected)); audit(st.session_state.auth['usuario'],'EDITAR_DEVOLUCION',f'ID {selected}'); st.success(f'Devolución actualizada correctamente: Número {selected}'); st.rerun()
            if is_dev(): delete_confirm('devoluciones_registros','devoluciones','ELIMINAR_DEVOLUCION',selected)

def page_muestras_retencion():
    periodos=[
        ('muestras_10_meses','10 Meses','🗓️'),
        ('muestras_12_meses_alergeno','12 Meses Alérgeno','⚠️'),
        ('muestras_12_meses_duvalin','12 Meses Duvalin','🍫'),
        ('muestras_12_meses_nave2','12 Meses Nave 2','🏭'),
        ('muestras_15_meses','15 Meses','📦'),
        ('muestras_18_meses','18 Meses','🧪'),
        ('muestras_24_meses','24 Meses','✅'),
    ]
    if 'muestra_tipo' not in st.session_state: st.session_state.muestra_tipo=None
    if 'muestra_nonce' not in st.session_state: st.session_state.muestra_nonce=0
    seleccionado=st.session_state.muestra_tipo
    if seleccionado is None:
        st.markdown('<div class="registro-landing-hero"><div class="registro-landing-title">Muestras de retención</div><div class="registro-landing-subtitle">Selecciona el periodo de retención que deseas registrar.</div></div>',unsafe_allow_html=True)
        for inicio in range(0,len(periodos),3):
            columnas=st.columns(3,gap='large')
            for col,(tabla,nombre,icono) in zip(columnas,periodos[inicio:inicio+3]):
                with col:
                    st.markdown('<span class="registro-card-slot"></span>',unsafe_allow_html=True)
                    if st.button(f'{icono}  **{nombre}**\n\nCaptura y seguimiento de muestras de retención.\n\n*Abrir registro*',key=f'card_{tabla}'):
                        st.session_state.muestra_tipo=tabla; st.rerun()
        return
    info=next((x for x in periodos if x[0]==seleccionado),None)
    if not info:
        st.session_state.muestra_tipo=None; st.rerun()
    tabla,nombre,icono=info
    if st.button('← Cambiar periodo de retención',key='volver_muestras'):
        st.session_state.muestra_tipo=None; st.rerun()
    st.markdown(f'<div class="registro-full-panel"><div class="registro-pill">Nuevo registro</div><div class="registro-full-title">{icono} {nombre}</div><div class="registro-full-subtitle">Los campos marcados con * son obligatorios.</div>',unsafe_allow_html=True)
    productos=read_df('SELECT item,descripcion FROM productos WHERE activo=1 ORDER BY descripcion')
    opciones=['']+[f'{r.item} | {r.descripcion}' for r in productos.itertuples()]
    nonce=st.session_state.muestra_nonce
    elegido=st.selectbox('ITEM *',opciones,key=f'mr_item_{tabla}_{nonce}')
    item=elegido.split('|')[0].strip() if elegido else ''
    fila=productos[productos['item'].astype(str)==item]
    descripcion=str(fila.iloc[0]['descripcion']) if not fila.empty else ''
    st.text_input('Descripción',value=descripcion,disabled=True)
    with st.form(f'mr_form_{tabla}_{nonce}'):
        c1,c2=st.columns(2)
        lote=c1.text_input('Lote *')
        destino=c2.selectbox('Destino *',['','Nacional','Exportación'])
        c3,c4=st.columns(2)
        muestras=c3.number_input('# De muestras *',min_value=0.0,step=1.0,format='%.2f')
        corrugado=c4.number_input('# Corrugado *',min_value=0.0,step=1.0,format='%.2f')
        responsable=st.selectbox('Responsable *',opt_blank(catalog('analista')))
        observaciones=st.text_area('Observaciones')
        guardar=st.form_submit_button('Guardar registro',type='primary')
    if guardar:
        req={'ITEM':item,'Descripción':descripcion,'Lote':lote,'Destino':destino,'# De muestras':muestras,'# Corrugado':corrugado,'Responsable':responsable}
        faltan=[k for k,v in req.items() if v is None or (isinstance(v,str) and not v.strip()) or (k in ['# De muestras','# Corrugado'] and float(v)<=0)]
        if faltan: st.error('Completa los siguientes campos obligatorios: '+', '.join(faltan)+'.')
        else:
            rid=exec_sql(f'INSERT INTO {tabla}(item,descripcion,lote,destino,numero_muestras,numero_corrugado,responsable,observaciones,creado_por,creado_en) VALUES(?,?,?,?,?,?,?,?,?,?)',(item,descripcion,lote.strip(),destino,float(muestras),float(corrugado),responsable,observaciones.strip(),st.session_state.auth['usuario'],now_iso()))
            audit(st.session_state.auth['usuario'],'CREAR_MUESTRA_RETENCION',f'{nombre} | ID {rid}')
            st.session_state.muestra_nonce+=1
            st.success(f'Registro guardado correctamente. Número {rid}')
            st.rerun()
    st.markdown('</div>',unsafe_allow_html=True)

def consulta_muestras_retencion():
    configuracion=[
        ('muestras_10_meses','10 Meses'),('muestras_12_meses_alergeno','12 Meses Alérgeno'),
        ('muestras_12_meses_duvalin','12 Meses Duvalin'),('muestras_12_meses_nave2','12 Meses Nave 2'),
        ('muestras_15_meses','15 Meses'),('muestras_18_meses','18 Meses'),('muestras_24_meses','24 Meses')]
    st.markdown('## Muestras de retención')
    tabs=st.tabs([n for _,n in configuracion])
    productos=read_df('SELECT item,descripcion FROM productos WHERE activo=1 ORDER BY descripcion')
    opciones_item=['']+[f'{r.item} | {r.descripcion}' for r in productos.itertuples()]
    for tab,(tabla,nombre) in zip(tabs,configuracion):
        with tab:
            df=read_df(f'SELECT * FROM {tabla} ORDER BY id DESC')
            buscar=st.text_input('Buscar registro',placeholder='Buscar por número, ITEM, descripción, lote, destino o responsable...',key=f'buscar_{tabla}')
            mostrado=df.copy()
            if buscar and not mostrado.empty:
                mask=mostrado.astype(str).apply(lambda col: col.str.contains(buscar,case=False,na=False)).any(axis=1)
                mostrado=mostrado[mask]
            if mostrado.empty:
                st.info('No se encontraron registros para mostrar.')
                continue
            vista=mostrado[['id','item','descripcion','lote','destino','numero_muestras','numero_corrugado','responsable','observaciones']].rename(columns={'id':'N°','item':'ITEM','descripcion':'Descripción','lote':'Lote','destino':'Destino','numero_muestras':'# De muestras','numero_corrugado':'# Corrugado','responsable':'Responsable','observaciones':'Observaciones'})
            nonce_key=f'table_nonce_{tabla}'
            nonce=st.session_state.get(nonce_key,0)
            evento=st.dataframe(vista,use_container_width=True,hide_index=True,on_select='rerun',selection_mode='single-row',key=f'table_{tabla}_{nonce}')
            st.download_button('Descargar CSV',vista.to_csv(index=False).encode('utf-8-sig'),f'{tabla}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv','text/csv',key=f'download_{tabla}')
            filas=getattr(evento,'selection',{}).get('rows',[]) if evento is not None else []
            posicion_valida=bool(filas) and isinstance(filas[0],int) and 0 <= filas[0] < len(vista)
            if not posicion_valida:
                if filas:
                    st.session_state[nonce_key]=nonce+1
                    st.rerun()
                continue
            rid=int(vista.iloc[filas[0]]['N°'])
            coincidencia=mostrado[mostrado['id']==rid]
            if coincidencia.empty:
                st.session_state[nonce_key]=nonce+1
                st.rerun()
            st.markdown(f'### Registro seleccionado: Número {rid}')
            original_df=read_df(f'SELECT * FROM {tabla} WHERE id=?',(rid,))
            if original_df.empty:
                st.session_state[nonce_key]=nonce+1
                st.rerun()
            original=original_df.iloc[0].to_dict()
            with st.expander(f'✏️ Editar registro N° {rid}',expanded=True):
                actual=next((x for x in opciones_item if x and x.split('|')[0].strip()==str(original['item'])),'')
                elegido=st.selectbox('ITEM *',opciones_item,index=idx_or_zero(opciones_item,actual),key=f'edit_item_{tabla}_{rid}')
                item=elegido.split('|')[0].strip() if elegido else ''
                fila=productos[productos['item'].astype(str)==item]
                descripcion=str(fila.iloc[0]['descripcion']) if not fila.empty else ''
                st.text_input('Descripción',value=descripcion,disabled=True,key=f'edit_desc_{tabla}_{rid}_{item}')
                with st.form(f'edit_form_{tabla}_{rid}'):
                    c1,c2=st.columns(2)
                    lote=c1.text_input('Lote *',value=str(original['lote'] or ''))
                    destinos=['','Nacional','Exportación']; dest=str(original['destino'] or '')
                    destino=c2.selectbox('Destino *',destinos,index=idx_or_zero(destinos,dest))
                    c3,c4=st.columns(2)
                    muestras=c3.number_input('# De muestras *',min_value=0.0,step=1.0,format='%.2f',value=float(original['numero_muestras'] or 0))
                    corrugado=c4.number_input('# Corrugado *',min_value=0.0,step=1.0,format='%.2f',value=float(original['numero_corrugado'] or 0))
                    analistas=opt_blank(catalog('analista')); resp=str(original['responsable'] or '')
                    if resp and resp not in analistas: analistas.append(resp)
                    responsable=st.selectbox('Responsable *',analistas,index=idx_or_zero(analistas,resp))
                    observaciones=st.text_area('Observaciones',value=str(original['observaciones'] or ''))
                    actualizar=st.form_submit_button('Guardar cambios',type='primary')
                if actualizar:
                    req={'ITEM':item,'Descripción':descripcion,'Lote':lote,'Destino':destino,'# De muestras':muestras,'# Corrugado':corrugado,'Responsable':responsable}
                    faltan=[k for k,v in req.items() if v is None or (isinstance(v,str) and not v.strip()) or (k in ['# De muestras','# Corrugado'] and float(v)<=0)]
                    if faltan: st.error('Completa los siguientes campos obligatorios: '+', '.join(faltan)+'.')
                    else:
                        exec_sql(f'UPDATE {tabla} SET item=?,descripcion=?,lote=?,destino=?,numero_muestras=?,numero_corrugado=?,responsable=?,observaciones=?,actualizado_por=?,actualizado_en=? WHERE id=?',(item,descripcion,lote.strip(),destino,float(muestras),float(corrugado),responsable,observaciones.strip(),st.session_state.auth['usuario'],now_iso(),rid))
                        audit(st.session_state.auth['usuario'],'EDITAR_MUESTRA_RETENCION',f'{nombre} | ID {rid}')
                        st.success(f'Registro actualizado correctamente. Número {rid}'); st.rerun()
            if is_dev():
                if st.button('Eliminar registro',key=f'del_{tabla}_{rid}'):
                    st.session_state[f'confirm_{tabla}']=rid
                if st.session_state.get(f'confirm_{tabla}')==rid:
                    st.warning(f'El registro número {rid} se eliminará de forma permanente. Esta acción no se puede deshacer.')
                    d1,d2=st.columns(2)
                    if d1.button('Confirmar eliminación',key=f'ok_{tabla}_{rid}'):
                        exec_sql(f'DELETE FROM {tabla} WHERE id=?',(rid,)); reset_autoincrement(tabla)
                        audit(st.session_state.auth['usuario'],'ELIMINAR_MUESTRA_RETENCION',f'{nombre} | ID {rid}')
                        st.session_state.pop(f'confirm_{tabla}',None); st.session_state[nonce_key]=st.session_state.get(nonce_key,0)+1; st.rerun()
                    if d2.button('Cancelar',key=f'cancel_{tabla}_{rid}'):
                        st.session_state.pop(f'confirm_{tabla}',None); st.session_state[nonce_key]=st.session_state.get(nonce_key,0)+1; st.rerun()

def page_entrega_turno():
    referencia='RE-CAL01-2301-00002-2013 Rev. 1'
    grupos=formato_entrega('Nave 1','PROCESO')
    if 'entrega_nave' not in st.session_state: st.session_state.entrega_nave=None
    if 'entrega_nonce' not in st.session_state: st.session_state.entrega_nonce=0
    if st.session_state.entrega_nave is None:
        st.markdown('<div class="registro-landing-hero"><div class="registro-landing-title">Entrega de turno</div><div class="registro-landing-subtitle">Selecciona la nave correspondiente para iniciar el registro.</div></div>',unsafe_allow_html=True)
        cols=st.columns(3,gap='large')
        for col,nave in zip(cols,['Nave 1','Nave 2','Nave 3']):
            with col:
                st.markdown('<span class="registro-card-slot"></span>',unsafe_allow_html=True)
                if st.button(f'🏭  **{nave}**\n\nEntrega y continuidad de actividades.\n\n*Abrir registro*',key=f'entrega_{nave}'):
                    st.session_state.entrega_nave=nave; st.rerun()
        return
    nave=st.session_state.entrega_nave
    if st.button('← Cambiar nave',key='entrega_volver'): st.session_state.entrega_nave=None; st.rerun()
    if nave in ['Nave 2','Nave 3']:
        n=st.session_state.entrega_nonce
        st.markdown(f'<div class="registro-full-panel"><div class="registro-pill">{referencia}</div><div class="registro-full-title">ENTREGA DE TURNO CALIDAD PROCESOS</div><div class="registro-full-subtitle">NAVE 2 Y 3 · Captura correspondiente a {nave}</div>',unsafe_allow_html=True)
        a,b,c=st.columns([2,1,1])
        analista=a.selectbox('ANALISTA *',opt_blank(catalog('analista')),key=f'et23_a_{nave}_{n}')
        fecha=b.date_input('Fecha *',date.today(),key=f'et23_f_{nave}_{n}')
        turno=c.selectbox('Turno *',opt_blank(catalog('turno')),key=f'et23_t_{nave}_{n}')
        catalogo_procesos={}
        orden_nave2=[]
        orden_nave3=[]
        procesos=formato_entrega(nave,'PROCESO')
        color_lineas='#0070C0' if nave=='Nave 2' else '#00A651'
        filas=[]
        st.markdown('### Líneas y procesos')
        st.caption('Producto es un solo campo libre dentro de cada proceso. Los controles − y + avanzan de uno en uno y permiten escribir decimales.')
        for gi,(grupo,lineas) in enumerate(procesos.items()):
            st.markdown(f'<div style="background:{color_lineas};color:white;padding:.55rem .8rem;border-radius:12px 12px 0 0;font-weight:900">{grupo}</div>',unsafe_allow_html=True)
            producto=st.text_input(f'Productos / Item / descripcion de la linea {grupo}',key=f'et23_prod_{nave}_{gi}_{n}',placeholder='Escritura libre')
            h=st.columns([2.2,1,1,2.2]);h[0].markdown('**Sector**');h[1].markdown('**Horas / cargas liberadas**');h[2].markdown('**Carga al SPAC**');h[3].markdown('**Observaciones**')
            for li,linea in enumerate(lineas):
                cols=st.columns([2.2,1,1,2.2])
                cols[0].text_input('Sector',linea,disabled=True,key=f'et23_l_{nave}_{gi}_{li}_{n}',label_visibility='collapsed')
                horas=cols[1].number_input('Horas',0.0,step=1.0,format='%.2f',key=f'et23_h_{nave}_{gi}_{li}_{n}',label_visibility='collapsed')
                carga=cols[2].number_input('Carga',0.0,step=1.0,format='%.2f',key=f'et23_c_{nave}_{gi}_{li}_{n}',label_visibility='collapsed')
                obs=cols[3].text_input('Observaciones',key=f'et23_o_{nave}_{gi}_{li}_{n}',label_visibility='collapsed')
                filas.append((grupo,linea,producto,float(horas),float(carga),obs,len(filas)))
            st.markdown('<div style="height:.7rem"></div>',unsafe_allow_html=True)
        st.markdown('### Análisis de laboratorio')
        analisis=formato_entrega(nave,'ANALISIS')
        resultados=[]
        for ai,(grupo,linea,tipo_analisis) in enumerate(analisis):
            cc=st.columns([1.2,2,1.5,1,2]);cc[0].text_input('Grupo',grupo,disabled=True,key=f'et23_ag_{nave}_{ai}_{n}',label_visibility='collapsed');cc[1].text_input('Línea',linea,disabled=True,key=f'et23_al_{nave}_{ai}_{n}',label_visibility='collapsed');cc[2].text_input('Análisis',tipo_analisis,disabled=True,key=f'et23_at_{nave}_{ai}_{n}',label_visibility='collapsed');resultado=cc[3].text_input('Resultado',key=f'et23_ar_{nave}_{ai}_{n}',label_visibility='collapsed');obs=cc[4].text_input('Observaciones',key=f'et23_ao_{nave}_{ai}_{n}',label_visibility='collapsed');contador=st.number_input(f'Cantidad de análisis - {linea} - {tipo_analisis}',min_value=0.0,step=1.0,format='%.2f',key=f'et23_cnt_{nave}_{ai}_{n}');resultados.append((grupo,linea,tipo_analisis,resultado,obs,float(contador)))
        seguimientos=[];st.markdown('### Seguimientos')
        for bi,bloque_cfg in enumerate(configuracion_seguimientos()):
            bloque=bloque_cfg['nombre']
            with st.expander(bloque,expanded=bi<2):
                ed,meta=editor_seguimiento_dinamico(bloque_cfg,f'et23_s_{nave}_{bi}_{n}')
                seguimientos.append((bloque,ed,meta))
        totales_nave,filas_clasificadas=clasificar_filas(filas);total_h=sum(totales_nave.values());total_c=sum(x[4] for x in filas)+sum(x[5] for x in resultados)
        m1,m2=st.columns(2);m1.metric('TOTAL DE CARGA DE DATOS',f'{total_c:.2f}');m2.metric('TOTAL GENERAL DE HORAS',f'{total_h:.2f}');q1,q2,q3=st.columns(3);q1.metric('HORAS NAVE 1',f"{totales_nave['Nave 1']:.2f}");q2.metric('HORAS NAVE 2',f"{totales_nave['Nave 2']:.2f}");q3.metric('HORAS NAVE 3',f"{totales_nave['Nave 3']:.2f}")
        if st.button('Guardar entrega de turno',type='primary',key=f'et23_g_{nave}_{n}'):
            errores_seguimiento=[]
            for bloque,df_seg,meta in seguimientos:
                for _,row_seg in df_seg.iterrows():
                    datos_seg={str(col):('' if pd.isna(val) else str(val).strip()) for col,val in row_seg.to_dict().items()}
                    if any(datos_seg.values()):
                        faltan_seg=[nombre for nombre,cfg in meta.items() if cfg.get('obligatorio') and not datos_seg.get(nombre,'')]
                        if faltan_seg: errores_seguimiento.append(f'{bloque}: '+', '.join(faltan_seg))
            repetida=not read_df('SELECT id FROM matriz_entrega WHERE fecha=? AND analista=?',(fecha.isoformat(),analista)).empty
            faltan=[]
            if not analista:faltan.append('Analista')
            if not turno:faltan.append('Turno')
            if errores_seguimiento: st.error('No se guardó la entrega. Completa los campos obligatorios de seguimiento: '+'; '.join(errores_seguimiento)+'.')
            elif repetida:st.error('Registro repetido. Ya existe una entrega de turno para el mismo analista y la misma fecha. Corrige la fecha, selecciona otro analista o elimina previamente ese registro desde la matriz.')
            elif faltan:st.error('Completa los campos obligatorios: '+', '.join(faltan)+'.')
            else:
                eid=exec_sql('INSERT INTO entregas_turno(nave,fecha,analista,turno,referencia,total_carga_datos,total_horas_trabajadas,creado_por,creado_en) VALUES(?,?,?,?,?,?,?,?,?)',(nave,fecha.isoformat(),analista,turno,referencia,total_c,total_h,st.session_state.auth['usuario'],now_iso()))
                for grupo,linea,producto,horas,carga,obs,orden,nvcat in filas_clasificadas:
                    if producto.strip() or horas or carga or obs.strip():exec_sql('INSERT INTO entregas_turno_lineas(entrega_id,grupo,linea,producto_descripcion,horas_trabajadas,carga_spac,observaciones,orden_fila,nave_catalogo) VALUES(?,?,?,?,?,?,?,?,?)',(eid,grupo,linea,producto.strip(),horas,carga,obs.strip(),orden,nvcat))
                base_orden=len(filas)
                for j,(grupo,linea,tipo_analisis,resultado,obs,contador) in enumerate(resultados):
                    if resultado.strip() or obs.strip():exec_sql('INSERT INTO entregas_turno_lineas(entrega_id,grupo,linea,producto_descripcion,horas_trabajadas,carga_spac,observaciones,orden_fila) VALUES(?,?,?,?,?,?,?,?)',(eid,'ANÁLISIS '+grupo,linea,tipo_analisis,0,contador,(resultado+' | '+obs).strip(' |'),base_orden+j))
                for bloque,df,meta in seguimientos:
                    for orden,row in df.iterrows():
                        datos={str(col):('' if pd.isna(val) else str(val)) for col,val in row.to_dict().items()}
                        if any(v.strip() for v in datos.values()):
                            obligatorios=[nombre for nombre,cfg in meta.items() if cfg.get('obligatorio') and not datos.get(nombre,'').strip()]
                            if obligatorios: st.warning(f'Seguimiento {bloque}: faltan campos obligatorios: '+', '.join(obligatorios)); continue
                            legado=[datos.get(c,'') for c in ['Registro #','Hoja física','Carga electrónica','Correo','Descripción del seguimiento']]
                            exec_sql('INSERT INTO entregas_turno_seguimientos(entrega_id,bloque,registro_numero,hoja_fisica,carga_electronica,correo,descripcion_seguimiento,datos_json,orden_fila) VALUES(?,?,?,?,?,?,?,?,?)',(eid,bloque,*legado,json.dumps(datos,ensure_ascii=False),int(orden)))
                guardar_matriz(eid,fecha.isoformat(),analista,total_c,totales_nave);audit(st.session_state.auth['usuario'],'CREAR_ENTREGA_TURNO',f'{nave} | ID {eid}');st.session_state.entrega_nonce+=1;st.success(f'Entrega de turno guardada correctamente. Número {eid}');st.rerun()
        st.markdown('</div>',unsafe_allow_html=True)
        return
    n=st.session_state.entrega_nonce
    st.markdown(f'<div class="registro-full-panel"><div class="registro-pill">{referencia}</div><div class="registro-full-title">ENTREGA DE TURNO CALIDAD PROCESOS</div><div class="registro-full-subtitle">NAVE 1</div>',unsafe_allow_html=True)
    a,b,c=st.columns([2,1,1]); analista=a.selectbox('ANALISTA *',opt_blank(catalog('analista')),key=f'et_a_{n}'); fecha=b.date_input('Fecha *',date.today(),key=f'et_f_{n}'); turno=c.selectbox('Turno *',opt_blank(catalog('turno')),key=f'et_t_{n}')
    st.markdown('### Líneas y procesos')
    st.caption('Producto es un único campo libre por proceso. Los botones − y + cambian horas y carga de uno en uno; también puedes escribir decimales directamente.')
    filas=[]
    for gi,(grupo,lineas) in enumerate(grupos.items()):
        st.markdown(f'<div style="background:#062C36;color:white;padding:.55rem .8rem;border-radius:12px 12px 0 0;font-weight:900">{grupo}</div>',unsafe_allow_html=True)
        producto=st.text_input(f'Producto / Item / descripcion de la linea {grupo}',key=f'et_prod_{gi}_{n}',placeholder='Escritura libre, no ligado a catálogos')
        h=st.columns([2.2,1,1,2.2]); h[0].markdown('**Sector**');h[1].markdown('**Horas trabajadas**');h[2].markdown('**Carga al SPAC**');h[3].markdown('**Observaciones**')
        for li,linea in enumerate(lineas):
            cols=st.columns([2.2,1,1,2.2])
            cols[0].text_input('Sector',linea,disabled=True,key=f'et_l_{gi}_{li}_{n}',label_visibility='collapsed')
            horas=cols[1].number_input('Horas',min_value=0.0,step=1.0,format='%.2f',key=f'et_h_{gi}_{li}_{n}',label_visibility='collapsed')
            carga=cols[2].number_input('Carga',min_value=0.0,step=1.0,format='%.2f',key=f'et_c_{gi}_{li}_{n}',label_visibility='collapsed')
            obs=cols[3].text_input('Observaciones',key=f'et_o_{gi}_{li}_{n}',label_visibility='collapsed')
            filas.append((grupo,linea,producto,float(horas),float(carga),obs,len(filas)))
        st.markdown('<div style="height:.7rem"></div>',unsafe_allow_html=True)
    seguimientos=[]; st.markdown('### Seguimientos')
    for bi,bloque_cfg in enumerate(configuracion_seguimientos()):
        bloque=bloque_cfg['nombre']
        with st.expander(bloque,expanded=bi<2):
            ed,meta=editor_seguimiento_dinamico(bloque_cfg,f'et_s_{bi}_{n}')
            seguimientos.append((bloque,ed,meta))
    totales_nave,filas_clasificadas=clasificar_filas(filas);total_h=sum(totales_nave.values());total_c=sum(x[4] for x in filas)
    m1,m2=st.columns(2);m1.metric('TOTAL DE CARGA DE DATOS',f'{total_c:.2f}');m2.metric('TOTAL GENERAL DE HORAS',f'{total_h:.2f}');q1,q2,q3=st.columns(3);q1.metric('HORAS NAVE 1',f"{totales_nave['Nave 1']:.2f}");q2.metric('HORAS NAVE 2',f"{totales_nave['Nave 2']:.2f}");q3.metric('HORAS NAVE 3',f"{totales_nave['Nave 3']:.2f}")
    if st.button('Guardar entrega de turno',type='primary',key=f'et_g_{n}'):
        errores_seguimiento=[]
        for bloque,df_seg,meta in seguimientos:
            for _,row_seg in df_seg.iterrows():
                datos_seg={str(col):('' if pd.isna(val) else str(val).strip()) for col,val in row_seg.to_dict().items()}
                if any(datos_seg.values()):
                    faltan_seg=[nombre for nombre,cfg in meta.items() if cfg.get('obligatorio') and not datos_seg.get(nombre,'')]
                    if faltan_seg: errores_seguimiento.append(f'{bloque}: '+', '.join(faltan_seg))
        repetida=not read_df('SELECT id FROM matriz_entrega WHERE fecha=? AND analista=?',(fecha.isoformat(),analista)).empty
        faltan=[]
        if not analista: faltan.append('Analista')
        if not turno: faltan.append('Turno')
        if errores_seguimiento: st.error('No se guardó la entrega. Completa los campos obligatorios de seguimiento: '+'; '.join(errores_seguimiento)+'.')
        elif repetida:st.error('Registro repetido. Ya existe una entrega de turno para el mismo analista y la misma fecha. Corrige la fecha, selecciona otro analista o elimina previamente ese registro desde la matriz.')
        elif faltan: st.error('Completa los campos obligatorios: '+', '.join(faltan)+'.')
        else:
            eid=exec_sql('INSERT INTO entregas_turno(nave,fecha,analista,turno,referencia,total_carga_datos,total_horas_trabajadas,creado_por,creado_en) VALUES(?,?,?,?,?,?,?,?,?)',(nave,fecha.isoformat(),analista,turno,referencia,total_c,total_h,st.session_state.auth['usuario'],now_iso()))
            for grupo,linea,producto,horas,carga,obs,orden,nvcat in filas_clasificadas:
                if producto.strip() or horas or carga or obs.strip(): exec_sql('INSERT INTO entregas_turno_lineas(entrega_id,grupo,linea,producto_descripcion,horas_trabajadas,carga_spac,observaciones,orden_fila,nave_catalogo) VALUES(?,?,?,?,?,?,?,?,?)',(eid,grupo,linea,producto.strip(),horas,carga,obs.strip(),orden,nvcat))
            for bloque,df,meta in seguimientos:
                for orden,row in df.iterrows():
                    datos={str(col):('' if pd.isna(val) else str(val)) for col,val in row.to_dict().items()}
                    if any(v.strip() for v in datos.values()):
                        obligatorios=[nombre for nombre,cfg in meta.items() if cfg.get('obligatorio') and not datos.get(nombre,'').strip()]
                        if obligatorios: st.warning(f'Seguimiento {bloque}: faltan campos obligatorios: '+', '.join(obligatorios)); continue
                        legado=[datos.get(c,'') for c in ['Registro #','Hoja física','Carga electrónica','Correo','Descripción del seguimiento']]
                        exec_sql('INSERT INTO entregas_turno_seguimientos(entrega_id,bloque,registro_numero,hoja_fisica,carga_electronica,correo,descripcion_seguimiento,datos_json,orden_fila) VALUES(?,?,?,?,?,?,?,?,?)',(eid,bloque,*legado,json.dumps(datos,ensure_ascii=False),int(orden)))
            guardar_matriz(eid,fecha.isoformat(),analista,total_c,totales_nave);audit(st.session_state.auth['usuario'],'CREAR_ENTREGA_TURNO',f'Nave 1 | ID {eid}');st.session_state.entrega_nonce+=1;st.success(f'Entrega de turno guardada correctamente. Número {eid}');st.rerun()
    st.markdown('</div>',unsafe_allow_html=True)

def admin_required():
    if not is_dev(): st.warning('Solo el usuario administrador puede modificar catálogos.'); return False
    return True


def page_catalogos():
    # El rol usuario dispone únicamente del catálogo de Productos y solo puede
    # consultar/agregar. La edición, eliminación y los demás catálogos continúan
    # reservados al rol desarrollador.
    if not is_dev():
        st.title('Catálogo de Productos')
        st.caption('Consulta los productos disponibles y agrega nuevos elementos al catálogo.')
        productos=read_df('SELECT id,item,descripcion,cliente,familia FROM productos WHERE activo=1 ORDER BY descripcion')
        vista=productos.rename(columns={'id':'ID','item':'ITEM','descripcion':'Descripción','cliente':'Cliente','familia':'Familia'})
        if vista.empty:
            st.info('Todavía no hay productos activos en el catálogo.')
        else:
            st.dataframe(vista,use_container_width=True,hide_index=True)
        with st.expander('Agregar producto',expanded=True):
            with st.form('usuario_agregar_producto',clear_on_submit=True):
                a,b=st.columns(2)
                item=a.text_input('ITEM *')
                descripcion=b.text_input('Descripción *')
                c,d=st.columns(2)
                cliente=c.text_input('Cliente *')
                familia=d.text_input('Familia *')
                agregar=st.form_submit_button('Agregar producto',type='primary')
            if agregar:
                valores=[item.strip(),descripcion.strip(),cliente.strip(),familia.strip()]
                etiquetas=['ITEM','Descripción','Cliente','Familia']
                faltantes=[etiquetas[i] for i,v in enumerate(valores) if not v]
                existente=read_df('SELECT id,activo FROM productos WHERE item=?',(valores[0],)) if valores[0] else pd.DataFrame()
                if faltantes:
                    st.error('Completa los campos obligatorios: '+', '.join(faltantes)+'.')
                elif not existente.empty and int(existente.iloc[0].get('activo',1) or 0)==1:
                    st.error('No fue posible agregar el producto porque el ITEM ya existe.')
                elif not existente.empty:
                    exec_sql('UPDATE productos SET descripcion=?,cliente=?,familia=?,activo=1 WHERE id=?',(valores[1],valores[2],valores[3],int(existente.iloc[0].id)))
                    audit(st.session_state.auth['usuario'],'REACTIVAR_PRODUCTO',f'ITEM {valores[0]}')
                    st.success('Producto agregado correctamente.'); st.rerun()
                else:
                    exec_sql('INSERT INTO productos(item,descripcion,cliente,familia,activo) VALUES(?,?,?,?,1)',tuple(valores))
                    audit(st.session_state.auth['usuario'],'AGREGAR_PRODUCTO',f'ITEM {valores[0]}')
                    st.success('Producto agregado correctamente.'); st.rerun()
        return
    st.title('Catálogos'); st.caption('Datos precargados desde el Excel adjunto. El administrador puede agregar o eliminar elementos.')
    tab1,tab2,tab3,tab4,tab5=st.tabs(['Productos','Defectos','Datos generales','Naves, líneas y sectores','Formatos entrega de turno'])
    def catalogo_seleccionable(tabla, consulta, columnas_vista, clave, titulo_singular, campos, insertar_sql, actualizar_sql, duplicado_sql, valores_opciones=None):
        df=read_df(consulta)
        vista=df.rename(columns=columnas_vista)
        nonce_key=f'cat_{clave}_nonce'
        nonce=st.session_state.get(nonce_key,0)
        evento=st.dataframe(vista,use_container_width=True,hide_index=True,on_select='rerun',selection_mode='single-row',key=f'cat_{clave}_tabla_{nonce}') if not vista.empty else None
        st.download_button(f'Descargar tabla de {titulo_singular.lower()}s',vista.to_csv(index=False).encode('utf-8-sig'),f'catalogo_{clave}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv','text/csv',key=f'cat_{clave}_csv')
        filas=getattr(evento,'selection',{}).get('rows',[]) if evento is not None else []
        valido=bool(filas) and isinstance(filas[0],int) and 0<=filas[0]<len(vista)
        rid=int(vista.iloc[filas[0]]['ID']) if valido else None
        with st.expander(f'Agregar {titulo_singular.lower()}',expanded=False):
            with st.form(f'cat_{clave}_agregar',clear_on_submit=True):
                valores={}
                cols=st.columns(min(4,len(campos)))
                for i,(campo,etiqueta) in enumerate(campos):
                    opciones=(valores_opciones or {}).get(campo)
                    valores[campo]=cols[i%len(cols)].selectbox(etiqueta,opciones) if opciones else cols[i%len(cols)].text_input(etiqueta)
                agregar=st.form_submit_button(f'Agregar {titulo_singular.lower()}',type='primary')
            if agregar:
                limpios={k:(v.strip() if isinstance(v,str) else v) for k,v in valores.items()}
                obligatorios=[etiqueta.replace(' *','') for campo,etiqueta in campos if etiqueta.endswith('*') and not str(limpios.get(campo,'')).strip()]
                if obligatorios: st.error('Completa los campos obligatorios: '+', '.join(obligatorios)+'.')
                elif not read_df(duplicado_sql,tuple(limpios[k] for k,_ in campos)).empty: st.error(f'No fue posible agregar el {titulo_singular.lower()} porque ya existe un registro con los mismos datos.')
                else:
                    exec_sql(insertar_sql,tuple(limpios[k] for k,_ in campos)); audit(st.session_state.auth['usuario'],f'AGREGAR_{clave.upper()}',str(limpios)); st.session_state[nonce_key]=nonce+1; st.success(f'{titulo_singular} agregado correctamente.'); st.rerun()
        if rid:
            actual=read_df(f'SELECT * FROM {tabla} WHERE id=? AND activo=1',(rid,))
            if not actual.empty:
                r=actual.iloc[0]
                with st.expander(f'Editar {titulo_singular.lower()} seleccionado',expanded=True):
                    with st.form(f'cat_{clave}_editar_{rid}'):
                        editados={}; cols=st.columns(min(4,len(campos)))
                        for i,(campo,etiqueta) in enumerate(campos):
                            valor=str(r[campo] or '')
                            opciones=(valores_opciones or {}).get(campo)
                            if opciones:
                                opciones=list(opciones)
                                if valor not in opciones: opciones.append(valor)
                                editados[campo]=cols[i%len(cols)].selectbox(etiqueta,opciones,index=idx_or_zero(opciones,valor))
                            else: editados[campo]=cols[i%len(cols)].text_input(etiqueta,valor)
                        guardar=st.form_submit_button('Guardar cambios',type='primary')
                    if guardar:
                        limpios={k:(v.strip() if isinstance(v,str) else v) for k,v in editados.items()}
                        obligatorios=[etiqueta.replace(' *','') for campo,etiqueta in campos if etiqueta.endswith('*') and not str(limpios.get(campo,'')).strip()]
                        params=tuple(limpios[k] for k,_ in campos)
                        if obligatorios: st.error('Completa los campos obligatorios: '+', '.join(obligatorios)+'.')
                        elif not read_df(duplicado_sql+' AND id<>?',params+(rid,)).empty: st.error(f'Ya existe otro {titulo_singular.lower()} activo con los mismos datos.')
                        else:
                            exec_sql(actualizar_sql,params+(rid,)); audit(st.session_state.auth['usuario'],f'EDITAR_{clave.upper()}',f'ID {rid}'); st.session_state[nonce_key]=nonce+1; st.success(f'{titulo_singular} actualizado correctamente.'); st.rerun()
                if st.button(f'Eliminar {titulo_singular.lower()}',key=f'cat_{clave}_eliminar_{rid}'): st.session_state[f'cat_{clave}_confirmar']=rid
                if st.session_state.get(f'cat_{clave}_confirmar')==rid:
                    st.warning(f'El {titulo_singular.lower()} dejará de estar disponible para nuevos registros. Los datos históricos se conservarán.')
                    x,y=st.columns(2)
                    if x.button('Confirmar eliminación',key=f'cat_{clave}_ok_{rid}'):
                        exec_sql(f'UPDATE {tabla} SET activo=0 WHERE id=?',(rid,)); audit(st.session_state.auth['usuario'],f'ELIMINAR_{clave.upper()}',f'ID {rid}'); st.session_state.pop(f'cat_{clave}_confirmar',None); st.session_state[nonce_key]=nonce+1; st.rerun()
                    if y.button('Cancelar',key=f'cat_{clave}_cancelar_{rid}'): st.session_state.pop(f'cat_{clave}_confirmar',None); st.session_state[nonce_key]=nonce+1; st.rerun()

    with tab1:
        st.subheader('Productos')
        st.caption('Selecciona una fila para editarla o eliminarla. Los cambios se aplican automáticamente en los formularios relacionados.')
        catalogo_seleccionable('productos','SELECT id,item,descripcion,cliente,familia FROM productos WHERE activo=1 ORDER BY descripcion',{'id':'ID','item':'ITEM','descripcion':'Descripción','cliente':'Cliente','familia':'Familia'},'productos','Producto',[('item','ITEM *'),('descripcion','Descripción *'),('cliente','Cliente *'),('familia','Familia *')],'INSERT INTO productos(item,descripcion,cliente,familia,activo) VALUES(?,?,?,?,1) ON CONFLICT(item) DO UPDATE SET descripcion=excluded.descripcion,cliente=excluded.cliente,familia=excluded.familia,activo=1','UPDATE productos SET item=?,descripcion=?,cliente=?,familia=? WHERE id=?','SELECT id FROM productos WHERE item=? AND descripcion=? AND cliente=? AND familia=? AND activo=1')
    with tab2:
        st.subheader('Defectos')
        st.caption('Selecciona una fila para editarla o eliminarla. Los cambios se aplican automáticamente en los formularios relacionados.')
        catalogo_seleccionable('defectos','SELECT id,codigo,defecto,tipo_defecto,clasificacion FROM defectos WHERE activo=1 ORDER BY CAST(codigo AS INTEGER)',{'id':'ID','codigo':'Código','defecto':'Defecto','tipo_defecto':'Tipo de defecto','clasificacion':'Clasificación'},'defectos','Defecto',[('codigo','Código *'),('defecto','Defecto *'),('tipo_defecto','Tipo de defecto *'),('clasificacion','Clasificación *')],'INSERT INTO defectos(codigo,defecto,tipo_defecto,clasificacion,activo) VALUES(?,?,?,?,1) ON CONFLICT(codigo) DO UPDATE SET defecto=excluded.defecto,tipo_defecto=excluded.tipo_defecto,clasificacion=excluded.clasificacion,activo=1','UPDATE defectos SET codigo=?,defecto=?,tipo_defecto=?,clasificacion=? WHERE id=?','SELECT id FROM defectos WHERE codigo=? AND defecto=? AND tipo_defecto=? AND clasificacion=? AND activo=1',{'tipo_defecto':catalog('tipo_defecto') or ['Funcional','Contaminación'],'clasificacion':['Inocuidad','Calidad','Salubridad','Legalidad']})
    with tab3:
        st.subheader('Datos generales')
        st.caption('Selecciona una fila para editarla o eliminarla. Los cambios se aplican automáticamente en los formularios relacionados.')
        etiquetas={'Supervisores':'supervisor','Analistas':'analista','Nave':'nave','Status':'status','Responsable de detectar PNC':'responsable_detecta','Etapa':'etapa','Línea':'linea_sector','Turno':'turno','Defecto':'tipo_defecto','Disposición':'disposicion'}
        inv={v:k for k,v in etiquetas.items()}
        raw=read_df('SELECT id,categoria,valor FROM catalogos WHERE activo=1 ORDER BY categoria,valor')
        raw['lista']=raw['categoria'].map(inv); raw=raw.dropna(subset=['lista'])
        vista=raw[['id','lista','valor']].rename(columns={'id':'ID','lista':'Lista','valor':'Valor'})
        nonce=st.session_state.get('cat_datos_nonce',0)
        evento=st.dataframe(vista,use_container_width=True,hide_index=True,on_select='rerun',selection_mode='single-row',key=f'cat_datos_tabla_{nonce}') if not vista.empty else None
        st.download_button('Descargar tabla de datos generales',vista.to_csv(index=False).encode('utf-8-sig'),f'datos_generales_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv','text/csv',key='cat_datos_csv')
        filas=getattr(evento,'selection',{}).get('rows',[]) if evento is not None else []
        rid=int(vista.iloc[filas[0]]['ID']) if filas and isinstance(filas[0],int) and 0<=filas[0]<len(vista) else None
        with st.expander('Agregar valor',expanded=False):
            with st.form('cat_datos_agregar',clear_on_submit=True):
                a,b=st.columns(2); lista=a.selectbox('Lista *',list(etiquetas)); valor=b.text_input('Valor *'); agregar=st.form_submit_button('Agregar valor',type='primary')
            if agregar:
                if not valor.strip(): st.error('Ingresa el valor que deseas agregar.')
                elif not read_df('SELECT id FROM catalogos WHERE categoria=? AND valor=? AND activo=1',(etiquetas[lista],valor.strip())).empty: st.error('El valor ya existe en la lista seleccionada.')
                else: exec_sql('INSERT INTO catalogos(categoria,valor,activo) VALUES(?,?,1) ON CONFLICT(categoria,valor) DO UPDATE SET activo=1',(etiquetas[lista],valor.strip())); audit(st.session_state.auth['usuario'],'AGREGAR_DATO_GENERAL',f'{lista} | {valor.strip()}'); st.session_state.cat_datos_nonce=nonce+1; st.rerun()
        if rid:
            r=raw[raw['id']==rid].iloc[0]
            with st.expander('Editar valor seleccionado',expanded=True):
                with st.form(f'cat_datos_editar_{rid}'):
                    a,b=st.columns(2); lista=a.selectbox('Lista *',list(etiquetas),index=idx_or_zero(list(etiquetas),str(r.lista))); valor=b.text_input('Valor *',str(r.valor)); guardar=st.form_submit_button('Guardar cambios',type='primary')
                if guardar:
                    cat=etiquetas[lista]; val=valor.strip(); dup=read_df('SELECT id FROM catalogos WHERE categoria=? AND valor=? AND id<>? AND activo=1',(cat,val,rid))
                    if not val: st.error('Ingresa un valor válido.')
                    elif not dup.empty: st.error('Ya existe otro valor igual en la lista seleccionada.')
                    else: exec_sql('UPDATE catalogos SET categoria=?,valor=? WHERE id=?',(cat,val,rid)); audit(st.session_state.auth['usuario'],'EDITAR_DATO_GENERAL',f'ID {rid}'); st.session_state.cat_datos_nonce=nonce+1; st.rerun()
            if st.button('Eliminar valor',key=f'cat_datos_eliminar_{rid}'): st.session_state.cat_datos_confirmar=rid
            if st.session_state.get('cat_datos_confirmar')==rid:
                st.warning('El valor dejará de estar disponible para nuevos registros. Los datos históricos se conservarán.')
                x,y=st.columns(2)
                if x.button('Confirmar eliminación',key=f'cat_datos_ok_{rid}'): exec_sql('UPDATE catalogos SET activo=0 WHERE id=?',(rid,)); audit(st.session_state.auth['usuario'],'ELIMINAR_DATO_GENERAL',f'ID {rid}'); st.session_state.pop('cat_datos_confirmar',None); st.session_state.cat_datos_nonce=nonce+1; st.rerun()
                if y.button('Cancelar',key=f'cat_datos_cancelar_{rid}'): st.session_state.pop('cat_datos_confirmar',None); st.session_state.cat_datos_nonce=nonce+1; st.rerun()
    with tab4:
        st.subheader('Naves, líneas y sectores')
        st.caption('Este catálogo clasifica cada Línea y Sector para calcular por separado las horas generales de Nave 1, Nave 2 y Nave 3. No modifica la estructura visible del formato.')
        f1,f2=st.columns([1,2])
        filtro_nave=f1.multiselect('Filtrar por nave',['Nave 1','Nave 2','Nave 3'],key='cat_general_naves')
        buscar=f2.text_input('Buscar línea o sector',key='cat_general_buscar')
        df_general=read_df('SELECT id,nave,linea,sector,orden FROM catalogo_naves_lineas WHERE activo=1 ORDER BY nave,orden,id')
        if filtro_nave:
            df_general=df_general[df_general['nave'].isin(filtro_nave)]
        if buscar.strip() and not df_general.empty:
            q=buscar.strip()
            df_general=df_general[df_general['linea'].fillna('').astype(str).str.contains(q,case=False,na=False)|df_general['sector'].fillna('').astype(str).str.contains(q,case=False,na=False)]
        vista_general=df_general.rename(columns={'id':'ID','nave':'Nave','linea':'Línea','sector':'Sector','orden':'Orden'})
        st.download_button('Descargar tabla de naves, líneas y sectores',vista_general.to_csv(index=False).encode('utf-8-sig'),f'catalogo_naves_lineas_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv','text/csv',key='cat_general_csv')
        nonce_general=st.session_state.get('cat_general_nonce',0)
        evento_general=st.dataframe(vista_general,use_container_width=True,hide_index=True,on_select='rerun',selection_mode='single-row',key=f'cat_general_tabla_{nonce_general}') if not vista_general.empty else None
        filas_general=getattr(evento_general,'selection',{}).get('rows',[]) if evento_general is not None else []
        posicion_general=bool(filas_general) and isinstance(filas_general[0],int) and 0<=filas_general[0]<len(vista_general)
        id_general=int(vista_general.iloc[filas_general[0]]['ID']) if posicion_general else None
        if admin_required():
            with st.expander('Agregar relación general',expanded=False):
                with st.form('cat_general_agregar',clear_on_submit=True):
                    a,bx,c=st.columns(3)
                    nueva_nave=a.selectbox('Nave *',['Nave 1','Nave 2','Nave 3'])
                    nueva_linea=bx.text_input('Línea *',placeholder='Ejemplo: BOB, BUTTER, CARAMELO')
                    nuevo_sector=c.text_input('Sector *',placeholder='Ejemplo: MD - BOB - CENTRO DE MASA')
                    agregar_general=st.form_submit_button('Agregar relación',type='primary')
                if agregar_general:
                    li=nueva_linea.strip();se=nuevo_sector.strip()
                    if not li or not se:
                        st.error('Completa Línea y Sector.')
                    else:
                        duplicado=read_df('SELECT id FROM catalogo_naves_lineas WHERE nave=? AND linea_norm=? AND sector_norm=? AND activo=1',(nueva_nave,normalizar_catalogo(li),normalizar_catalogo(se)))
                        if not duplicado.empty:
                            st.error('La relación ya existe en el catálogo general.')
                        else:
                            orden=int(read_df('SELECT COALESCE(MAX(orden),-1)+1 AS n FROM catalogo_naves_lineas WHERE nave=?',(nueva_nave,)).iloc[0]['n'])
                            exec_sql('INSERT INTO catalogo_naves_lineas(nave,linea,sector,linea_norm,sector_norm,orden,activo) VALUES(?,?,?,?,?,?,1)',(nueva_nave,li,se,normalizar_catalogo(li),normalizar_catalogo(se),orden))
                            audit(st.session_state.auth['usuario'],'AGREGAR_RELACION_NAVE',f'{nueva_nave} | {li} | {se}')
                            st.session_state.cat_general_nonce=nonce_general+1
                            st.rerun()
            if id_general:
                actual=read_df('SELECT * FROM catalogo_naves_lineas WHERE id=? AND activo=1',(id_general,))
                if not actual.empty:
                    r=actual.iloc[0]
                    with st.expander('Editar relación general seleccionada',expanded=True):
                        with st.form(f'cat_general_editar_{id_general}'):
                            naves=['Nave 1','Nave 2','Nave 3']
                            a,bx,c,d=st.columns([1,2,2,1])
                            nave_e=a.selectbox('Nave *',naves,index=idx_or_zero(naves,str(r.nave)))
                            linea_e=bx.text_input('Línea *',str(r.linea or ''))
                            sector_e=c.text_input('Sector *',str(r.sector or ''))
                            orden_e=d.number_input('Orden',0,value=int(r.orden or 0),step=1)
                            guardar_general=st.form_submit_button('Guardar cambios',type='primary')
                        if guardar_general:
                            li=linea_e.strip();se=sector_e.strip()
                            duplicado=read_df('SELECT id FROM catalogo_naves_lineas WHERE nave=? AND linea_norm=? AND sector_norm=? AND id<>? AND activo=1',(nave_e,normalizar_catalogo(li),normalizar_catalogo(se),id_general))
                            if not li or not se:
                                st.error('Completa Línea y Sector.')
                            elif not duplicado.empty:
                                st.error('Ya existe otra relación activa con esos datos.')
                            else:
                                exec_sql('UPDATE catalogo_naves_lineas SET nave=?,linea=?,sector=?,linea_norm=?,sector_norm=?,orden=? WHERE id=?',(nave_e,li,se,normalizar_catalogo(li),normalizar_catalogo(se),int(orden_e),id_general))
                                audit(st.session_state.auth['usuario'],'EDITAR_RELACION_NAVE',f'ID {id_general} | {nave_e} | {li} | {se}')
                                st.session_state.cat_general_nonce=nonce_general+1
                                st.rerun()
                    if st.button('Eliminar relación general',key=f'cat_general_eliminar_{id_general}'):
                        st.session_state.cat_general_confirmar=id_general
                    if st.session_state.get('cat_general_confirmar')==id_general:
                        st.warning('Esta relación dejará de utilizarse en las nuevas sumatorias generales por nave.')
                        x,y=st.columns(2)
                        if x.button('Sí, eliminar',key=f'cat_general_ok_{id_general}'):
                            exec_sql('UPDATE catalogo_naves_lineas SET activo=0 WHERE id=?',(id_general,))
                            audit(st.session_state.auth['usuario'],'ELIMINAR_RELACION_NAVE',f'ID {id_general}')
                            st.session_state.pop('cat_general_confirmar',None)
                            st.session_state.cat_general_nonce=nonce_general+1
                            st.rerun()
                        if y.button('Cancelar',key=f'cat_general_cancelar_{id_general}'):
                            st.session_state.pop('cat_general_confirmar',None)
                            st.session_state.cat_general_nonce=nonce_general+1
                            st.rerun()

    with tab5:
        st.subheader('Formatos de entrega de turno')
        st.caption('Este catálogo define la estructura visible de cada formato y el TOTAL DE CARGA DE DATOS individual del analista. Cada Línea se muestra como encabezado y sus Sectores como filas.')
        nave_fmt=st.radio('Formato',['Nave 1','Nave 2','Nave 3'],horizontal=True,key='fmt_nave')
        tipo_fmt=st.radio('Sección',['PROCESO','ANALISIS'],format_func=lambda x:'Líneas y sectores' if x=='PROCESO' else 'Análisis fisicoquímicos',horizontal=True,key='fmt_tipo')
        df=read_df('SELECT id,linea,sector,tipo_analisis,orden_linea,orden_sector FROM catalogo_formatos_entrega WHERE formato_nave=? AND tipo=? AND activo=1 ORDER BY orden_linea,orden_sector,id',(nave_fmt,tipo_fmt))
        vista=df.rename(columns={'id':'ID','linea':'Línea','sector':'Sector','tipo_analisis':'Tipo de análisis','orden_linea':'Orden línea','orden_sector':'Orden sector'})
        st.download_button('Descargar tabla del formato seleccionado',vista.to_csv(index=False).encode('utf-8-sig'),f'formato_entrega_{nave_fmt}_{tipo_fmt}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv','text/csv',key=f'fmt_csv_{nave_fmt}_{tipo_fmt}')
        nonce=st.session_state.get('fmt_nonce',0)
        evento=st.dataframe(vista,use_container_width=True,hide_index=True,on_select='rerun',selection_mode='single-row',key=f'fmt_table_{nave_fmt}_{tipo_fmt}_{nonce}') if not vista.empty else None
        rows=getattr(evento,'selection',{}).get('rows',[]) if evento is not None else []
        valido=bool(rows) and isinstance(rows[0],int) and 0<=rows[0]<len(vista)
        rid=int(vista.iloc[rows[0]]['ID']) if valido else None
        if admin_required():
            with st.expander('Agregar elemento al formato',expanded=False):
                with st.form(f'fmt_add_{nave_fmt}_{tipo_fmt}',clear_on_submit=True):
                    a,b=st.columns(2);linea=a.text_input('Línea *',placeholder='Ejemplo: BOB, BUTTER, CARAMELO');sector=b.text_input('Sector *',placeholder='Ejemplo: MD - BOB - CENTRO DE MASA')
                    c,d,e=st.columns(3);analisis=c.text_input('Tipo de análisis *' if tipo_fmt=='ANALISIS' else 'Tipo de análisis',disabled=tipo_fmt!='ANALISIS');ol=d.number_input('Orden línea',0,step=1);os=e.number_input('Orden sector',0,step=1);agregar=st.form_submit_button('Agregar',type='primary')
                if agregar:
                    faltan=not linea.strip() or not sector.strip() or (tipo_fmt=='ANALISIS' and not analisis.strip())
                    if faltan: st.error('Completa Línea, Sector y Tipo de análisis cuando corresponda.')
                    else:
                        existente=read_df('SELECT id FROM catalogo_formatos_entrega WHERE formato_nave=? AND tipo=? AND linea=? AND sector=? AND tipo_analisis=?',(nave_fmt,tipo_fmt,linea.strip(),sector.strip(),analisis.strip() if tipo_fmt=='ANALISIS' else ''))
                        if existente.empty: exec_sql('INSERT INTO catalogo_formatos_entrega(formato_nave,tipo,linea,sector,tipo_analisis,orden_linea,orden_sector,activo) VALUES(?,?,?,?,?,?,?,1)',(nave_fmt,tipo_fmt,linea.strip(),sector.strip(),analisis.strip() if tipo_fmt=='ANALISIS' else '',int(ol),int(os)))
                        else: exec_sql('UPDATE catalogo_formatos_entrega SET activo=1,orden_linea=?,orden_sector=? WHERE id=?',(int(ol),int(os),int(existente.iloc[0].id)))
                        audit(st.session_state.auth['usuario'],'AGREGAR_FORMATO_ENTREGA',f'{nave_fmt} | {tipo_fmt} | {linea} | {sector}');st.session_state.fmt_nonce=nonce+1;st.success('Elemento agregado. Ya aparecerá en el formato correspondiente.');st.rerun()
            if rid:
                r=read_df('SELECT * FROM catalogo_formatos_entrega WHERE id=? AND activo=1',(rid,))
                if not r.empty:
                    r=r.iloc[0]
                    with st.expander('Editar elemento seleccionado',expanded=True):
                        with st.form(f'fmt_edit_{rid}'):
                            a,b=st.columns(2);linea=a.text_input('Línea *',str(r.linea));sector=b.text_input('Sector *',str(r.sector));c,d,e=st.columns(3);analisis=c.text_input('Tipo de análisis',str(r.tipo_analisis or ''),disabled=tipo_fmt!='ANALISIS');ol=d.number_input('Orden línea',0,value=int(r.orden_linea));os=e.number_input('Orden sector',0,value=int(r.orden_sector));guardar=st.form_submit_button('Guardar cambios',type='primary')
                        if guardar:
                            linea_editada=linea.strip()
                            sector_editado=sector.strip()
                            analisis_editado=analisis.strip() if tipo_fmt=='ANALISIS' else ''
                            if not linea_editada or not sector_editado or (tipo_fmt=='ANALISIS' and not analisis_editado):
                                st.error('Completa Línea, Sector y Tipo de análisis cuando corresponda.')
                            else:
                                duplicado=read_df('''SELECT id FROM catalogo_formatos_entrega
                                    WHERE formato_nave=? AND tipo=? AND linea=? AND sector=?
                                    AND tipo_analisis=? AND id<>? AND activo=1 LIMIT 1''',
                                    (nave_fmt,tipo_fmt,linea_editada,sector_editado,analisis_editado,rid))
                                if not duplicado.empty:
                                    st.error('No se puede guardar porque ya existe otro elemento activo con la misma Nave, Línea, Sector y Tipo de análisis.')
                                else:
                                    try:
                                        exec_sql('''UPDATE catalogo_formatos_entrega
                                            SET linea=?,sector=?,tipo_analisis=?,orden_linea=?,orden_sector=?
                                            WHERE id=?''',
                                            (linea_editada,sector_editado,analisis_editado,int(ol),int(os),rid))
                                    except DatabaseIntegrityError:
                                        st.error('No se pudo guardar porque esa combinación ya existe en el catálogo. Modifica la Línea, el Sector o el Tipo de análisis.')
                                    else:
                                        audit(st.session_state.auth['usuario'],'EDITAR_FORMATO_ENTREGA',f'ID {rid} | {nave_fmt} | {tipo_fmt} | {linea_editada} | {sector_editado} | {analisis_editado}')
                                        st.session_state.fmt_nonce=nonce+1
                                        st.success('Elemento actualizado correctamente.')
                                        st.rerun()
                    if st.button('Eliminar del formato',key=f'fmt_del_{rid}'):st.session_state.fmt_confirm=rid
                    if st.session_state.get('fmt_confirm')==rid:
                        st.warning('El elemento dejará de aparecer en el formato correspondiente.')
                        x,y=st.columns(2)
                        if x.button('Sí, eliminar',key=f'fmt_ok_{rid}'):exec_sql('UPDATE catalogo_formatos_entrega SET activo=0 WHERE id=?',(rid,));audit(st.session_state.auth['usuario'],'ELIMINAR_FORMATO_ENTREGA',f'ID {rid}');st.session_state.pop('fmt_confirm',None);st.session_state.fmt_nonce=nonce+1;st.rerun()
                        if y.button('Cancelar',key=f'fmt_cancel_{rid}'):st.session_state.pop('fmt_confirm',None);st.session_state.fmt_nonce=nonce+1;st.rerun()

        st.markdown('---'); st.subheader('Seguimientos generales de entrega de turno')
        st.caption('Catálogo general para Nave 1, Nave 2 y Nave 3. Los cambios se reflejan automáticamente en los formatos nuevos.')
        seg=read_df('SELECT id,nombre,orden FROM catalogo_seguimientos_entrega WHERE activo=1 ORDER BY orden,id')
        vista_seg=seg.rename(columns={'id':'ID','nombre':'Seguimiento','orden':'Orden'}); ns=st.session_state.get('seg_nonce',0)
        ev=st.dataframe(vista_seg,use_container_width=True,hide_index=True,on_select='rerun',selection_mode='single-row',key=f'seg_tabla_{ns}') if not vista_seg.empty else None
        rs=getattr(ev,'selection',{}).get('rows',[]) if ev is not None else []; sid=int(vista_seg.iloc[rs[0]].ID) if rs and isinstance(rs[0],int) and 0<=rs[0]<len(vista_seg) else None
        with st.expander('Agregar seguimiento'):
            with st.form('seg_agregar',clear_on_submit=True):
                a,b=st.columns([3,1]); nom=a.text_input('Seguimiento *'); orden=b.number_input('Orden',0,step=1); alta=st.form_submit_button('Agregar',type='primary')
            if alta:
                nombre=nom.strip(); ex=read_df('SELECT id FROM catalogo_seguimientos_entrega WHERE UPPER(TRIM(nombre))=UPPER(TRIM(?))',(nombre,)) if nombre else pd.DataFrame()
                if not nombre: st.error('Ingresa el seguimiento.')
                elif not ex.empty: exec_sql('UPDATE catalogo_seguimientos_entrega SET nombre=?,orden=?,activo=1 WHERE id=?',(nombre,int(orden),int(ex.iloc[0].id)));audit(st.session_state.auth['usuario'],'REACTIVAR_SEGUIMIENTO_ENTREGA',nombre);st.session_state.seg_nonce=ns+1;st.rerun()
                else: exec_sql('INSERT INTO catalogo_seguimientos_entrega(nombre,orden,activo) VALUES(?,?,1)',(nombre,int(orden)));audit(st.session_state.auth['usuario'],'AGREGAR_SEGUIMIENTO_ENTREGA',nombre);st.session_state.seg_nonce=ns+1;st.rerun()
        if sid:
            rr=read_df('SELECT * FROM catalogo_seguimientos_entrega WHERE id=?',(sid,)).iloc[0]
            st.markdown('#### Campos configurables del seguimiento')
            st.caption('Define el nombre del campo, el tipo de captura y sus opciones. Solo administradores pueden modificar esta configuración.')
            campos_seg=read_df('SELECT id,nombre,tipo_campo,opciones,obligatorio,orden FROM catalogo_seguimientos_campos WHERE seguimiento_id=? AND activo=1 ORDER BY orden,id',(sid,))
            vista_campos=campos_seg.rename(columns={'id':'ID','nombre':'Nombre del campo','tipo_campo':'Tipo','opciones':'Opciones','obligatorio':'Obligatorio','orden':'Orden'})
            nc=st.session_state.get('seg_campos_nonce',0)
            ec=st.dataframe(vista_campos,use_container_width=True,hide_index=True,on_select='rerun',selection_mode='single-row',key=f'seg_campos_{sid}_{nc}') if not vista_campos.empty else None
            rc=getattr(ec,'selection',{}).get('rows',[]) if ec is not None else []
            cid=int(vista_campos.iloc[rc[0]].ID) if rc and isinstance(rc[0],int) and 0<=rc[0]<len(vista_campos) else None
            tipos_campo=['Texto','Texto largo','Número','Fecha','Lista desplegable','Sí / No / N/A']
            with st.expander('Agregar campo al seguimiento'):
                with st.form(f'seg_campo_add_{sid}',clear_on_submit=True):
                    a,b,c,d=st.columns([2,1.4,2,1]); cn=a.text_input('Nombre del campo *'); ct=b.selectbox('Tipo de campo *',tipos_campo); co=c.text_input('Opciones',help='Para lista desplegable, separa las opciones con |'); cob=d.checkbox('Obligatorio'); cord=st.number_input('Orden del campo',0,step=1); ca=st.form_submit_button('Agregar campo',type='primary')
                if ca:
                    nombre_campo=cn.strip(); opciones_campo=co.strip() if ct=='Lista desplegable' else ('Sí|No|N/A' if ct=='Sí / No / N/A' else '')
                    if not nombre_campo: st.error('Ingresa el nombre del campo.')
                    elif ct=='Lista desplegable' and not opciones_campo: st.error('Captura al menos una opción para la lista desplegable.')
                    else:
                        exec_sql('INSERT INTO catalogo_seguimientos_campos(seguimiento_id,nombre,tipo_campo,opciones,obligatorio,orden,activo) VALUES(?,?,?,?,?,?,1) ON CONFLICT(seguimiento_id,nombre) DO UPDATE SET tipo_campo=excluded.tipo_campo,opciones=excluded.opciones,obligatorio=excluded.obligatorio,orden=excluded.orden,activo=1',(sid,nombre_campo,ct,opciones_campo,int(cob),int(cord)))
                        audit(st.session_state.auth['usuario'],'AGREGAR_CAMPO_SEGUIMIENTO',f'Seguimiento {sid} | {nombre_campo}'); st.session_state.seg_campos_nonce=nc+1; st.rerun()
            if cid:
                cr=campos_seg[campos_seg.id==cid].iloc[0]
                with st.expander('Editar campo seleccionado',expanded=True):
                    with st.form(f'seg_campo_edit_{cid}'):
                        a,b,c,d=st.columns([2,1.4,2,1]); cen=a.text_input('Nombre del campo *',str(cr.nombre)); cet=b.selectbox('Tipo de campo *',tipos_campo,index=idx_or_zero(tipos_campo,str(cr.tipo_campo))); ceo=c.text_input('Opciones',str(cr.opciones or ''),help='Separa las opciones con |'); ceob=d.checkbox('Obligatorio',value=bool(cr.obligatorio)); ceord=st.number_input('Orden del campo',0,value=int(cr.orden or 0),step=1); ceg=st.form_submit_button('Guardar campo',type='primary')
                    if ceg:
                        opciones_edit=ceo.strip() if cet=='Lista desplegable' else ('Sí|No|N/A' if cet=='Sí / No / N/A' else '')
                        if not cen.strip(): st.error('Ingresa el nombre del campo.')
                        elif cet=='Lista desplegable' and not opciones_edit: st.error('Captura las opciones de la lista desplegable.')
                        else: exec_sql('UPDATE catalogo_seguimientos_campos SET nombre=?,tipo_campo=?,opciones=?,obligatorio=?,orden=? WHERE id=?',(cen.strip(),cet,opciones_edit,int(ceob),int(ceord),cid)); audit(st.session_state.auth['usuario'],'EDITAR_CAMPO_SEGUIMIENTO',f'ID {cid}'); st.session_state.seg_campos_nonce=nc+1; st.rerun()
                if st.button('Eliminar campo seleccionado',key=f'seg_campo_del_{cid}'):
                    exec_sql('UPDATE catalogo_seguimientos_campos SET activo=0 WHERE id=?',(cid,)); audit(st.session_state.auth['usuario'],'ELIMINAR_CAMPO_SEGUIMIENTO',f'ID {cid}'); st.session_state.seg_campos_nonce=nc+1; st.rerun()
            with st.expander('Editar seguimiento seleccionado',expanded=True):
                with st.form(f'seg_editar_{sid}'):
                    a,b=st.columns([3,1]); ne=a.text_input('Seguimiento *',str(rr.nombre)); oe=b.number_input('Orden',0,value=int(rr.orden or 0),step=1); guardar=st.form_submit_button('Guardar cambios',type='primary')
                if guardar and ne.strip(): exec_sql('UPDATE catalogo_seguimientos_entrega SET nombre=?,orden=? WHERE id=?',(ne.strip(),int(oe),sid));audit(st.session_state.auth['usuario'],'EDITAR_SEGUIMIENTO_ENTREGA',f'ID {sid}');st.session_state.seg_nonce=ns+1;st.rerun()
            if st.button('Eliminar seguimiento',key=f'seg_del_{sid}'): st.session_state.seg_confirm=sid
            if st.session_state.get('seg_confirm')==sid:
                st.warning('Dejará de aparecer en entregas nuevas. Los registros históricos se conservarán.')
                a,b=st.columns(2)
                if a.button('Confirmar eliminación',key=f'seg_ok_{sid}'): exec_sql('UPDATE catalogo_seguimientos_entrega SET activo=0 WHERE id=?',(sid,));audit(st.session_state.auth['usuario'],'ELIMINAR_SEGUIMIENTO_ENTREGA',f'ID {sid}');st.session_state.pop('seg_confirm',None);st.session_state.seg_nonce=ns+1;st.rerun()
                if b.button('Cancelar',key=f'seg_no_{sid}'): st.session_state.pop('seg_confirm',None);st.rerun()

def page_usuarios():
    if not is_dev():
        st.warning('Esta sección está disponible únicamente para administradores.')
        return
    st.title('Administración de usuarios')
    st.caption('Selecciona directamente un registro de la tabla para editar, habilitar, inhabilitar o eliminar la cuenta.')
    if 'usuarios_nonce' not in st.session_state: st.session_state.usuarios_nonce=0
    bloqueados=read_df("SELECT id,usuario,nombre,rol,COALESCE(intentos_fallidos,0) AS intentos_fallidos FROM usuarios WHERE activo=0 ORDER BY intentos_fallidos DESC,usuario")
    with st.expander(f'🔓 Habilitación de cuentas bloqueadas ({len(bloqueados)})',expanded=not bloqueados.empty):
        if bloqueados.empty:
            st.info('No hay usuarios bloqueados.')
        else:
            st.dataframe(bloqueados.rename(columns={'id':'ID','usuario':'Usuario','nombre':'Nombre','rol':'Rol','intentos_fallidos':'Intentos fallidos'}),use_container_width=True,hide_index=True)
            opciones_bloqueados={f"{r.usuario} | {r.nombre}":int(r.id) for r in bloqueados.itertuples()}
            seleccion_bloqueado=st.selectbox('Cuenta bloqueada',list(opciones_bloqueados),key='usuario_bloqueado_reactivar')
            if st.button('Reactivar cuenta y reiniciar intentos',type='primary',key='reactivar_usuario_bloqueado'):
                bid=opciones_bloqueados[seleccion_bloqueado]
                usuario_bloqueado=str(bloqueados[bloqueados.id==bid].iloc[0].usuario)
                exec_sql('UPDATE usuarios SET activo=1,intentos_fallidos=0 WHERE id=?',(bid,))
                audit(st.session_state.auth['usuario'],'REACTIVAR_USUARIO_BLOQUEADO',f'ID {bid} | {usuario_bloqueado}')
                st.session_state.usuarios_nonce+=1
                st.success('Cuenta reactivada correctamente. Los intentos fallidos se reiniciaron.')
                st.rerun()
    st.subheader('Crear usuario')
    with st.expander('Agregar nuevo usuario',expanded=False):
        with st.form('crear_usuario_form',clear_on_submit=True):
            c1,c2=st.columns(2)
            nuevo_usuario=c1.text_input('Usuario *',placeholder='Ejemplo: jperez')
            nuevo_nombre=c2.text_input('Nombre completo *',placeholder='Ejemplo: Juan Pérez')
            c3,c4=st.columns(2)
            nueva_password=c3.text_input('Contraseña inicial *',type='password')
            nuevo_rol=c4.selectbox('Rol *',['usuario','desarrollador'],format_func=lambda x:'Usuario' if x=='usuario' else 'Administrador / desarrollador')
            crear=st.form_submit_button('Crear usuario',type='primary')
        if crear:
            usuario=nuevo_usuario.strip(); nombre=nuevo_nombre.strip(); password=nueva_password.strip()
            faltan=[]
            if not usuario:faltan.append('Usuario')
            if not nombre:faltan.append('Nombre completo')
            if not password:faltan.append('Contraseña inicial')
            if faltan: st.error('Completa los campos obligatorios: '+', '.join(faltan)+'.')
            elif not read_df('SELECT id FROM usuarios WHERE usuario=?',(usuario,)).empty: st.error('No fue posible crear la cuenta porque el nombre de usuario ya está registrado.')
            else:
                exec_sql('INSERT INTO usuarios(usuario,nombre,password_hash,rol,activo,intentos_fallidos,requiere_cambio_pass,creado_en) VALUES(?,?,?,?,1,0,1,?)',(usuario,nombre,hash_password(password),nuevo_rol,now_iso()))
                audit(st.session_state.auth['usuario'],'CREAR_USUARIO',f'Usuario {usuario} | Rol {nuevo_rol}')
                st.session_state.usuarios_nonce+=1; st.success('Usuario creado correctamente.'); st.rerun()
    st.subheader('Usuarios registrados')
    usuarios=read_df("""SELECT id,usuario,nombre,rol,COALESCE(intentos_fallidos,0) AS intentos_fallidos,COALESCE(requiere_cambio_pass,0) AS requiere_cambio_pass,CASE WHEN activo=1 THEN 'Habilitado' ELSE 'Bloqueado / Inhabilitado' END AS estado,creado_en FROM usuarios ORDER BY id""")
    if usuarios.empty:
        st.info('No se encontraron usuarios registrados.'); return
    vista=usuarios.rename(columns={'id':'ID','usuario':'Usuario','nombre':'Nombre completo','rol':'Rol','intentos_fallidos':'Intentos fallidos','requiere_cambio_pass':'Cambio de contraseña pendiente','estado':'Estado','creado_en':'Fecha de creación'})
    evento=st.dataframe(vista,use_container_width=True,hide_index=True,on_select='rerun',selection_mode='single-row',key=f'usuarios_tabla_{st.session_state.usuarios_nonce}')
    filas=getattr(evento,'selection',{}).get('rows',[]) if evento is not None else []
    valido=bool(filas) and isinstance(filas[0],int) and 0<=filas[0]<len(vista)
    if not valido:
        st.info('Selecciona un usuario en la tabla para administrar la cuenta.')
        return
    rid=int(vista.iloc[filas[0]]['ID'])
    actual_df=read_df('SELECT * FROM usuarios WHERE id=?',(rid,))
    if actual_df.empty:
        st.session_state.usuarios_nonce+=1; st.rerun()
    r=actual_df.iloc[0]
    st.markdown(f"### Usuario seleccionado: {r['usuario']}")
    with st.expander('Editar información y rol',expanded=True):
        with st.form(f'editar_usuario_{rid}'):
            a,b=st.columns(2)
            usuario_editado=a.text_input('Usuario *',value=str(r['usuario'] or ''))
            nombre_editado=b.text_input('Nombre completo *',value=str(r['nombre'] or ''))
            c,d=st.columns(2)
            roles=['usuario','desarrollador']
            rol_editado=c.selectbox('Rol *',roles,index=idx_or_zero(roles,str(r['rol'] or 'usuario')),format_func=lambda x:'Usuario' if x=='usuario' else 'Administrador / desarrollador')
            password_nueva=d.text_input('Nueva contraseña',type='password',help='Déjala vacía para conservar la contraseña actual.')
            guardar=st.form_submit_button('Guardar cambios',type='primary')
        if guardar:
            usuario=usuario_editado.strip(); nombre=nombre_editado.strip(); nueva=password_nueva.strip()
            if not usuario or not nombre: st.error('Completa el usuario y el nombre completo.')
            elif not read_df('SELECT id FROM usuarios WHERE usuario=? AND id<>?',(usuario,rid)).empty: st.error('No fue posible guardar porque el nombre de usuario ya está asignado a otra cuenta.')
            elif str(r['usuario'])==ADMIN_USER and rol_editado!='desarrollador': st.error('La cuenta administradora principal debe conservar el rol de administrador/desarrollador.')
            else:
                if nueva:
                    exec_sql('UPDATE usuarios SET usuario=?,nombre=?,rol=?,password_hash=?,requiere_cambio_pass=1,intentos_fallidos=0 WHERE id=?',(usuario,nombre,rol_editado,hash_password(nueva),rid))
                else:
                    exec_sql('UPDATE usuarios SET usuario=?,nombre=?,rol=? WHERE id=?',(usuario,nombre,rol_editado,rid))
                audit(st.session_state.auth['usuario'],'EDITAR_USUARIO',f'ID {rid} | Usuario {usuario} | Rol {rol_editado}')
                if rid==int(read_df('SELECT id FROM usuarios WHERE usuario=?',(st.session_state.auth['usuario'],)).iloc[0].id):
                    st.session_state.auth.update({'usuario':usuario,'nombre':nombre,'rol':rol_editado})
                st.session_state.usuarios_nonce+=1; st.success('La información del usuario se actualizó correctamente.'); st.rerun()
    activo=int(r['activo'] or 0)==1
    c1,c2=st.columns(2)
    if activo:
        inhabilitar=c1.button('Inhabilitar usuario',key=f'inhabilitar_usuario_{rid}',use_container_width=True)
        habilitar=False
    else:
        habilitar=c1.button('Habilitar usuario',key=f'habilitar_usuario_{rid}',use_container_width=True)
        inhabilitar=False
    eliminar=c2.button('Eliminar usuario',key=f'eliminar_usuario_{rid}',use_container_width=True)
    if inhabilitar:
        if str(r['usuario'])==st.session_state.auth['usuario']: st.error('No puedes inhabilitar la cuenta utilizada en la sesión actual.')
        elif str(r['usuario'])==ADMIN_USER: st.error('La cuenta administradora principal no puede inhabilitarse.')
        else:
            exec_sql('UPDATE usuarios SET activo=0 WHERE id=?',(rid,)); audit(st.session_state.auth['usuario'],'INHABILITAR_USUARIO',f'ID {rid} | {r["usuario"]}'); st.session_state.usuarios_nonce+=1; st.success('Usuario inhabilitado correctamente.'); st.rerun()
    if habilitar:
        exec_sql('UPDATE usuarios SET activo=1,intentos_fallidos=0 WHERE id=?',(rid,)); audit(st.session_state.auth['usuario'],'HABILITAR_USUARIO',f'ID {rid} | {r["usuario"]}'); st.session_state.usuarios_nonce+=1; st.success('Usuario habilitado correctamente.'); st.rerun()
    if eliminar: st.session_state.usuario_confirmar_eliminacion=rid
    if st.session_state.get('usuario_confirmar_eliminacion')==rid:
        st.warning('La cuenta se eliminará de forma permanente. Los registros históricos y de auditoría asociados conservarán el nombre de usuario.')
        x,y=st.columns(2)
        if x.button('Confirmar eliminación',key=f'confirmar_usuario_{rid}',use_container_width=True):
            if str(r['usuario'])==st.session_state.auth['usuario']: st.error('No puedes eliminar la cuenta utilizada en la sesión actual.')
            elif str(r['usuario'])==ADMIN_USER: st.error('La cuenta administradora principal no puede eliminarse.')
            else:
                audit(st.session_state.auth['usuario'],'ELIMINAR_USUARIO',f'ID {rid} | {r["usuario"]}')
                exec_sql('DELETE FROM usuarios WHERE id=?',(rid,)); reset_autoincrement('usuarios'); st.session_state.pop('usuario_confirmar_eliminacion',None); st.session_state.usuarios_nonce+=1; st.success('Usuario eliminado correctamente.'); st.rerun()
        if y.button('Cancelar',key=f'cancelar_usuario_{rid}',use_container_width=True):
            st.session_state.pop('usuario_confirmar_eliminacion',None); st.session_state.usuarios_nonce+=1; st.rerun()
def page_auditoria(): st.title('Auditoría'); st.dataframe(read_df('SELECT * FROM auditoria ORDER BY id DESC LIMIT 1000'),use_container_width=True,hide_index=True)

def main():
    init_state(); init_db()
    if FORCE_RESET_ADMIN and ADMIN_PASS:
        reset_admin()
    user=login()
    cambiar_password_obligatoria()
    styles(False)
    left_col, right_col = st.columns([0.19, 0.81], gap='large')
    with left_col:
        left_menu()
    with right_col:
        topbar(user)
        page=st.session_state.page
        if page=='Inicio': page_inicio()
        elif page=='Nuevo registro': page_registro()
        elif page=='Consulta y descarga': page_consulta()
        elif page=='Muestras de retención': page_muestras_retencion()
        elif page=='Entrega de turno': page_entrega_turno()
        elif page=='Catálogos': page_catalogos()
        elif page=='Usuarios': page_usuarios()
        elif page=='Auditoría': page_auditoria()

if __name__=='__main__': main()
