import base64
from datetime import date, datetime
from io import BytesIO

import pandas as pd
from openpyxl.styles import Font, PatternFill

from base_datos import consultar
from configuracion import LOGO_MUNDO_DULCE_BASE64
from servicios import normalizar_catalogo

def _consultar_df(tabla, filtros=None, columnas="*", ordenar_por=None, limite=None):
    """Consulta general para reportes, independiente del motor de base de datos.

    reportes.py solo depende del contrato publico base_datos.consultar().
    Si el motor cambia, se sustituye la implementacion de base_datos.py y los
    reportes permanecen sin cambios.
    """
    resultado = consultar(
        tabla=tabla,
        columnas=columnas,
        filtros=filtros or {},
        ordenar_por=ordenar_por or [],
        limite=limite,
    )
    if isinstance(resultado, pd.DataFrame):
        return resultado
    if resultado is None:
        return pd.DataFrame()
    if isinstance(resultado, dict):
        return pd.DataFrame([resultado])
    return pd.DataFrame(resultado)


def excel_matriz(df):
    b = BytesIO()
    with pd.ExcelWriter(b, engine='openpyxl') as w:
        base = df.rename(columns={'fecha': 'Fecha', 'analista': 'Analista', 'total_carga_datos': 'Total carga de datos', 'horas_nave1': 'Horas Nave 1', 'horas_nave2': 'Horas Nave 2', 'horas_nave3': 'Horas Nave 3'})
        base.to_excel(w, sheet_name='Base indicadores', index=False)
        for c, n in [('total_carga_datos', 'Total carga'), ('horas_nave1', 'Horas Nave 1'), ('horas_nave2', 'Horas Nave 2'), ('horas_nave3', 'Horas Nave 3')]:
            df.pivot_table(index='analista', columns='fecha', values=c, aggfunc='sum', fill_value=0).to_excel(w, sheet_name=n)
        for ws in w.book.worksheets:
            ws.freeze_panes = 'A2'
            for cell in ws[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='062C36')
    return b.getvalue()

def pdf_pnc(rid):
    """Genera el formato corporativo de Producto No Conforme del registro seleccionado."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase.pdfmetrics import stringWidth
        from reportlab.lib.utils import ImageReader
    except ModuleNotFoundError:
        return None
    df = _consultar_df('pnc_registros', {'id': rid}, limite=1)
    if df.empty:
        return None
    r = df.iloc[0].to_dict()

    def valor(campo):
        v = r.get(campo, '')
        return '' if v is None or pd.isna(v) else str(v)

    def texto(c, contenido, x, y, tam=8, negrita=False):
        c.setFont('Helvetica-Bold' if negrita else 'Helvetica', tam)
        c.drawString(x, y, str(contenido or ''))

    def ajustar(c, contenido, x, y, ancho, tam=8, negrita=False):
        contenido = str(contenido or '').replace('\n', ' / ')
        fuente = 'Helvetica-Bold' if negrita else 'Helvetica'
        actual = tam
        while actual > 5.5 and stringWidth(contenido, fuente, actual) > ancho:
            actual -= 0.25
        if stringWidth(contenido, fuente, actual) > ancho:
            while contenido and stringWidth(contenido + '...', fuente, actual) > ancho:
                contenido = contenido[:-1]
            contenido += '...'
        c.setFont(fuente, actual)
        c.drawString(x, y, contenido)

    def parrafo(c, contenido, x, y, ancho, tam=8.5, interlineado=11, max_lineas=7):
        lineas = []
        for bloque in str(contenido or '').splitlines() or ['']:
            actual = ''
            for palabra in bloque.split():
                prueba = (actual + ' ' + palabra).strip()
                if stringWidth(prueba, 'Helvetica', tam) <= ancho:
                    actual = prueba
                else:
                    if actual:
                        lineas.append(actual)
                    actual = palabra
            if actual:
                lineas.append(actual)
        c.setFont('Helvetica', tam)
        for i, linea in enumerate(lineas[:max_lineas]):
            c.drawString(x, y - i * interlineado, linea)
        if len(lineas) > max_lineas:
            c.drawRightString(x + ancho, y - (max_lineas - 1) * interlineado, '...')
    b = BytesIO()
    c = canvas.Canvas(b, pagesize=A4)
    W, H = A4
    c.setTitle('Informe de Producto No Conforme')
    c.setLineWidth(1.1)
    c.rect(10, 10, W - 20, H - 20)
    izq = 58
    der = W - 58
    ancho = der - izq
    sup = 742
    alto = 52
    logo = 132
    revision = 58
    c.setLineWidth(0.7)
    c.rect(izq, sup - alto, ancho, alto)
    c.line(izq + logo, sup - alto, izq + logo, sup)
    c.line(der - revision, sup - alto, der - revision, sup)
    try:
        logo_bytes = BytesIO(base64.b64decode(LOGO_MUNDO_DULCE_BASE64))
        c.drawImage(ImageReader(logo_bytes), izq + 8, sup - alto + 7, width=logo - 16, height=alto - 14, preserveAspectRatio=True, anchor='c', mask='auto')
    except Exception:
        texto(c, 'Mundo Dulce', izq + 28, sup - 31, 11, True)
    centro_izq = izq + logo
    centro_der = der - revision
    c.setFont('Helvetica', 8.3)
    c.drawCentredString((centro_izq + centro_der) / 2, sup - 19, 'Anexo 5. Informe de Producto No Conforme')
    c.drawCentredString((centro_izq + centro_der) / 2, sup - 39, 'PG-CAL01-2301-01760-2007')
    c.drawCentredString(der - revision / 2, sup - 31, 'Rev. 0')
    y = 665
    texto(c, 'PNC No.', der - 100, y + 5, 8.5)
    c.rect(der - 58, y, 58, 18)
    fecha_pnc = valor('fecha_apertura')
    try:
        anio_pnc = str(pd.to_datetime(fecha_pnc).year)
    except Exception:
        anio_pnc = str(datetime.now().year)
    ajustar(c, f'{rid}/{anio_pnc}', der - 55, y + 5, 52, 8, True)
    producto = ' - '.join((x for x in [valor('item'), valor('descripcion_producto')] if x))
    cantidad = valor('cantidad_observada')
    try:
        cantidad = f'{float(cantidad):.2f} kg'
    except Exception:
        pass
    responsables = ' / '.join((x for x in [valor('supervisor'), valor('analista')] if x))
    codigo_nc = '_'.join((x.strip() for x in [valor('etapa'), valor('codigo_defecto')] if x.strip()))
    filas = [('Fecha:', valor('fecha_apertura')), ('Sector:', valor('linea_sector')), ('Código de la No Conformidad:', codigo_nc), ('Item y Producto o SE:', producto), ('Lote:', valor('lote')), ('Cantidad observada:', cantidad), ('Responsables:', responsables)]
    arriba = 646
    fila_alto = 19
    etiqueta = 150
    c.rect(izq, arriba - fila_alto * len(filas), ancho, fila_alto * len(filas))
    c.line(izq + etiqueta, arriba - fila_alto * len(filas), izq + etiqueta, arriba)
    for i, (nombre, dato) in enumerate(filas):
        limite = arriba - i * fila_alto
        if i:
            c.line(izq, limite, der, limite)
        texto(c, nombre, izq + 2, limite - fila_alto + 6, 7.8)
        ajustar(c, dato, izq + etiqueta + 3, limite - fila_alto + 6, ancho - etiqueta - 6, 7.8)
    acciones_sup = 495
    acciones_alto = 112
    c.rect(izq, acciones_sup - acciones_alto, ancho, acciones_alto)
    c.line(izq, acciones_sup - 18, der, acciones_sup - 18)
    c.setFont('Helvetica', 8.5)
    c.drawCentredString(izq + ancho / 2, acciones_sup - 13, 'Acciones inmediatas')
    parrafo(c, valor('acciones_inmediatas'), izq + 6, acciones_sup - 32, ancho - 12)
    responsable_y = 348
    c.rect(izq, responsable_y, ancho, 18)
    ajustar(c, 'Responsable: ' + (valor('analista') or valor('supervisor')), izq + 3, responsable_y + 5, ancho - 6, 8)
    obs_sup = 331
    obs_alto = 115
    c.rect(izq, obs_sup - obs_alto, ancho, obs_alto)
    c.line(izq, obs_sup - 18, der, obs_sup - 18)
    c.setFont('Helvetica', 8.5)
    c.drawCentredString(izq + ancho / 2, obs_sup - 13, 'Observaciones')
    observaciones = valor('observaciones')
    if valor('descripcion_defecto'):
        observaciones = ('Descripción de la no conformidad: ' + valor('descripcion_defecto') + '\n' + observaciones).strip()
    parrafo(c, observaciones, izq + 6, obs_sup - 32, ancho - 12)
    c.showPage()
    c.save()
    return b.getvalue()

def pdf_pnc_fisico(rid):
    """Genera el Registro Fisico PNC en A4, llenado unicamente hasta Disposicion."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase.pdfmetrics import stringWidth
        from reportlab.lib.utils import ImageReader
    except ModuleNotFoundError:
        return None
    df = _consultar_df('pnc_registros', {'id': rid}, limite=1)
    if df.empty:
        return None
    r = df.iloc[0].to_dict()

    def v(campo):
        x = r.get(campo, '')
        return '' if x is None or pd.isna(x) else str(x).strip()

    def fecha_es(x):
        try:
            return pd.to_datetime(x).strftime('%d/%m/%Y')
        except Exception:
            return str(x or '')

    def t(c, texto, x, y, tam=10, negrita=False):
        c.setFont('Helvetica-Bold' if negrita else 'Helvetica', tam)
        c.drawString(x, y, str(texto or ''))

    def fit(c, texto, x, y, ancho, tam=10, negrita=False, centrado=False):
        texto = str(texto or '').replace('\n', ' / ')
        fuente = 'Helvetica-Bold' if negrita else 'Helvetica'
        actual = tam
        while actual > 5.5 and stringWidth(texto, fuente, actual) > ancho:
            actual -= 0.25
        if stringWidth(texto, fuente, actual) > ancho:
            while texto and stringWidth(texto + '...', fuente, actual) > ancho:
                texto = texto[:-1]
            texto += '...'
        c.setFont(fuente, actual)
        if centrado:
            c.drawCentredString(x + ancho / 2, y, texto)
        else:
            c.drawString(x, y, texto)

    def logo(c, x, y, w, h):
        try:
            data = BytesIO(base64.b64decode(LOGO_MUNDO_DULCE_BASE64))
            c.drawImage(ImageReader(data), x, y, width=w, height=h, preserveAspectRatio=True, anchor='c', mask='auto')
        except Exception:
            fit(c, 'Mundo Dulce', x, y + h / 2, w, 11, True, True)

    def encabezado(c, reverso=False):
        x0, x1, x2, x3 = (38, 132, 494, 557)
        top, bot = (805, 756)
        c.setLineWidth(0.8)
        c.rect(x0, bot, x3 - x0, top - bot)
        c.line(x1, bot, x1, top)
        c.line(x2, bot, x2, top)
        logo(c, x0 + 6, bot + 5, x1 - x0 - 12, top - bot - 10)
        fit(c, 'Anexo 1. Registro de Productos No Conformes', x1 + 3, top - 17, 278, 10, False, True)
        fit(c, 'Rev. 3', 412, top - 17, 72, 10, False, True)
        fit(c, 'PG-CAL01-2301-01760-2007', x1 + 4, bot + 7, x2 - x1 - 8, 9.5, False, True)
        fit(c, 'Reverso' if reverso else 'NUMERO', x2 + 2, top - 28, x3 - x2 - 4, 10, False, True)
    b = BytesIO()
    c = canvas.Canvas(b, pagesize=A4)
    W, H = A4
    c.setTitle('Registro Fisico de Producto No Conforme')
    encabezado(c)
    try:
        anio = str(pd.to_datetime(v('fecha_apertura')).year)
    except Exception:
        anio = str(datetime.now().year)
    fit(c, f'{rid:03d}/{anio[-2:]}', 496, 765, 59, 10, True, True)
    t(c, 'Sector:', 39, 733, 10)
    c.line(79, 731, 404, 731)
    fit(c, v('linea_sector'), 82, 734, 319, 10)
    x0, x3 = (38, 557)
    top_id, bot_id = (711, 498)
    c.setLineWidth(0.8)
    c.rect(x0, bot_id, x3 - x0, top_id - bot_id)
    c.rect(x0, top_id - 18, x3 - x0, 18)
    c.setFont('Helvetica-Bold', 10)
    c.drawCentredString((x0 + x3) / 2, top_id - 13, 'IDENTIFICACION')
    codigo = '_'.join((x for x in [v('etapa'), v('codigo_defecto')] if x))
    cantidad = ''
    if v('cantidad_observada'):
        try:
            cantidad = f"{float(v('cantidad_observada')):.2f} kg"
        except Exception:
            cantidad = v('cantidad_observada')
    campos = [('Código:', codigo, 681), ('Fecha:', fecha_es(v('fecha_apertura')), 648), ('ITEM:', v('item'), 615), ('Lote:', v('lote'), 582), ('Cantidad observada:', cantidad, 549), ('Responsable / Persona que detecta:', ' / '.join((x for x in [v('supervisor'), v('analista')] if x)), 516)]
    for etiqueta, valor, y in campos:
        tam_campo = 9.6
        t(c, etiqueta, 40, y, tam_campo)
        valor_x = 40 + stringWidth(etiqueta, 'Helvetica', tam_campo) + 9
        fit(c, valor, valor_x, y, x3 - valor_x - 6, tam_campo)
    etiqueta_producto = 'Producto o Semielaborado:'
    t(c, etiqueta_producto, 300, 615, 9.6)
    producto_x = 300 + stringWidth(etiqueta_producto, 'Helvetica', 9.6) + 9
    fit(c, v('descripcion_producto'), producto_x, 615, x3 - producto_x - 6, 9.2)
    top_d, bot_d = (482, 350)
    medio = 300
    c.rect(x0, bot_d, x3 - x0, top_d - bot_d)
    c.rect(x0, top_d - 18, x3 - x0, 18)
    c.setFont('Helvetica-Bold', 10)
    c.drawCentredString((x0 + x3) / 2, top_d - 13, 'DISPOSICION')
    c.line(medio, bot_d, medio, top_d - 18)
    t(c, 'Marque con una X el recuadro correspondiente:', 40, 451, 9.3)
    t(c, 'Fecha, nombre y firma de responsable de disposición:', 304, 451, 9.1)
    opciones = ['Reproceso', 'Retrabajo', 'Decomiso', 'Inspeccion', 'Aprobado en segunda instancia', 'Otro']
    actual = v('disposicion').lower()
    for i, op in enumerate(opciones):
        y = 430 - i * 16
        c.rect(40, y - 3, 82, 16)
        t(c, op + ':', 124, y + 1, 9.3)
        normal_op = op.lower().replace('inspeccion', 'inspección')
        if normal_op == actual or op.lower() == actual:
            c.setFont('Helvetica-Bold', 12)
            c.drawCentredString(81, y, 'X')
    fit(c, fecha_es(v('fecha_apertura')), 307, 430, 242, 10, False, True)
    fit(c, v('analista'), 307, 365, 242, 10, False, True)
    top_t, bot_t = (334, 216)
    c.rect(x0, bot_t, x3 - x0, top_t - bot_t)
    c.rect(x0, top_t - 18, x3 - x0, 18)
    c.setFont('Helvetica-Bold', 10)
    c.drawCentredString((x0 + x3) / 2, top_t - 13, 'TRATAMIENTO')
    ancho_t = x3 - x0
    col1 = x0 + ancho_t / 3
    col2 = x0 + ancho_t * 2 / 3
    for xx in (col1, col2):
        c.line(xx, bot_t, xx, top_t - 18)
    separador_y = top_t - 67
    c.line(x0, separador_y, x3, separador_y)
    titulos = [('Kg Aprobados en\nsegunda instancia', x0 + 3, top_t - 31), ('Kg Decomiso', col1 + 3, top_t - 31), ('Kg Reproceso', col2 + 3, top_t - 31), ('Fecha comienzo del tratamiento', x0 + 3, separador_y - 15), ('Fecha final de tratamiento', col1 + 3, separador_y - 15), ('Nombre y firma del responsable del\ntratamiento:', col2 + 3, separador_y - 15)]
    for texto, x, y in titulos:
        for j, linea in enumerate(texto.split('\n')):
            t(c, linea, x, y - j * 11, 8.9)
    top_c, bot_c = (198, 76)
    c.rect(x0, bot_c, x3 - x0, top_c - bot_c)
    c.rect(x0, top_c - 18, x3 - x0, 18)
    c.setFont('Helvetica-Bold', 10)
    c.drawCentredString((x0 + x3) / 2, top_c - 13, 'CONSIDERACIONES DE TRABAJO')
    filas = [('Dias de retrabajo:', top_c - 34), ('No. de personas:', top_c - 52), ('Materiales:', top_c - 70), ('Otro:', top_c - 88)]
    for etiqueta, y in filas:
        t(c, etiqueta, x0 + 2, y, 9)
        c.line(x0, y - 5, x3, y - 5)
    c.line(300, bot_c, 300, top_c - 93)
    t(c, 'Fecha:', x0 + 2, bot_c + 7, 9)
    t(c, 'Nombre y firma:', 302, bot_c + 7, 9)
    c.showPage()
    encabezado(c, True)
    TAM_REVERSO = 10.5
    X_REVERSO = 38
    ANCHO_REVERSO = 519

    def lineas_reverso(texto, sangria=0, negrita=False, espacio_despues=0):
        nonlocal yy
        fuente = 'Helvetica-Bold' if negrita else 'Helvetica'
        palabras = str(texto or '').split()
        linea = ''
        ancho = ANCHO_REVERSO - sangria
        for palabra in palabras:
            prueba = (linea + ' ' + palabra).strip()
            if stringWidth(prueba, fuente, TAM_REVERSO) <= ancho:
                linea = prueba
            else:
                c.setFont(fuente, TAM_REVERSO)
                c.drawString(X_REVERSO + sangria, yy, linea)
                yy -= 15
                linea = palabra
        if linea:
            c.setFont(fuente, TAM_REVERSO)
            c.drawString(X_REVERSO + sangria, yy, linea)
            yy -= 15
        yy -= espacio_despues
    yy = 727
    lineas_reverso('Instrucciones para el uso del Registro de Productos No Conformes:', espacio_despues=21)
    lineas_reverso('ENCABEZADO:', negrita=True, espacio_despues=3)
    lineas_reverso('En el cuadro superior derecho se coloca el número consecutivo del registro que se genera. La numeración se compone de la siguiente manera:', sangria=22)
    c.setFont('Helvetica-Bold', TAM_REVERSO)
    c.drawCentredString((38 + 557) / 2, yy, 'NNN / AA')
    yy -= 21
    lineas_reverso('dónde:', espacio_despues=0)
    lineas_reverso('- NNN es el número consecutivo', sangria=16)
    lineas_reverso('- AA son los dos últimos dígitos del año en curso.', sangria=16, espacio_despues=8)
    lineas_reverso('IDENTIFICACIÓN:', negrita=True, espacio_despues=3)
    for texto_info in ['- Código: se coloca el código de retención de acuerdo con el Anexo 4. Codificación de retención', '- Fecha en la que se produjo el Producto No Conforme', '- ITEM de Semielaborado o Producto Terminado', '- Lote(s) del producto', '- Categoría de retención inicial', '- Cantidad observada: Producto que potencialmente o no cumple con las especificaciones']:
        lineas_reverso(texto_info, sangria=16)
    yy -= 8
    lineas_reverso('DISPOSICIÓN:', negrita=True, espacio_despues=3)
    lineas_reverso('En el casillero de disposición marcar con una X acorde a la definición y colocar fecha, nombre y firma de dicha disposición.', sangria=22, espacio_despues=8)
    lineas_reverso('TRATAMIENTO:', negrita=True, espacio_despues=3)
    lineas_reverso('En el casillero la fecha del comienzo y final del tratamiento, firma de la persona que llevó a cabo el tratamiento del producto.', sangria=22)
    lineas_reverso('Detallar Kg aprobados en segunda instancia, Kg decomiso y kg reproceso acorde con el tratamiento realizado.', sangria=22, espacio_despues=4)
    linea_categoria_y = yy
    lineas_reverso('Categoría final: Después del tratamiento a qué categoría pasó')
    c.line(X_REVERSO, linea_categoria_y - 3, X_REVERSO + 390, linea_categoria_y - 3)
    yy -= 10
    lineas_reverso('CONSIDERACIONES DE TRABAJO:', negrita=True, espacio_despues=3)
    lineas_reverso('Colocar la información que para determinar el costo de la NO Calidad (Personas, materiales y tiempo del retrabajo) más la fecha y la firma de la persona que tuvo a cargo de validar dicha tarea. En caso de que el retrabajo se realice en diferentes días o turnos detallar información.', sangria=22)
    lineas_reverso('Este casillero solo lo puede completar personal del área de calidad.')
    c.showPage()
    c.save()
    return b.getvalue()

def _pdf_hallazgo_base(tabla, rid, tipo_formato):
    """Genera los formatos ME y Detector/RX sin alterar los datos ni la logica operativa."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase.pdfmetrics import stringWidth
        from reportlab.lib.utils import ImageReader
    except ModuleNotFoundError:
        return None
    df = _consultar_df(tabla, {'id': rid}, limite=1)
    if df.empty:
        return None
    r = df.iloc[0].to_dict()

    def v(campo):
        x = r.get(campo, '')
        return '' if x is None or pd.isna(x) else str(x)

    def fecha_registro():
        try:
            return date(int(float(v('anio'))), int(float(v('mes'))), int(float(v('dia')))).strftime('%d/%m/%Y')
        except Exception:
            return ''

    def t(c, texto, x, y, tam=8, negrita=False):
        c.setFont('Helvetica-Bold' if negrita else 'Helvetica', tam)
        c.drawString(x, y, str(texto or ''))

    def fit(c, texto, x, y, ancho, tam=8, negrita=False):
        texto = str(texto or '').replace('\n', ' / ')
        fuente = 'Helvetica-Bold' if negrita else 'Helvetica'
        actual = tam
        while actual > 5.5 and stringWidth(texto, fuente, actual) > ancho:
            actual -= 0.25
        if stringWidth(texto, fuente, actual) > ancho:
            while texto and stringWidth(texto + '...', fuente, actual) > ancho:
                texto = texto[:-1]
            texto += '...'
        c.setFont(fuente, actual)
        c.drawString(x, y, texto)

    def wrap(c, texto, x, y, ancho, tam=8, leading=11, max_lines=8):
        lineas = []
        for bloque in str(texto or '').splitlines() or ['']:
            linea = ''
            for palabra in bloque.split():
                prueba = (linea + ' ' + palabra).strip()
                if stringWidth(prueba, 'Helvetica', tam) <= ancho:
                    linea = prueba
                else:
                    if linea:
                        lineas.append(linea)
                    linea = palabra
            if linea:
                lineas.append(linea)
        c.setFont('Helvetica', tam)
        for i, linea in enumerate(lineas[:max_lines]):
            c.drawString(x, y - i * leading, linea)
        if len(lineas) > max_lines:
            c.drawRightString(x + ancho, y - (max_lines - 1) * leading, '...')

    def line_value(c, etiqueta, valor, y, label_x=45, value_x=108, right=557, bold=False):
        t(c, etiqueta, label_x, y, 7.7, bold)
        c.line(value_x, y - 2, right, y - 2)
        fit(c, valor, value_x + 3, y, right - value_x - 6, 7.7)

    def logo(c, x, y, w, h):
        try:
            data = BytesIO(base64.b64decode(LOGO_MUNDO_DULCE_BASE64))
            c.drawImage(ImageReader(data), x, y, width=w, height=h, preserveAspectRatio=True, anchor='c', mask='auto')
        except Exception:
            t(c, 'Mundo Dulce', x + 12, y + h / 2, 10, True)
    b = BytesIO()
    c = canvas.Canvas(b, pagesize=A4)
    W, H = A4
    c.setLineWidth(0.8)
    if tipo_formato == 'ME':
        c.setTitle('Reporte de Hallazgos de Materia Extraña')
        x0 = 45
        x1 = 198
        x2 = 508
        x3 = 558
        top = 800
        bot = 758
        c.rect(x0, bot, x3 - x0, top - bot)
        c.line(x1, bot, x1, top)
        c.line(x2, bot, x2, top)
        logo(c, x0 + 15, bot + 6, x1 - x0 - 30, top - bot - 12)
        c.setFont('Helvetica', 7.7)
        c.drawCentredString((x1 + x2) / 2, top - 17, 'Anexo 1. Reporte de Hallazgos de Materia Extraña')
        c.drawCentredString((x1 + x2) / 2, top - 32, 'IT-CAL03-2301-03')
        c.drawCentredString((x2 + x3) / 2, top - 25, 'Rev. 0')
        t(c, 'N° Hallazgo:', 46, 733, 7.5)
        fit(c, f"{rid}/{v('anio') or datetime.now().year}", 108, 733, 140, 7.5, True)
        t(c, 'Fecha:', 46, 718, 7.5)
        fit(c, fecha_registro(), 108, 718, 140, 7.5)
        line_value(c, 'Línea:', v('linea_sector'), 688)
        line_value(c, 'Familia:', v('familia'), 673)
        line_value(c, 'Equipo:', v('equipo_hallazgo'), 658)
        line_value(c, 'Producto:', f"{v('item')} - {v('producto')}".strip(' -'), 643)
        line_value(c, 'Lote:', v('lote'), 628)
        t(c, 'Descripción de material hallado:', 46, 596, 7.5)
        c.line(222, 594, 557, 594)
        c.line(45, 579, 557, 579)
        c.line(45, 564, 557, 564)
        wrap(c, v('descripcion_hallazgo'), 225, 596, 329, 7.5, 14, 3)
        c.rect(45, 339, 512, 210)
        t(c, 'Muestra Hallazgo de Materia Extraña:', 48, 537, 7.2)
        detalle = 'Tipo: ' + v('tipo') + ' | Partículas halladas: ' + v('particulas_halladas')
        fit(c, detalle, 48, 522, 500, 7)
        t(c, 'Acción contingente:', 46, 312, 7.5)
        c.line(45, 296, 557, 296)
        c.line(45, 281, 557, 281)
        wrap(c, v('accion_contingente') or v('acciones_inmediatas'), 48, 299, 506, 7.5, 14, 2)
        t(c, 'Investigación del origen:', 46, 254, 7.5)
        c.line(45, 238, 557, 238)
        c.line(45, 223, 557, 223)
        c.line(45, 208, 557, 208)
        wrap(c, v('investigacion_origen'), 48, 241, 506, 7.5, 14, 3)
        for cx, titulo in [(165, 'Calidad'), (306, 'Producción'), (448, 'Mantenimiento')]:
            c.line(cx - 57, 142, cx + 57, 142)
            c.setFont('Helvetica', 7)
            c.drawCentredString(cx, 129, titulo)
            c.drawCentredString(cx, 117, '(Nombre y Firma)')
    else:
        c.setTitle('Registro de Materiales Segregados por Detector de Metales')
        x0 = 27
        x1 = 187
        x2 = 506
        x3 = 568
        top = 805
        mid = 772
        bot = 720
        c.rect(x0, bot, x3 - x0, top - bot)
        c.line(x1, bot, x1, top)
        c.line(x2, bot, x2, top)
        c.line(x1, mid, x2, mid)
        logo(c, x0 + 18, bot + 16, x1 - x0 - 36, top - bot - 30)
        c.setFont('Helvetica-Bold', 7.5)
        c.drawCentredString((x1 + x2) / 2, top - 18, 'Anexo 1. Registro de Materiales Segregados por el detector de')
        c.drawCentredString((x1 + x2) / 2, top - 30, 'metales')
        c.setFont('Helvetica-Bold', 7.1)
        c.drawCentredString((x1 + x2) / 2, mid - 23, 'IT-CAL03-2301-09240-2007 Instructivo para el Tratamiento de')
        c.drawCentredString((x1 + x2) / 2, mid - 36, 'producto separado por Detector de Metales')
        c.drawCentredString((x2 + x3) / 2, top - 44, 'REV: 02')
        t(c, 'RECHAZO No. :', 28, 643, 8, True)
        fit(c, f"{rid}/{v('anio') or datetime.now().year}", 119, 643, 180, 8, True)
        t(c, 'FECHA:', 28, 602, 8, True)
        fit(c, fecha_registro(), 119, 602, 180, 8)
        t(c, 'LINEA DE ELABORACIÓN / EQUIPO:', 28, 547, 8, True)
        c.line(270, 545, x3, 545)
        fit(c, v('linea_sector') + ' / ' + v('equipo_hallazgo'), 273, 547, x3 - 276, 7.8)
        t(c, 'SECTOR DEL HALLAZGO:', 28, 520, 8, True)
        ubic = (v('descripcion_hallazgo') + ' ' + v('equipo_hallazgo')).upper()
        sector_explicito = normalizar_catalogo(v('sector_hallazgo'))
        ubicacion_explicita = normalizar_catalogo(v('ubicacion_material'))
        opciones = [('CONFORMADO', 219, 286), ('ENVOLTURA', 361, 425), ('EMPAQUE', 485, 532)]
        for nombre, label_x, box_x in opciones:
            t(c, nombre, label_x, 515, 7.0)
            c.rect(box_x, 507, 34, 15)
            if normalizar_catalogo(nombre) == sector_explicito or (not sector_explicito and nombre in ubic):
                c.setFont('Helvetica-Bold', 11)
                c.drawCentredString(box_x + 17, 509, 'X')
        line_value(c, 'PRODUCTO:', f"{v('item')} - {v('producto')}".strip(' -'), 468, label_x=28, value_x=119, right=x3, bold=True)
        line_value(c, 'LOTE:', v('lote'), 400, label_x=28, value_x=119, right=x3, bold=True)
        t(c, 'DESCRIPCIÓN DEL MATERIAL HALLADO:', 28, 361, 8, True)
        c.line(320, 359, x3, 359)
        c.line(x0, 334, x3, 334)
        wrap(c, v('descripcion_hallazgo'), 323, 362, x3 - 326, 7.6, 12, 2)
        t(c, 'UBICACIÓN:', 28, 308, 8, True)
        for nombre, label_x, box_x in [('MASA', 190, 226), ('RELLENO', 337, 383), ('RECUBIERTO', 478, 535)]:
            t(c, nombre, label_x, 303, 7.0)
            c.rect(box_x, 295, 30, 15)
            if normalizar_catalogo(nombre) == ubicacion_explicita or (not ubicacion_explicita and nombre in ubic):
                c.setFont('Helvetica-Bold', 11)
                c.drawCentredString(box_x + 15, 297, 'X')
        t(c, 'ADJUNTAR AQUÍ MUESTRA DEL MATERIAL HALLADO', 28, 267, 8, True)
        c.rect(x0, 143, x3 - x0, 105)
        t(c, 'ACCIÓN INMEDIATA:', 28, 112, 8, True)
        c.line(187, 110, x3, 110)
        fit(c, v('acciones_inmediatas') or v('accion_contingente'), 190, 112, x3 - 194, 7.5)
        t(c, 'INVESTIGACIÓN DEL ORIGEN:', 28, 85, 7.5)
        c.line(187, 83, x3, 83)
        fit(c, v('investigacion_origen'), 190, 85, x3 - 194, 7.5)
        for cx, l1, l2 in [(112, 'Nombre y Firma del Analista de', 'Calidad'), (298, 'Nombre y Firma del', 'Supervisor de Producción'), (484, 'Nombre y Firma de', 'Mantenimiento')]:
            c.line(cx - 78, 44, cx + 78, 44)
            c.setFont('Helvetica', 6.7)
            c.drawCentredString(cx, 31, l1)
            c.drawCentredString(cx, 20, l2)
    c.showPage()
    c.save()
    return b.getvalue()

def pdf_materia_extrana(rid):
    return _pdf_hallazgo_base('me_registros', rid, 'ME')

def pdf_detector_metales_rx(rid):
    return _pdf_hallazgo_base('ddm_rx_registros', rid, 'DDM')

def pdf_entrega(eid):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except ModuleNotFoundError:
        return None
    h = _consultar_df('entregas_turno', {'id': eid}, limite=1)
    d = _consultar_df('entregas_turno_lineas', {'entrega_id': eid}, ordenar_por=[('orden_fila', 'asc')])
    m = _consultar_df('matriz_entrega', {'entrega_id': eid}, limite=1)
    sg = _consultar_df('entregas_turno_seguimientos', {'entrega_id': eid}, ordenar_por=[('bloque', 'asc'), ('orden_fila', 'asc'), ('id', 'asc')])
    if h.empty:
        return None
    r = h.iloc[0]
    b = BytesIO()
    doc = SimpleDocTemplate(b, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28)
    sty = getSampleStyleSheet()
    story = [Paragraph('REPORTE DE ENTREGA DE TURNO', sty['Title']), Spacer(1, 8)]
    meta = [['Registro', eid, 'Fecha', r.fecha, 'Analista', r.analista, 'Turno', r.turno], ['Nave', r.nave, 'Referencia', r.referencia, 'Creado por', r.creado_por, 'Creado en', r.creado_en]]
    t = Table(meta)
    t.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.4, colors.grey), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DDF4EF'))]))
    story += [t, Spacer(1, 10)]
    if not m.empty:
        x = m.iloc[0]
        mt = Table([['Total carga', 'Horas Nave 1', 'Horas Nave 2', 'Horas Nave 3'], [f'{x.total_carga_datos:.2f}', f'{x.horas_nave1:.2f}', f'{x.horas_nave2:.2f}', f'{x.horas_nave3:.2f}']], colWidths=[165] * 4)
        mt.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#062C36')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('ALIGN', (0, 0), (-1, -1), 'CENTER'), ('GRID', (0, 0), (-1, -1), 0.4, colors.grey)]))
        story += [mt, Spacer(1, 10)]
    data = [['Nave', 'Grupo', 'Línea/Sector', 'Producto/Análisis', 'Horas', 'Carga', 'Observaciones']]
    for q in d.itertuples():
        data.append([str(getattr(q, 'nave_catalogo', '') or ''), q.grupo or '', q.linea or '', q.producto_descripcion or '', f'{float(q.horas_trabajadas or 0):.2f}', f'{float(q.carga_spac or 0):.2f}', q.observaciones or ''])
    dt = Table(data, colWidths=[55, 85, 145, 145, 45, 45, 190], repeatRows=1)
    dt.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0A4652')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('FONTSIZE', (0, 0), (-1, -1), 7), ('GRID', (0, 0), (-1, -1), 0.3, colors.grey), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))
    story.append(dt)
    if not sg.empty:
        columnas = ['registro_numero', 'hoja_fisica', 'carga_electronica', 'correo', 'descripcion_seguimiento']
        sg = sg[sg[columnas].fillna('').astype(str).apply(lambda c: c.str.strip()).ne('').any(axis=1)]
        if not sg.empty:
            story += [Spacer(1, 12), Paragraph('SEGUIMIENTOS REGISTRADOS', sty['Heading2']), Spacer(1, 6)]
            filas = [['Seguimiento', 'Registro #', 'Hoja física', 'Carga electrónica', 'Correo', 'Descripción']]
            for q in sg.itertuples():
                filas.append([q.bloque or '', q.registro_numero or '', q.hoja_fisica or '', q.carga_electronica or '', q.correo or '', q.descripcion_seguimiento or ''])
            ts = Table(filas, colWidths=[130, 66, 65, 78, 52, 300], repeatRows=1)
            ts.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#062C36')), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('FONTSIZE', (0, 0), (-1, -1), 7), ('GRID', (0, 0), (-1, -1), 0.3, colors.grey), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
            story.append(ts)
    doc.build(story)
    return b.getvalue()
